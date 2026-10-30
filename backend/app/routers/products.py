"""产品信息 + 产品销售订单"""
import io
import re
from datetime import datetime, date
from decimal import Decimal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user
from ..models import Product, ProductSale, Customer, Template
from ..pagination import paginate, ok_page
from ..regions import PROVINCES, CITIES
from ..utils import fmt_dt

router = APIRouter(prefix="/api", tags=["products"])

ALL_REGIONS = set(PROVINCES) | {c for cities in CITIES.values() for c in cities}


def _ulabel(user):
    return (user.nickname or user.username) if user else ""


# ============ 产品信息维护 ============

@router.get("/products")
def list_products(db: Session = Depends(get_db), _=Depends(get_current_user),
                  q: str = "", category: str = "", source: str = "",
                  page: int = 1, per_page: int = 100):
    qy = db.query(Product).filter(Product.deleted_at.is_(None))
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Product.name.like(like), Product.category.like(like)))
    if category:
        qy = qy.filter(Product.category == category)
    if source:
        qy = qy.filter(Product.source == source)
    total, rows = paginate(qy.order_by(Product.id.desc()), page, per_page)
    data = [{"id": p.id, "name": p.name, "category": p.category or "", "source": p.source or "",
             "purchase_price": float(p.purchase_price or 0),
             "default_sales_price": float(p.default_sales_price or 0),
             "note": p.note or "", "status": p.status,
             "created_at": fmt_dt(p.created_at), "created_by": p.created_by or "",
             "updated_at": fmt_dt(p.updated_at), "updated_by": p.updated_by or ""} for p in rows]
    return ok_page(data, total)


@router.post("/products")
def create_product(body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    name = (body.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=422, detail="产品名称不能为空")
    p = Product(name=name,
                category=(body.get("category") or "").strip(),
                source=(body.get("source") or "").strip(),
                purchase_price=_num(body.get("purchase_price")),
                default_sales_price=_num(body.get("default_sales_price")),
                note=(body.get("note") or "").strip(),
                status=body.get("status", 1))
    p.created_by = _ulabel(user)
    p.updated_by = _ulabel(user)
    db.add(p)
    db.commit()
    return {"code": 0, "data": {"id": p.id}, "msg": "产品已创建"}


@router.put("/products/{pid}")
def update_product(pid: int, body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = db.query(Product).filter(Product.id == pid).first()
    if not p:
        raise HTTPException(status_code=404, detail="产品不存在")
    if body.get("name") is not None:
        name = (body.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=422, detail="产品名称不能为空")
        p.name = name
    if body.get("category") is not None:
        p.category = (body.get("category") or "").strip()
    if body.get("source") is not None:
        p.source = (body.get("source") or "").strip()
    if body.get("purchase_price") is not None:
        p.purchase_price = _num(body.get("purchase_price"))
    if body.get("default_sales_price") is not None:
        p.default_sales_price = _num(body.get("default_sales_price"))
    if body.get("note") is not None:
        p.note = (body.get("note") or "").strip()
    if body.get("status") is not None:
        p.status = body.get("status")
    p.updated_by = _ulabel(user)
    db.commit()
    return {"code": 0, "data": None, "msg": "产品已更新"}


@router.delete("/products/{pid}")
def delete_product(pid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    p = db.query(Product).filter(Product.id == pid).first()
    if not p:
        raise HTTPException(status_code=404, detail="产品不存在")
    if db.query(ProductSale).filter(ProductSale.product_id == pid, ProductSale.deleted_at.is_(None)).count() > 0:
        raise HTTPException(status_code=400, detail="该产品已有销售记录，不能删除")
    p.deleted_at = datetime.now()
    db.commit()
    return {"code": 0, "data": None, "msg": "已删除"}


# ============ 产品销售订单 ============

@router.get("/product-sales")
def list_product_sales(db: Session = Depends(get_db), _=Depends(get_current_user),
                       start_date: str = "", end_date: str = "", customer_id: int = 0,
                       page: int = 1, per_page: int = 100):
    qy = db.query(ProductSale).filter(ProductSale.deleted_at.is_(None))
    if customer_id:
        qy = qy.filter(ProductSale.customer_id == customer_id)
    if start_date:
        qy = qy.filter(ProductSale.sale_date >= start_date)
    if end_date:
        qy = qy.filter(ProductSale.sale_date <= end_date)
    total, rows = paginate(qy.order_by(ProductSale.id.desc()), page, per_page)
    prod_map = {p.id: p for p in db.query(Product).filter(Product.deleted_at.is_(None)).all()}
    cust_map = {c.id: c for c in db.query(Customer).filter(Customer.deleted_at.is_(None)).all()}
    data = [_sale_dict(s, prod_map, cust_map) for s in rows]
    return ok_page(data, total)


def _sale_dict(s, prod_map, cust_map):
    p = prod_map.get(s.product_id)
    c = cust_map.get(s.customer_id)
    return {
        "id": s.id, "sale_date": s.sale_date.strftime("%Y-%m-%d") if s.sale_date else "",
        "customer_id": s.customer_id,
        "customer": f"{c.code} {c.name}" if c else "",
        "product_id": s.product_id,
        "product": p.name if p else "",
        "region": s.region or "",
        "qty": s.qty,
        "purchase_price": float(s.purchase_price or 0),
        "sales_price": float(s.sales_price or 0),
        "sales_amount": float(s.sales_amount or 0),
        "profit": float(s.profit or 0),
        "note": s.note or "",
        "created_at": fmt_dt(s.created_at), "created_by": s.created_by or "",
        "updated_at": fmt_dt(s.updated_at), "updated_by": s.updated_by or "",
    }


def _calc(qty, purchase_price, sales_price):
    qty = int(qty or 0)
    pp = Decimal(str(purchase_price or 0))
    sp = Decimal(str(sales_price or 0))
    sales_amount = round(Decimal(qty) * sp, 2)
    profit = round(Decimal(qty) * (sp - pp), 2)
    return qty, float(pp), float(sp), float(sales_amount), float(profit)


def _sync_customer_bill(db, customer_id, sale_date, user_label=""):
    """产品数据变化后，重新生成该客户当日账单并同步余额"""
    from ..routers.data import _regen_downstream_bills
    if not customer_id or not sale_date:
        return
    _regen_downstream_bills(db, sale_date, [customer_id], user_label)


@router.post("/product-sales")
def create_product_sale(body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    sale_date = (body.get("sale_date") or "").strip()
    if not sale_date:
        raise HTTPException(status_code=422, detail="请选择日期")
    if not body.get("customer_id"):
        raise HTTPException(status_code=422, detail="请选择客户")
    if not body.get("product_id"):
        raise HTTPException(status_code=422, detail="请选择产品")
    region = (body.get("region") or "").strip()
    if not region:
        raise HTTPException(status_code=422, detail="请选择地区")
    if region not in ALL_REGIONS:
        raise HTTPException(status_code=422, detail=f"地区不存在：{region}")
    qty, pp, sp, sales_amount, profit = _calc(body.get("qty"),
                                              body.get("purchase_price"),
                                              body.get("sales_price"))
    s = ProductSale(sale_date=_date(sale_date), customer_id=body.get("customer_id"),
                    product_id=body.get("product_id"), region=region, qty=qty,
                    purchase_price=pp, sales_price=sp,
                    sales_amount=sales_amount, profit=profit,
                    note=(body.get("note") or "").strip())
    s.created_by = _ulabel(user)
    s.updated_by = _ulabel(user)
    db.add(s)
    db.commit()
    _sync_customer_bill(db, body.get("customer_id"), _date(sale_date), _ulabel(user))
    return {"code": 0, "data": {"id": s.id}, "msg": "产品销售已录入"}


@router.put("/product-sales/{sid}")
def update_product_sale(sid: int, body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    s = db.query(ProductSale).filter(ProductSale.id == sid).first()
    if not s:
        raise HTTPException(status_code=404, detail="产品销售记录不存在")
    old_date = s.sale_date
    old_customer = s.customer_id
    if body.get("sale_date") is not None:
        s.sale_date = _date(body.get("sale_date"))
    if body.get("customer_id") is not None:
        s.customer_id = body.get("customer_id")
    if body.get("product_id") is not None:
        s.product_id = body.get("product_id")
    if body.get("region") is not None:
        region = (body.get("region") or "").strip()
        if not region:
            raise HTTPException(status_code=422, detail="请选择地区")
        if region not in ALL_REGIONS:
            raise HTTPException(status_code=422, detail=f"地区不存在：{region}")
        s.region = region
    qty = body.get("qty", s.qty)
    pp = body.get("purchase_price", s.purchase_price)
    sp = body.get("sales_price", s.sales_price)
    qty, pp, sp, sales_amount, profit = _calc(qty, pp, sp)
    s.qty, s.purchase_price, s.sales_price = qty, pp, sp
    s.sales_amount, s.profit = sales_amount, profit
    if body.get("note") is not None:
        s.note = (body.get("note") or "").strip()
    s.updated_by = _ulabel(user)
    db.commit()
    # 产品数据变化：更新新日期账单；若日期/客户变化还需更新原日期账单
    _sync_customer_bill(db, s.customer_id, s.sale_date, _ulabel(user))
    if (old_date != s.sale_date) or (old_customer != s.customer_id):
        _sync_customer_bill(db, old_customer, old_date, _ulabel(user))
    return {"code": 0, "data": None, "msg": "产品销售已更新"}


@router.delete("/product-sales/{sid}")
def delete_product_sale(sid: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    s = db.query(ProductSale).filter(ProductSale.id == sid).first()
    if not s:
        raise HTTPException(status_code=404, detail="产品销售记录不存在")
    customer_id = s.customer_id
    sale_date = s.sale_date
    s.deleted_at = datetime.now()
    db.commit()
    _sync_customer_bill(db, customer_id, sale_date, _ulabel(user))
    return {"code": 0, "data": None, "msg": "已删除"}


@router.post("/product-sales/save-batch")
def save_product_sales_batch(body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """批量保存产品销售（按 日期+客户+产品+地区 去重 upsert），并重新生成受影响客户当日账单数据"""
    date_str = (body.get("date") or "").strip()
    rows = body.get("rows") or []
    if not date_str or not rows:
        raise HTTPException(status_code=422, detail="请选择日期并填写数据")
    sale_date = _date(date_str)
    if not sale_date:
        raise HTTPException(status_code=422, detail="日期格式错误")

    affected_customers = set()
    for r in rows:
        region = (r.get("region") or "").strip()
        if not region:
            raise HTTPException(status_code=422, detail="请选择地区")
        if region not in ALL_REGIONS:
            raise HTTPException(status_code=422, detail=f"地区不存在：{region}")
        customer_id = r.get("customer_id")
        product_id = r.get("product_id")
        if not customer_id or not product_id:
            raise HTTPException(status_code=422, detail="请选择客户和产品")
        qty, pp, sp, sales_amount, profit = _calc(r.get("qty"),
                                                  r.get("purchase_price"),
                                                  r.get("sales_price"))
        sid = r.get("id")
        s = None
        if sid:
            s = db.query(ProductSale).filter(ProductSale.id == sid).first()
        else:
            s = db.query(ProductSale).filter(ProductSale.sale_date == sale_date,
                                             ProductSale.customer_id == customer_id,
                                             ProductSale.product_id == product_id,
                                             ProductSale.region == region).first()
        if s:
            s.customer_id = customer_id
            s.product_id = product_id
            s.region = region
            s.qty, s.purchase_price, s.sales_price = qty, pp, sp
            s.sales_amount, s.profit = sales_amount, profit
            if r.get("note") is not None:
                s.note = (r.get("note") or "").strip()
            s.updated_by = _ulabel(user)
        else:
            s = ProductSale(sale_date=sale_date, customer_id=customer_id,
                            product_id=product_id, region=region, qty=qty,
                            purchase_price=pp, sales_price=sp,
                            sales_amount=sales_amount, profit=profit,
                            note=(r.get("note") or "").strip())
            s.created_by = _ulabel(user)
            s.updated_by = _ulabel(user)
            db.add(s)
        affected_customers.add(customer_id)
    db.commit()

    # 生成账单数据：重新生成受影响客户当日的下游账单（分数+产品销售合并）并同步余额
    if affected_customers:
        from ..routers.data import _regen_downstream_bills
        _regen_downstream_bills(db, sale_date, affected_customers, _ulabel(user))
    return {"code": 0, "data": {"customers": sorted(affected_customers)}, "msg": f"已保存并生成账单数据"}


@router.post("/product-sales/import")
async def import_product_sales(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    """Excel 导入产品销售订单（入库并生成账单数据）：表头 日期 客户 地区 产品 进货单价 出货单价 数量"""
    content = await file.read()
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    header, rows = None, []
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = ["" if v is None else str(v).strip() for v in values]
            continue
        rows.append(values)
    wb.close()

    def _norm(h):
        return re.sub(r"\s+", "", str(h))

    idx = {_norm(h): i for i, h in enumerate(header)}
    date_col = idx.get(_norm("日期"))
    cust_col = idx.get(_norm("客户"))
    region_col = idx.get(_norm("地区"))
    prod_col = idx.get(_norm("产品"))
    pp_col = idx.get(_norm("进货单价"))
    sp_col = idx.get(_norm("出货单价"))
    qty_col = idx.get(_norm("数量"))

    if date_col is None or cust_col is None or prod_col is None or region_col is None:
        raise HTTPException(status_code=422, detail="缺少必需表头：日期、客户、地区、产品")

    customers = db.query(Customer).filter(Customer.ctype == "downstream").all()
    cust_by_code = {c.code: c for c in customers if c.code}
    products = db.query(Product).all()
    prod_by_name = {p.name: p for p in products if p.name}

    errors = []
    imported = 0
    updated = 0
    affected = set()  # (customer_id, sale_date)
    for values in rows:
        def getv(i):
            return values[i] if i is not None and i < len(values) else None

        date_str = "" if date_col is None else str(getv(date_col) or "").strip()
        if date_str == "":
            continue  # 日期列为空的行直接跳过，不检查不导入
        cust_str = "" if cust_col is None else str(getv(cust_col) or "").strip()
        prod_str = "" if prod_col is None else str(getv(prod_col) or "").strip()
        region_str = "" if region_col is None else str(getv(region_col) or "").strip()

        sale_date = _date(date_str)
        if not sale_date:
            errors.append(f"日期无效：{date_str}")
            continue
        cust = cust_by_code.get(cust_str)
        if not cust:
            errors.append(f"客户编号不存在：{cust_str}")
            continue
        if not region_str:
            errors.append(f"地区为空")
            continue
        if region_str not in ALL_REGIONS:
            errors.append(f"地区不存在：{region_str}")
            continue
        prod = prod_by_name.get(prod_str)
        if not prod:
            errors.append(f"产品不存在：{prod_str}")
            continue

        qty_raw = getv(qty_col)
        qty_str = "" if qty_raw is None else str(qty_raw).strip()
        qty_num = _parse_num(qty_raw)
        if qty_num is None or qty_num <= 0 or qty_num != int(qty_num):
            errors.append(f"数量无效：{qty_str}")
            continue
        qty = int(qty_num)

        pp_raw = getv(pp_col)
        pp_str = "" if pp_raw is None else str(pp_raw).strip()
        if pp_str == "":
            pp = 0.0
        else:
            pp_num = _parse_num(pp_raw)
            if pp_num is None or pp_num < 0:
                errors.append(f"进货单价无效：{pp_str}")
                continue
            pp = pp_num

        sp_raw = getv(sp_col)
        sp_str = "" if sp_raw is None else str(sp_raw).strip()
        sp_num = _parse_num(sp_raw)
        if sp_str == "" or sp_num is None or sp_num <= 0:
            errors.append(f"出货单价无效：{sp_str}")
            continue
        sp = sp_num

        sales_amount = round(qty * sp, 2)
        profit = round(qty * (sp - pp), 2)
        existing = db.query(ProductSale).filter(ProductSale.sale_date == sale_date,
                                                 ProductSale.customer_id == cust.id,
                                                 ProductSale.product_id == prod.id,
                                                 ProductSale.region == region_str).first()
        if existing:
            existing.qty = qty
            existing.purchase_price = pp
            existing.sales_price = sp
            existing.sales_amount = sales_amount
            existing.profit = profit
            existing.region = region_str
            existing.updated_by = _ulabel(user)
            updated += 1
        else:
            s = ProductSale(sale_date=sale_date, customer_id=cust.id, product_id=prod.id,
                            region=region_str, qty=qty, purchase_price=pp, sales_price=sp,
                            sales_amount=sales_amount, profit=profit)
            s.created_by = _ulabel(user)
            s.updated_by = _ulabel(user)
            db.add(s)
            imported += 1
        affected.add((cust.id, sale_date))
    db.commit()

    # 生成账单数据：按日期分组重新生成受影响客户账单
    from collections import defaultdict
    by_date = defaultdict(set)
    for cid, d in affected:
        by_date[d].add(cid)
    for d, cids in by_date.items():
        from ..routers.data import _regen_downstream_bills
        _regen_downstream_bills(db, d, cids, _ulabel(user))

    msg = f"导入 {imported} 条，更新 {updated} 条"
    if errors:
        msg += f"，失败 {len(errors)} 条：{'；'.join(errors[:10])}"
    return {"code": 0, "data": {"imported": imported, "updated": updated, "errors": errors,
                                "dates": sorted([d.strftime("%Y-%m-%d") for d in by_date.keys()])}, "msg": msg}


@router.get("/product-sales/export")
def export_product_sales(db: Session = Depends(get_db), _=Depends(get_current_user), date: str = ""):
    """导出产品销售为 Excel（基于导入模版 MB-PS-IMPORT 文件，保留其格式，可再导入）：传 date 只导该日期，不传导全部"""
    qy = db.query(ProductSale)
    if date:
        try:
            d = datetime.strptime(date, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(status_code=422, detail="日期格式错误")
        qy = qy.filter(ProductSale.sale_date == d)
    rows = qy.order_by(ProductSale.id).all()
    import openpyxl
    tpl = ensure_product_sale_template(db)
    data = tpl.file_data if tpl.file_data else _build_product_sale_template_xlsx()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)  # 删除示例行，保留表头与格式
    prod_map = {p.id: p.name for p in db.query(Product).all()}
    cust_code = {c.id: c.code for c in db.query(Customer).all()}
    for s in rows:
        ws.append([s.sale_date.strftime("%Y-%m-%d") if s.sale_date else "",
                   cust_code.get(s.customer_id, ""), s.region or "",
                   prod_map.get(s.product_id, ""),
                   float(s.purchase_price or 0), float(s.sales_price or 0), s.qty])
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"产品销售{('_' + date) if date else ''}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


PS_IMPORT_TEMPLATE_CODE = "MB-PS-IMPORT"


def _build_product_sale_template_xlsx():
    """生成产品销售导入模版（表头 + 示例行）"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "产品销售导入"
    ws.append(["日期", "客户", "地区", "产品", "进货单价", "出货单价", "数量"])
    ws.append(["2026-01-01", "660001", "浙江", "产品名称", 1.00, 2.00, 10])
    for col in ("A", "B", "C", "D", "E", "F", "G"):
        ws.column_dimensions[col].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ensure_product_sale_template(db):
    """确保「产品销售导入模版」存在于模版中心（幂等），返回模版记录"""
    t = db.query(Template).filter(Template.code == PS_IMPORT_TEMPLATE_CODE).first()
    if not t:
        t = Template(ttype="产品销售导入模版", tpl_type="其他", code=PS_IMPORT_TEMPLATE_CODE,
                     name="产品销售导入模版.xlsx", file_data=_build_product_sale_template_xlsx(),
                     file_type="xlsx",
                     description="产品销售订单导入模版：表头为 日期/客户/地区/产品/进货单价/出货单价/数量。"
                                 "日期支持 YYYY-MM-DD、YYYY/MM/DD、YYYYMMDD、YYMMDD、MMDD；"
                                 "客户填客户管理中的客户编号；地区填省份或地市名称（需在地区主数据中存在）；"
                                 "产品填产品管理中的产品名称。",
                     status=1)
        db.add(t)
        db.commit()
    else:
        # 模版为自动生成的简单表头，若缺少地区列则同步为最新列
        needs_update = t.file_data is None
        if not needs_update:
            import openpyxl
            try:
                wb = openpyxl.load_workbook(io.BytesIO(t.file_data), read_only=True)
                hdr = [c.value for c in wb.active[1]]
                needs_update = "地区" not in hdr
                wb.close()
            except Exception:
                needs_update = True
        if needs_update:
            t.file_data = _build_product_sale_template_xlsx()
            t.description = "产品销售订单导入模版：表头为 日期/客户/地区/产品/进货单价/出货单价/数量。" \
                            "日期支持 YYYY-MM-DD、YYYY/MM/DD、YYYYMMDD、YYMMDD、MMDD；" \
                            "客户填客户管理中的客户编号；地区填省份或地市名称（需在地区主数据中存在）；" \
                            "产品填产品管理中的产品名称。"
            db.commit()
    return t


@router.get("/product-sales/template")
def export_product_sale_template(db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = ensure_product_sale_template(db)
    data = t.file_data if t.file_data else _build_product_sale_template_xlsx()
    filename = t.name or "产品销售导入模版.xlsx"
    if "." not in filename:
        filename = f"{filename}.xlsx"
    return Response(content=bytes(data),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


def _num(v):
    try:
        return float(v) if v not in (None, "") else 0
    except (TypeError, ValueError):
        return 0


def _parse_num(v):
    """解析数字：支持千分位逗号/货币符号/空格，失败返回 None（用于区分 0 与非法值）"""
    if v is None:
        return None
    s = str(v).strip().replace(",", "").replace("，", "")
    for ch in ("￥", "¥", "元"):
        s = s.replace(ch, "")
    s = s.strip()
    if s == "":
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _date(v):
    """解析日期：支持 YYYY-MM-DD / YYYY/MM/DD / YYYYMMDD / YYMMDD / MMDD / datetime 对象"""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    digits = s.replace(".", "")
    if digits.isdigit():
        n = len(digits)
        if n == 4:  # MMDD → 用当前年
            return _safe_date(date.today().year, int(digits[:2]), int(digits[2:4]))
        if n == 6:  # YYMMDD → 2000 + YY
            return _safe_date(2000 + int(digits[:2]), int(digits[2:4]), int(digits[4:6]))
        if n == 8:  # YYYYMMDD
            return _safe_date(int(digits[:4]), int(digits[4:6]), int(digits[6:8]))
    s2 = s.replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    try:
        return datetime.strptime(s2, "%Y-%m-%d").date()
    except ValueError:
        return None


def _safe_date(y, m, d):
    try:
        return date(y, m, d)
    except ValueError:
        return None
