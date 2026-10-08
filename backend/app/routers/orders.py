"""订单路由"""
import csv
import io
import json
import re
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session, aliased

from ..database import get_db
from ..deps import get_current_user, require_admin
from ..models import Order, OrderUrl, Channel, Customer, Alert, OrderTemplate, Template, TidabiaoHistory, OperationLog, Operator, Platform, Category
from ..schemas import OrderBody, StopBody, BatchStopBody, OrderDistConfigBody, GenerateBody, BatchGroupBody
from ..pagination import paginate, ok_page
from ..utils import fmt_dt
from ..tidabiao import _order_type, _region_mode

router = APIRouter(prefix="/api/orders", tags=["orders"])

ORDER_HEADER_MAP = {
    "甲方": "upstream", "上游": "upstream", "下游": "customer", "一级代理": "customer", "客户": "customer",
    "二级代理": "secondary_agent",
    "提单日": "order_date", "更新日期": "order_date", "日期": "order_date",
    "订单号": "order_no", "订单编号": "order_no", "编号": "order_no",
    "定价": "price", "单价": "price",
    "分组": "group_name", "小组": "group_name", "小组名": "group_name",
    "模版编号": "tpl_code", "出数模版": "tpl_code",
    "导出文件名": "export_filename", "出数excel命名": "export_filename",
    "是否加名": "add_name",
    "是否撞库": "check_collision",
    "状态": "status",
    "平台": "platform",
    "渠道": "channel", "类型": "channel",
    "运营商": "operator",
    "任务名": "task_name", "任务ID": "task_id", "工单号": "task_id", "工单": "task_id", "甲方订单号": "task_id",
    "URL": "url", "url": "url", "106拓展码": "url", "拓展码": "url",
    "数量": "qty", "提取数量": "qty",
    "时长": "duration",
    "省份": "province", "地市": "city", "排除省": "excl_province", "排除市": "excl_city",
    "排除省份": "excl_province", "排除地市": "excl_city", "排除地": "excl_city",
    "年龄上限": "age_max", "年龄下限": "age_min", "PV": "pv",
    "开始日期": "start_date", "截止日期": "end_date", "结束日期": "end_date",
    "开始时间": "start_date", "截止时间": "end_date",
}

MB019_HEADERS = ["开始日期", "截止日期", "更新日期", "状态", "甲方", "订单号", "甲方订单号",
                 "一级代理", "二级代理", "定价", "平台", "小组", "出数模版", "是否加名", "是否撞库",
                 "渠道", "运营商", "任务名", "url", "数量", "时长",
                 "年龄\n下限", "年龄\n上限", "pv", "省份", "排除省份", "地市", "排除地市"]


def match_template(db: Session, o) -> int:
    """根据订单属性自动匹配下发模版，返回 template_id 或 None"""
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first() if o.upstream_id else None
    party = "牛" if (up and up.name == "牛") else "新"
    typ = _order_type((o.task_id or "") + " " + (o.task_name or ""))
    region = _region_mode(o)
    tmpls = db.query(OrderTemplate).all()
    best = None
    best_score = -1
    for t in tmpls:
        try:
            mr = json.loads(t.match_rule_json or "{}")
        except Exception:
            mr = {}
        score = 0
        if mr.get("party"):
            score += 1 if mr["party"] == party else -100
        if mr.get("type"):
            score += 1 if mr["type"] == typ else -100
        if mr.get("region"):
            score += 1 if mr["region"] == region else -100
        if score > best_score:
            best_score = score
            best = t
    if best is None or best_score < 0:
        return None
    return best.id


def _region_summary(o: Order) -> str:
    parts = []
    if o.province:
        parts.append(o.province if o.province != "全国" else "全国")
    if o.city:
        parts.append(f"地市：{o.city}")
    if o.excl_province:
        parts.append(f"排除省：{o.excl_province}")
    if o.excl_city:
        parts.append(f"排除市：{o.excl_city}")
    return "｜".join(parts) if parts else "全国"


def _order_regions(o: Order) -> list:
    """订单地区列表（省份 + 地市，逗号拆分）"""
    regions = []
    for v in (o.province or "").split(","):
        v = v.strip()
        if v:
            regions.append(v)
    for v in (o.city or "").split(","):
        v = v.strip()
        if v:
            regions.append(v)
    return regions


def _operator_name(db, o: Order) -> str:
    if o.operator_id:
        op = db.query(Operator).filter(Operator.id == o.operator_id).first()
        if op:
            return op.name
    return ""


def _order_dict(db, o: Order, alert_tasks=None) -> dict:
    ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
    c = db.query(Customer).filter(Customer.id == o.customer_id).first()
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first()
    tpl = db.query(Template).filter(Template.id == o.tpl_id).first() if o.tpl_id else None
    operator = _operator_name(db, o)
    order_urls = list(o.urls)
    urls = [u.url for u in order_urls]
    url_items = [{"url": u.url} for u in order_urls]
    return {
        "id": o.id, "order_no": o.order_no, "status": o.status,
        "batch_no": o.batch_no or "",
        "dup_order_nos": o.dup_order_nos or "",
        "order_date": str(o.order_date) if o.order_date else None,
        "start_date": str(o.start_date) if o.start_date else None,
        "end_date": str(o.end_date) if o.end_date else None,
        "updated_at": fmt_dt(o.updated_at),
        "upstream": up.name if up else "", "upstream_id": o.upstream_id,
        "customer": f"{c.code} {c.name}" if c else "", "customer_id": o.customer_id,
        "channel": ch.name if ch else "", "channel_id": o.channel_id, "task_name": o.task_name,
        "task_id": o.task_id or "", "client_order_no": o.task_id or "", "duration": _duration_days(o.start_date, o.end_date) or o.duration or "",
        "operator": operator, "operator_id": o.operator_id,
        "price": float(o.price) if o.price is not None else None,
        "secondary_agent": o.secondary_agent or "",
        "platform": o.platform or "",
        "template_id": o.template_id, "filename_rule": o.filename_rule or "",
        "template_name": (db.query(OrderTemplate).filter(OrderTemplate.id == o.template_id).first().name
                          if o.template_id else ""),
        "tpl_id": o.tpl_id, "tpl_code": tpl.code if tpl else "", "tpl_name": tpl.name if tpl else "",
        "group_name": o.group_name or "",
        "export_filename": o.export_filename or "",
        "add_name": bool(o.add_name),
        "check_collision": bool(o.check_collision),
        "first_output_date": str(o.start_date + timedelta(days=2)) if o.start_date else None,
        "last_output_date": str(o.end_date + timedelta(days=1)) if o.end_date else None,
        "dist_config": json.loads(o.dist_config_json) if o.dist_config_json else None,
        "url": urls[0] + f" 等{len(urls)}个" if len(urls) > 1 else (urls[0] if urls else ""),
        "urls": urls, "url_items": url_items, "region": _region_summary(o),
        "province": o.province, "city": o.city, "excl_province": o.excl_province, "excl_city": o.excl_city,
        "qty": o.qty, "age_min": o.age_min, "age_max": o.age_max, "pv": o.pv,
        "created_at": fmt_dt(o.created_at), "stop_date": str(o.stop_date) if o.stop_date else None,
        "has_alert": (o.task_name in alert_tasks) if alert_tasks is not None else False,
    }


def _op_log(db, user, action, target, before, after):
    """写操作日志（不单独 commit，随业务事务一起提交）"""
    db.add(OperationLog(
        user_id=user.id if user else None,
        username=user.username if user else None,
        module="订单", action=action, target=target,
        detail=json.dumps({"before": before, "after": after}, ensure_ascii=False, default=str),
        created_at=datetime.now(),
    ))


def _snapshot(db, o: Order, urls=None) -> dict:
    """订单快照（用于操作前后对比）"""
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first()
    cust = db.query(Customer).filter(Customer.id == o.customer_id).first()
    ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
    tpl = db.query(Template).filter(Template.id == o.tpl_id).first() if o.tpl_id else None
    return {
        "order_no": o.order_no, "task_name": o.task_name, "task_id": o.task_id or "",
        "status": o.status, "order_date": str(o.order_date) if o.order_date else None,
        "upstream": up.name if up else "", "customer": (f"{cust.code} {cust.name}" if cust else ""),
        "channel": ch.name if ch else "", "operator": _operator_name(db, o), "tpl_code": tpl.code if tpl else "",
        "qty": o.qty, "duration": o.duration, "province": o.province, "city": o.city,
        "excl_province": o.excl_province, "excl_city": o.excl_city,
        "age_min": o.age_min, "age_max": o.age_max, "pv": o.pv,
        "start_date": str(o.start_date) if o.start_date else None,
        "end_date": str(o.end_date) if o.end_date else None,
        "price": float(o.price) if o.price is not None else None,
        "secondary_agent": o.secondary_agent, "platform": o.platform,
        "group_name": o.group_name, "add_name": bool(o.add_name),
        "check_collision": bool(o.check_collision),
        "urls": urls if urls is not None else [u.url for u in o.urls],
    }


def _norm(v):
    return None if v in (None, "") else v


def _validate_region(body: OrderBody):
    provs = [p for p in (body.province or "").split(",") if p.strip()]
    if "全国" in provs and body.city:
        raise HTTPException(status_code=422, detail="省份为全国时不可选择地市")
    if body.city and (body.excl_province or body.excl_city):
        raise HTTPException(status_code=422, detail="选择地市时不可设置排除省/排除市")
    if body.excl_province and "全国" not in provs:
        raise HTTPException(status_code=422, detail="仅省份=全国时可设置排除省")


@router.get("")
def list_orders(db: Session = Depends(get_db), _=Depends(get_current_user),
                status: str = "", q: str = "", order_no: str = "", upstream: str = "",
                customer: str = "", channel: str = "", task_name: str = "", province: str = "",
                city: str = "", excl_province: str = "", excl_city: str = "", qty: str = "",
                age_min: str = "", age_max: str = "", pv: str = "",
                start_date: str = "", end_date: str = "", order_date: str = "",
                platform: str = "", secondary_agent: str = "",
                task_id: str = "", duration: str = "", group_name: str = "", tpl_code: str = "",
                add_name: str = "", url: str = "", operator: str = "", price: str = "",
                dup: str = "", dup_order_no: str = "",
                check_collision: str = "", no_cat2: str = "",
                page: int = 1, per_page: int = 10):
    Upstream = aliased(Customer)
    qy = (db.query(Order)
          .outerjoin(Customer, Order.customer_id == Customer.id)
          .outerjoin(Upstream, Order.upstream_id == Upstream.id)
          .outerjoin(Channel, Order.channel_id == Channel.id)
          .outerjoin(Operator, Order.operator_id == Operator.id)
          .outerjoin(Template, Order.tpl_id == Template.id))
    if status:
        qy = qy.filter(Order.status == status)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Order.order_no.like(like), Order.task_name.like(like),
                           Customer.name.like(like), Upstream.name.like(like),
                           Channel.name.like(like)))
    if order_no:
        qy = qy.filter(Order.order_no.like(f"%{order_no}%"))
    if upstream:
        qy = qy.filter(Upstream.name.like(f"%{upstream}%"))
    if customer:
        qy = qy.filter(Customer.code.like(f"%{customer}%"))
    if channel:
        qy = qy.filter(Channel.name.like(f"%{channel}%"))
    if task_name:
        qy = qy.filter(Order.task_name.like(f"%{task_name}%"))
    if task_id:
        qy = qy.filter(or_(Order.task_id.like(f"%{task_id}%"), Order.order_no.like(f"%{task_id}%")))
    if duration:
        qy = qy.filter(Order.duration == duration)
    if group_name:
        qy = qy.filter(Order.group_name.like(f"%{group_name}%"))
    if tpl_code:
        qy = qy.filter(Template.code.like(f"%{tpl_code}%"))
    if add_name:
        qy = qy.filter(Order.add_name == (add_name in ("是", "true", "1", "True", "yes", "YES")))
    if check_collision:
        qy = qy.filter(Order.check_collision == (check_collision in ("是", "true", "1", "True", "yes", "YES")))
    if url:
        url_sub = db.query(OrderUrl.order_id).filter(OrderUrl.url.like(f"%{url}%")).subquery()
        qy = qy.filter(Order.id.in_(url_sub))
    if operator:
        qy = qy.filter(Operator.name.like(f"%{operator}%"))
    if dup in ("有", "是", "true", "1", "True", "yes", "YES"):
        qy = qy.filter(Order.dup_order_nos.isnot(None), Order.dup_order_nos != "")
    elif dup in ("无", "否", "false", "0", "False", "no", "NO"):
        qy = qy.filter(or_(Order.dup_order_nos.is_(None), Order.dup_order_nos == ""))
    if dup_order_no:
        qy = qy.filter(Order.dup_order_nos.like(f"%{dup_order_no}%"))
    if price:
        qy = qy.filter(Order.price == _float(price))
    if province:
        provs = [p for p in province.split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[Order.province.like(f"%{p}%") for p in provs]))
    if city:
        cities = [c for c in city.split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[Order.city.like(f"%{c}%") for c in cities]))
    if excl_province:
        qy = qy.filter(Order.excl_province.like(f"%{excl_province}%"))
    if excl_city:
        qy = qy.filter(Order.excl_city.like(f"%{excl_city}%"))
    if qty:
        qy = qy.filter(Order.qty == _num(qty))
    if age_min:
        qy = qy.filter(Order.age_min == _num(age_min))
    if age_max:
        qy = qy.filter(Order.age_max == _num(age_max))
    if pv:
        qy = qy.filter(Order.pv == _num(pv))
    if start_date:
        qy = qy.filter(Order.start_date >= start_date)
    if end_date:
        qy = qy.filter(Order.end_date <= end_date)
    if order_date:
        qy = qy.filter(Order.order_date == order_date)
    if platform:
        qy = qy.filter(Order.platform.like(f"%{platform}%"))
    if no_cat2:
        Cat2 = aliased(Category)
        cat2_platforms = db.query(Platform.name).join(Cat2, Cat2.id == Platform.cat_id).subquery()
        if no_cat2 in ("无", "否", "false", "0", "False", "no", "NO"):
            qy = qy.filter(or_(Order.platform.is_(None), Order.platform == "",
                               ~Order.platform.in_(cat2_platforms)))
        elif no_cat2 in ("有", "是", "true", "1", "True", "yes", "YES"):
            qy = qy.filter(Order.platform.in_(cat2_platforms))
    if secondary_agent:
        qy = qy.filter(Order.secondary_agent.like(f"%{secondary_agent}%"))
    total, rows = paginate(qy.order_by(Order.id.desc()), page, per_page)
    alert_tasks = {r[0] for r in db.query(Alert.task_name).filter(Alert.type == "停单提醒", Alert.status == "未处理").all() if r[0]}
    return ok_page([_order_dict(db, o, alert_tasks) for o in rows], total)


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _next_order_no(db, order_date=None):
    """订单编号 = LM- + 更新日期(年月日) + 4 位顺序号"""
    d = order_date or date.today()
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, str):
        s = d.strip()
        try:
            d = datetime.strptime(s, "%Y-%m-%d").date() if s else date.today()
        except ValueError:
            d = date.today()
    prefix = f"LM-{d.strftime('%Y%m%d')}"
    max_seq = 0
    for (no,) in db.query(Order.order_no).filter(Order.order_no.like(prefix + "%")).all():
        m = re.match(r"^LM-\d{8}(\d{4})$", no or "")
        if m:
            max_seq = max(max_seq, int(m.group(1)))
    return f"{prefix}{max_seq + 1:04d}"


def _gen_new_task_id(db, order_date, channel_name, operator_name, province, city, excl_city):
    """甲方=新 时自动生成工单号：147-{前缀}-{mmdd}-LM{mmdd}{四位顺序号}
    前缀规则：106→kz，小程序→xcx，dpi(省份/全国)→yd/ltd/dxdpi，dpi(地市/排除地市)→移动/联通/电信dpi
    """
    d = order_date or date.today()
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, str):
        try:
            d = datetime.strptime(d.strip(), "%Y-%m-%d").date()
        except ValueError:
            d = date.today()
    mmdd = d.strftime("%m%d")
    max_seq = 0
    for (tid,) in db.query(Order.task_id).filter(Order.task_id.like(f"%LM{mmdd}%")).all():
        m = re.match(rf".*LM{mmdd}(\d{{4}})$", tid or "")
        if m:
            max_seq = max(max_seq, int(m.group(1)))
    seq = f"{max_seq + 1:04d}"
    ch = (channel_name or "").strip()
    op = (operator_name or "").strip()
    if ch == "106":
        mid = "kz"
    elif ch == "小程序":
        mid = "xcx"
    elif "dpi" in ch:
        city_mode = bool((city or "").strip()) or bool((excl_city or "").strip())
        if city_mode:
            op_map = {"移动": "移动dpi", "联通": "联通dpi", "电信": "电信dpi"}
        else:
            op_map = {"移动": "yddpi", "联通": "ltdpi", "电信": "dxdpi"}
        mid = op_map.get(op) or (op + "dpi" if op else "dpi")
    else:
        return None
    return f"147-{mid}-{mmdd}-LM{mmdd}{seq}"


ACTIVE_STATUSES = ("未提", "在执", "改单")


def _find_dup_task(db, upstream_id, order_date, task_name, status="未提", exclude_id=None, operator_id=None):
    """唯一性校验：
    1. 同一(甲方+更新日期+任务名+运营商) 不可重复；
    2. 同一(甲方+任务名+运营商) 只能有一个未停订单（未提/在执/改单，已停后才可新建）
    """
    if not task_name:
        return None
    q = db.query(Order).filter(Order.upstream_id == upstream_id,
                               Order.order_date == order_date,
                               Order.task_name == task_name,
                               Order.operator_id == operator_id)
    if exclude_id is not None:
        q = q.filter(Order.id != exclude_id)
    dup = q.first()
    if dup:
        return dup
    if status in ACTIVE_STATUSES:
        q2 = db.query(Order).filter(Order.upstream_id == upstream_id,
                                    Order.task_name == task_name,
                                    Order.operator_id == operator_id,
                                    Order.status.in_(ACTIVE_STATUSES))
        if exclude_id is not None:
            q2 = q2.filter(Order.id != exclude_id)
        return q2.first()
    return None


@router.post("")
def create_order(body: OrderBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    _validate_region(body)
    missing = []
    if not body.order_date:
        missing.append("更新日期")
    if not (body.task_name or "").strip():
        missing.append("任务名")
    if body.qty is None:
        missing.append("数量")
    if not body.start_date:
        missing.append("开始日期")
    if not body.end_date:
        missing.append("截止日期")
    if not body.upstream_id:
        missing.append("上游")
    if not body.customer_id:
        missing.append("一级代理")
    if missing:
        raise HTTPException(status_code=422, detail=f"缺少必填字段：{'、'.join(missing)}")
    # 甲方=新 时自动生成工单号（牛则保持原样）
    up = db.query(Customer).filter(Customer.id == body.upstream_id).first()
    task_id = body.task_id or ""
    if up and up.name == "新":
        ch = db.query(Channel).filter(Channel.id == body.channel_id).first()
        op = db.query(Operator).filter(Operator.id == body.operator_id).first()
        gen = _gen_new_task_id(db, body.order_date, ch.name if ch else "", op.name if op else "",
                               body.province, body.city, body.excl_city)
        if gen:
            task_id = gen
    warning = None
    dup = _find_dup_task(db, body.upstream_id, body.order_date, body.task_name, status="未提", operator_id=body.operator_id)
    if dup:
        warning = f"提示：工单号「{task_id or '—'}」与 工单号「{dup.task_id or '—'}」重复"
    province = body.province or ("全国" if not body.city else "")
    order_no = _next_order_no(db, body.order_date)
    o = Order(order_no=order_no, customer_id=body.customer_id, upstream_id=body.upstream_id,
              channel_id=body.channel_id, operator_id=body.operator_id, task_name=body.task_name, task_id=task_id,
              qty=body.qty, duration=body.duration,
              province=province, city=body.city, excl_province=body.excl_province,
              excl_city=body.excl_city, age_min=body.age_min, age_max=body.age_max, pv=body.pv,
              start_date=body.start_date, end_date=body.end_date, status="未提",
              filename_rule=body.filename_rule,
              order_date=body.order_date, price=body.price,
              secondary_agent=body.secondary_agent, platform=body.platform, tpl_id=body.tpl_id,
               group_name=body.group_name, export_filename=body.export_filename, add_name=body.add_name,
               check_collision=body.check_collision,
               dist_config_json=json.dumps(body.dist_config, ensure_ascii=False) if body.dist_config else None)
    o.template_id = body.template_id if body.template_id is not None else match_template(db, o)
    db.add(o)
    db.flush()
    url_list = []
    for i, u in enumerate(body.urls):
        url = (u.get("url") or "").strip()
        url_list.append(url)
        db.add(OrderUrl(order_id=o.id, url=url, level=u.get("level", "中"), sort_no=i))
    _op_log(db, user, "create_order", order_no, None, _snapshot(db, o, urls=url_list))
    db.commit()
    return {"code": 0, "data": {"id": o.id, "order_no": order_no}, "msg": "已保存（状态：未提）", "warning": warning}


@router.put("/{oid}")
def update_order(oid: int, body: OrderBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    o = db.query(Order).filter(Order.id == oid).first()
    if not o:
        raise HTTPException(status_code=404, detail="订单不存在")
    before = _snapshot(db, o)
    _validate_region(body)
    is_edit = o.status == "未提"

    changed_fields = []
    if is_edit:
        new_status = "未提"
    else:
        # 在执/已停/改单：渠道/类型/运营商不允许修改
        if _norm(body.channel_id) != _norm(o.channel_id) or _norm(body.operator_id) != _norm(o.operator_id):
            raise HTTPException(status_code=400, detail="在执/已停订单不允许修改渠道/类型/运营商")
        # 只有 Url/数量/地区/年龄上下限/pv 变化才算改单
        before_urls = sorted(before.get("urls") or [])
        new_urls = sorted((u.get("url") or "").strip() for u in (body.urls or []))
        if before_urls != new_urls:
            changed_fields.append("url")
        for f in ("qty", "province", "city", "excl_province", "excl_city", "age_min", "age_max", "pv"):
            if _norm(getattr(body, f)) != _norm(before.get(f)):
                changed_fields.append(f)
        new_status = "改单" if changed_fields else o.status

    # 甲方=新 且 工单号为空 时自动生成工单号
    up = db.query(Customer).filter(Customer.id == body.upstream_id).first()
    if up and up.name == "新" and not (body.task_id or "").strip():
        ch = db.query(Channel).filter(Channel.id == body.channel_id).first()
        op = db.query(Operator).filter(Operator.id == body.operator_id).first()
        gen = _gen_new_task_id(db, body.order_date, ch.name if ch else "", op.name if op else "",
                               body.province, body.city, body.excl_city)
        if gen:
            body.task_id = gen

    warning = None
    dup = _find_dup_task(db, body.upstream_id, body.order_date, body.task_name, status=new_status, exclude_id=oid, operator_id=body.operator_id)
    if dup:
        warning = f"提示：工单号「{body.task_id or '—'}」与 工单号「{dup.task_id or '—'}」重复"
    body.province = body.province or ("全国" if not body.city else "")
    for k in ["customer_id", "upstream_id", "channel_id", "operator_id", "task_name", "task_id", "qty", "duration",
              "province", "city", "excl_province", "excl_city", "age_min", "age_max", "pv",
              "start_date", "end_date", "filename_rule",
              "order_date", "price", "secondary_agent", "platform", "tpl_id",
              "group_name", "export_filename", "add_name", "check_collision"]:
        setattr(o, k, getattr(body, k))
    o.dist_config_json = json.dumps(body.dist_config, ensure_ascii=False) if body.dist_config else None
    o.template_id = body.template_id if body.template_id is not None else match_template(db, o)
    o.status = new_status
    o.change_fields_json = json.dumps(changed_fields, ensure_ascii=False) if changed_fields else None
    o.stop_date = date.today() if o.status in ("已停", "待停") else o.stop_date
    o.updated_at = datetime.now()
    db.query(OrderUrl).filter(OrderUrl.order_id == oid).delete()
    url_list = []
    for i, u in enumerate(body.urls):
        url = (u.get("url") or "").strip()
        url_list.append(url)
        db.add(OrderUrl(order_id=oid, url=url, level=u.get("level", "中"), sort_no=i))
    _op_log(db, user, "update_order", o.order_no or str(oid), before, _snapshot(db, o, urls=url_list))
    db.commit()
    if is_edit:
        msg = "已保存（状态：未提）"
    elif changed_fields:
        msg = "改单已保存（状态：改单）"
    else:
        msg = "已保存（状态不变）"
    return {"code": 0, "data": None, "msg": msg, "warning": warning}


@router.put("/{oid}/dist-config")
def save_dist_config(oid: int, body: OrderDistConfigBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    o = db.query(Order).filter(Order.id == oid).first()
    if not o:
        raise HTTPException(status_code=404, detail="订单不存在")
    o.template_id = body.template_id
    o.filename_rule = body.filename_rule or ""
    o.dist_config_json = json.dumps(body.dist_config, ensure_ascii=False) if body.dist_config else None
    o.updated_at = datetime.now()
    db.commit()
    return {"code": 0, "data": None, "msg": "下发配置已保存"}


@router.delete("/{oid}")
def delete_order(oid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    o = db.query(Order).filter(Order.id == oid).first()
    if not o:
        raise HTTPException(status_code=404, detail="订单不存在")
    db.query(OrderUrl).filter(OrderUrl.order_id == oid).delete()
    db.delete(o)
    db.commit()
    return {"code": 0, "data": None, "msg": "已删除"}


@router.post("/{oid}/stop")
def stop_order(oid: int, body: StopBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    o = db.query(Order).filter(Order.id == oid).first()
    if not o:
        raise HTTPException(status_code=404, detail="订单不存在")
    if body.order_date:
        if body.order_date > date.today():
            raise HTTPException(status_code=400, detail="更新日期不能是未来日期，请修改更新日期")
        o.order_date = body.order_date
    o.status = "待停"
    o.stop_date = date.today()
    o.updated_at = datetime.now()
    db.query(Alert).filter(Alert.type == "停单提醒", Alert.task_name == o.task_name, Alert.status == "未处理").delete()
    db.commit()
    return {"code": 0, "data": None, "msg": "已提交停单（状态：待停）"}


@router.post("/batch-stop")
def batch_stop(body: BatchStopBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    if body.order_date:
        if body.order_date > date.today():
            raise HTTPException(status_code=400, detail="更新日期不能是未来日期，请修改更新日期")
    rows = db.query(Order).filter(Order.id.in_(body.ids)).all()
    task_names = []
    for o in rows:
        if body.order_date:
            o.order_date = body.order_date
        o.status = "待停"
        o.stop_date = date.today()
        o.updated_at = datetime.now()
        task_names.append(o.task_name)
    if task_names:
        db.query(Alert).filter(Alert.type == "停单提醒", Alert.task_name.in_(task_names), Alert.status == "未处理").delete()
    db.commit()
    return {"code": 0, "data": {"stopped": len(rows)}, "msg": f"已提交停单 {len(rows)} 单（状态：待停）"}


@router.post("/batch-reopen")
def batch_reopen(body: BatchStopBody, db: Session = Depends(get_db), user=Depends(require_admin)):
    """批量复提：已停订单改回未提，并可更新更新日期（仅管理员）"""
    if body.order_date:
        if body.order_date > date.today():
            raise HTTPException(status_code=400, detail="更新日期不能是未来日期，请修改更新日期")
    rows = db.query(Order).filter(Order.id.in_(body.ids)).all()
    order_nos = []
    for o in rows:
        if o.status == "已停":
            if body.order_date:
                o.order_date = body.order_date
            o.status = "未提"
            o.stop_date = None
            o.updated_at = datetime.now()
            order_nos.append(o.order_no)
    if order_nos:
        _op_log(db, user, "batch_reopen_orders", f"批量复提 {len(order_nos)} 单", None, {"order_nos": order_nos})
    db.commit()
    return {"code": 0, "data": {"reopened": len(order_nos)}, "msg": f"已复提 {len(order_nos)} 单（状态：未提）"}


@router.post("/batch-delete")
def batch_delete(body: BatchStopBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(Order).filter(Order.id.in_(body.ids)).all()
    ids = [o.id for o in rows]
    order_nos = [o.order_no for o in rows]
    if ids:
        db.query(OrderUrl).filter(OrderUrl.order_id.in_(ids)).delete(synchronize_session=False)
        for o in rows:
            db.delete(o)
    _op_log(db, user, "batch_delete_orders", f"批量删除 {len(rows)} 单", None, {"order_nos": order_nos})
    db.commit()
    return {"code": 0, "data": {"deleted": len(rows)}, "msg": f"已删除 {len(rows)} 单"}


@router.post("/batch-group")
def batch_group(body: BatchGroupBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    rows = db.query(Order).filter(Order.id.in_(body.ids)).all()
    for o in rows:
        o.group_name = body.group_name or ""
        o.updated_at = datetime.now()
    db.commit()
    return {"code": 0, "data": {"updated": len(rows)}, "msg": f"已修改 {len(rows)} 单小组"}


@router.post("/check-duplicates")
def check_duplicates(db: Session = Depends(get_db), _=Depends(get_current_user)):
    """查重：只验证「在执」订单，运营商 + url集合 + 地区集合 全部一致才算重复"""
    orders = db.query(Order).filter(Order.status == "在执").all()
    for o in orders:
        o.dup_order_nos = None

    def split_regions(province, city):
        rs = []
        for v in re.split(r"[|｜,，;；]", (province or "")):
            v = v.strip()
            if v:
                rs.append(v)
        for v in re.split(r"[|｜,，;；]", (city or "")):
            v = v.strip()
            if v:
                rs.append(v)
        return rs

    sig_map = {}
    order_sig = {}
    for o in orders:
        op = _operator_name(db, o)
        urls = frozenset(u.url.strip() for u in o.urls if (u.url or "").strip())
        regions = frozenset(split_regions(o.province, o.city))
        sig = (op, urls, regions)
        order_sig[o.id] = sig
        sig_map.setdefault(sig, []).append(o.id)

    for o in orders:
        others = [oid for oid in sig_map.get(order_sig[o.id], []) if oid != o.id]
        if others:
            nos = []
            for oid in sorted(others):
                other = db.query(Order).filter(Order.id == oid).first()
                if other:
                    nos.append(other.order_no)
            o.dup_order_nos = "\n".join(nos)
    db.commit()
    dup_count = sum(1 for o in orders if o.dup_order_nos)
    return {"code": 0, "data": {"total": len(orders), "duplicates": dup_count},
            "msg": f"查重完成：共 {len(orders)} 单（在执），{dup_count} 单存在重复"}


def _parse_channel(name):
    n = name or ""
    if "106" in n:
        return "106", ""
    if "小程序" in n:
        return "小程序", ""
    if "直播" in n:
        return "直播间", ""
    typ = "dpi-白" if "白" in n else ("dpi-灰" if "灰" in n else n)
    for isp in ("移动", "联通", "电信"):
        if isp in n:
            return typ, isp
    return typ, ""


def _export_row(db, o):
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first()
    cust = db.query(Customer).filter(Customer.id == o.customer_id).first()
    ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
    tpl = db.query(Template).filter(Template.id == o.tpl_id).first() if o.tpl_id else None
    ch_name = ch.name if ch else ""
    operator = _operator_name(db, o)
    urls = "\n".join(u.url for u in o.urls)
    return [
        o.start_date.strftime("%m%d") if o.start_date else "",
        o.end_date.strftime("%m%d") if o.end_date else "",
        o.order_date.strftime("%m%d") if o.order_date else "",
        o.status or "",
        up.name if up else "",
        o.order_no or "",
        o.task_id or "",
        cust.code if cust else "",
        o.secondary_agent or "",
        (str(o.price) if o.price is not None else ""),
        o.platform or "",
        o.group_name or "",
        tpl.code if tpl else "",
        "是" if o.add_name else "否",
        "是" if o.check_collision else "否",
        ch_name,
        operator,
        o.task_name or "",
        urls,
        (o.qty if o.qty is not None else ""),
        _duration_days(o.start_date, o.end_date) or o.duration or "",
        (o.age_min if o.age_min is not None else ""),
        (o.age_max if o.age_max is not None else ""),
        (o.pv if o.pv is not None else ""),
        o.province or "",
        o.excl_province or "",
        o.city or "",
        o.excl_city or "",
    ]


def _fill_mb019_export(db, orders):
    """用 MB-019 模版文件填充订单数据，保留模版的样式/下拉框/列宽等结构"""
    import openpyxl
    t = db.query(Template).filter(Template.code == "MB-019").first()
    if t and t.file_data:
        wb = openpyxl.load_workbook(io.BytesIO(bytes(t.file_data)))
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "订单表"
        ws.append(MB019_HEADERS)
    row = 2
    for o in orders:
        vals = _export_row(db, o)
        for j, v in enumerate(vals, 1):
            ws.cell(row=row, column=j, value=v)
        row += 1
    # 删除数据之后的空白行（模版里预置的下拉框/格式行），避免导入时被当成空数据行
    if ws.max_row > row - 1:
        ws.delete_rows(row, ws.max_row - row + 1)
    return wb


@router.get("/export-template")
def export_template(db: Session = Depends(get_db), _=Depends(get_current_user)):
    """导出模版：直接下载 MB-019 模版文件"""
    from urllib.parse import quote
    t = db.query(Template).filter(Template.code == "MB-019").first()
    if not t or not t.file_data:
        raise HTTPException(status_code=404, detail="模版 MB-019 不存在")
    filename = t.name or "MB-019.xlsx"
    return Response(content=bytes(t.file_data),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.get("/export")
def export_orders(db: Session = Depends(get_db), _=Depends(get_current_user),
                  status: str = "", q: str = "", order_no: str = "", upstream: str = "",
                  customer: str = "", channel: str = "", task_name: str = "",
                  province: str = "", city: str = "", excl_province: str = "",
                  excl_city: str = "", start_date: str = "", end_date: str = "",
                  order_date: str = "", platform: str = "", secondary_agent: str = "",
                  task_id: str = "", duration: str = "", group_name: str = "", tpl_code: str = "",
                  add_name: str = "", url: str = "", operator: str = "", price: str = "",
                  dup: str = "", dup_order_no: str = "",
                  check_collision: str = "",
                  qty: str = "", age_min: str = "", age_max: str = "", pv: str = ""):
    """按 MB-019 模版结构导出订单，文件名固定 订单表.xlsx"""
    import openpyxl
    from urllib.parse import quote
    Upstream = aliased(Customer)
    qy = (db.query(Order)
          .outerjoin(Customer, Order.customer_id == Customer.id)
          .outerjoin(Upstream, Order.upstream_id == Upstream.id)
          .outerjoin(Channel, Order.channel_id == Channel.id)
          .outerjoin(Operator, Order.operator_id == Operator.id)
          .outerjoin(Template, Order.tpl_id == Template.id))
    if status:
        qy = qy.filter(Order.status == status)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Order.order_no.like(like), Order.task_name.like(like),
                           Customer.name.like(like), Upstream.name.like(like),
                           Channel.name.like(like)))
    if order_no:
        qy = qy.filter(Order.order_no.like(f"%{order_no}%"))
    if upstream:
        qy = qy.filter(Upstream.name.like(f"%{upstream}%"))
    if customer:
        qy = qy.filter(Customer.code.like(f"%{customer}%"))
    if channel:
        qy = qy.filter(Channel.name.like(f"%{channel}%"))
    if task_name:
        qy = qy.filter(Order.task_name.like(f"%{task_name}%"))
    if task_id:
        qy = qy.filter(or_(Order.task_id.like(f"%{task_id}%"), Order.order_no.like(f"%{task_id}%")))
    if duration:
        qy = qy.filter(Order.duration == duration)
    if group_name:
        qy = qy.filter(Order.group_name.like(f"%{group_name}%"))
    if tpl_code:
        qy = qy.filter(Template.code.like(f"%{tpl_code}%"))
    if add_name:
        qy = qy.filter(Order.add_name == (add_name in ("是", "true", "1", "True", "yes", "YES")))
    if check_collision:
        qy = qy.filter(Order.check_collision == (check_collision in ("是", "true", "1", "True", "yes", "YES")))
    if url:
        url_sub = db.query(OrderUrl.order_id).filter(OrderUrl.url.like(f"%{url}%")).subquery()
        qy = qy.filter(Order.id.in_(url_sub))
    if operator:
        qy = qy.filter(Operator.name.like(f"%{operator}%"))
    if dup in ("有", "是", "true", "1", "True", "yes", "YES"):
        qy = qy.filter(Order.dup_order_nos.isnot(None), Order.dup_order_nos != "")
    elif dup in ("无", "否", "false", "0", "False", "no", "NO"):
        qy = qy.filter(or_(Order.dup_order_nos.is_(None), Order.dup_order_nos == ""))
    if dup_order_no:
        qy = qy.filter(Order.dup_order_nos.like(f"%{dup_order_no}%"))
    if price:
        qy = qy.filter(Order.price == _float(price))
    if province:
        provs = [p for p in province.split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[Order.province.like(f"%{p}%") for p in provs]))
    if city:
        cities = [c for c in city.split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[Order.city.like(f"%{c}%") for c in cities]))
    if excl_province:
        qy = qy.filter(Order.excl_province.like(f"%{excl_province}%"))
    if excl_city:
        qy = qy.filter(Order.excl_city.like(f"%{excl_city}%"))
    if start_date:
        qy = qy.filter(Order.start_date >= start_date)
    if end_date:
        qy = qy.filter(Order.end_date <= end_date)
    if order_date:
        qy = qy.filter(Order.order_date == order_date)
    if platform:
        qy = qy.filter(Order.platform.like(f"%{platform}%"))
    if secondary_agent:
        qy = qy.filter(Order.secondary_agent.like(f"%{secondary_agent}%"))
    if qty:
        qy = qy.filter(Order.qty == _num(qty))
    if age_min:
        qy = qy.filter(Order.age_min == _num(age_min))
    if age_max:
        qy = qy.filter(Order.age_max == _num(age_max))
    if pv:
        qy = qy.filter(Order.pv == _num(pv))

    rows = qy.order_by(Order.id.desc()).all()
    wb = _fill_mb019_export(db, rows)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = "订单表.xlsx"
    return Response(
        content=bio.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def _map_import_status(raw_status):
    """把文件里的状态原样映射为系统状态：停相关→已停，其余原样，空→未提"""
    s = (raw_status or "").strip()
    if s in ("已停", "停", "停单", "已停单"):
        return "已停"
    if s in ("未提", "在执", "改单", "待停"):
        return s
    return "未提"


TPL_CODE_ALIAS = {"A": "MB-002", "B": "MB-003", "C": "MB-004", "D": "MB-005", "E": "MB-006", "F": "MB-007"}
VALID_TPL_CODES = {"MB-002", "MB-003", "MB-004", "MB-005", "MB-006", "MB-007"}


def _resolve_tpl_code(raw):
    """出数模版简称 A~F 映射为 MB-002~MB-007"""
    s = (raw or "").strip()
    if not s:
        return ""
    return TPL_CODE_ALIAS.get(s.upper(), s)


@router.post("/import")
async def import_orders(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    fn = file.filename or ""
    content = await file.read()
    if fn.lower().endswith(".csv"):
        rows = _parse_csv(content)
    elif fn.lower().endswith((".xlsx", ".xls")):
        rows = _parse_xlsx(content)
    else:
        raise HTTPException(status_code=422, detail="仅支持 csv / xlsx 文件")

    imported = 0
    updated = 0
    imported_nos = []
    updated_nos = []
    errors = []
    tpl_code_map = {t.id: t.code for t in db.query(Template).all()}
    cust_code_map = {c.id: c.code for c in db.query(Customer).all()}
    group_consistency = {}
    # 预加载已有订单号，导入按订单号匹配（订单号相同才更新，无订单号则插入）
    existing_order_no = {o.order_no: o for o in db.query(Order).all() if o.order_no}
    for idx, r in enumerate(rows):
        if not any((str(v) if v is not None else "").strip() for v in r.values()):
            continue
        try:
            cval = str(r.get("customer") or "").strip()
            cust = db.query(Customer).filter(Customer.ctype == "downstream",
                                             or_(Customer.code == cval, Customer.name == cval)).first() if cval else None
            uval = str(r.get("upstream") or "").strip()
            up = db.query(Customer).filter(Customer.ctype == "upstream",
                                           or_(Customer.code == uval, Customer.name == uval)).first() if uval else None
            ch = db.query(Channel).filter(Channel.name == r.get("channel")).first() if r.get("channel") else None
            op = db.query(Operator).filter(Operator.name == r.get("operator")).first() if r.get("operator") else None
            tpl_code = _resolve_tpl_code(r.get("tpl_code"))
            tpl = db.query(Template).filter(Template.code == tpl_code).first() if tpl_code else None
            task_name = r.get("task_name")
            start_date = _date(r.get("start_date"))
            end_date = _date(r.get("end_date"))
            order_date = _date(r.get("order_date"))
            qty = _int(r.get("qty"))
            province = r.get("province") or ("全国" if not r.get("city") else "")
            # 必填字段校验
            missing = []
            if not order_date:
                missing.append("更新日期")
            if not task_name:
                missing.append("任务名")
            if qty is None:
                missing.append("数量")
            if not start_date:
                missing.append("开始日期")
            if not end_date:
                missing.append("截止日期")
            if not up:
                missing.append("上游")
            if not cust:
                missing.append("一级代理")
            if missing:
                errors.append({"row": idx + 2, "reason": f"缺少必填字段：{'、'.join(missing)}"})
                continue
            gname = (r.get("group_name") or "").strip()
            cust_code = cust.code if cust else ""
            if gname:
                known = group_consistency.get(gname)
                if known:
                    if tpl_code and known["tpl_code"] and known["tpl_code"] != tpl_code:
                        errors.append({"row": idx + 2, "reason": f"小组「{gname}」的出数模版必须一致（文件内已有 {known['tpl_code']}，本次 {tpl_code}）"})
                        continue
                    if tpl_code and not known["tpl_code"]:
                        known["tpl_code"] = tpl_code
                    if known["cust_code"] != cust_code:
                        errors.append({"row": idx + 2, "reason": f"小组「{gname}」的一级代理必须一致（文件内已有 {known['cust_code'] or '空'}，本次 {cust_code}）"})
                        continue
                else:
                    group_consistency[gname] = {"tpl_code": tpl_code, "cust_code": cust_code}
            raw_status = str(r.get("status") or "").strip()
            import_status = _map_import_status(raw_status)
            up_id = up.id
            op_id = op.id if op else None
            duration = _duration_days(start_date, end_date) or r.get("duration")
            url_str = str(r.get("url") or "")
            add_name = str(r.get("add_name") or "").strip() in ("是", "true", "1", "True", "yes", "YES")
            check_collision = str(r.get("check_collision") or "").strip() in ("是", "true", "1", "True", "yes", "YES")

            # 按订单号匹配：订单号相同才更新，无订单号则插入
            order_no = (r.get("order_no") or "").strip()
            existing = existing_order_no.get(order_no) if order_no else None

            if existing:
                before_urls = sorted((u.url or "").strip() for u in db.query(OrderUrl).filter(OrderUrl.order_id == existing.id).all())
                before_vals = {"qty": existing.qty, "province": existing.province, "city": existing.city,
                               "excl_province": existing.excl_province, "excl_city": existing.excl_city,
                               "age_min": existing.age_min, "age_max": existing.age_max, "pv": existing.pv}
                existing.channel_id = ch.id if ch else None
                existing.operator_id = op.id if op else None
                existing.order_date = order_date
                raw_task_id = (r.get("task_id") or "").strip()
                if raw_task_id:
                    existing.task_id = raw_task_id
                elif up.name == "新":
                    existing.task_id = _gen_new_task_id(db, order_date, ch.name if ch else "", op.name if op else "",
                                                        province, r.get("city"), r.get("excl_city")) or existing.order_no
                else:
                    existing.task_id = "" if (up.name == "牛") else existing.order_no
                existing.duration = duration
                existing.qty = qty
                existing.province = province
                existing.city = r.get("city")
                existing.excl_province = r.get("excl_province")
                existing.excl_city = r.get("excl_city")
                existing.age_min = _int(r.get("age_min"))
                existing.age_max = _int(r.get("age_max"))
                existing.pv = _int(r.get("pv"))
                existing.start_date = start_date
                existing.end_date = end_date
                existing.price = _float(r.get("price"))
                existing.secondary_agent = r.get("secondary_agent")
                existing.platform = r.get("platform")
                existing.tpl_id = tpl.id if tpl else None
                existing.group_name = r.get("group_name")
                existing.export_filename = r.get("export_filename")
                existing.add_name = add_name
                existing.check_collision = check_collision
                if existing.status != "已停":
                    existing.status = import_status
                    if import_status == "已停":
                        existing.stop_date = date.today()
                existing.updated_at = datetime.now()
                db.query(OrderUrl).filter(OrderUrl.order_id == existing.id).delete()
                if url_str:
                    urls = [u for u in re.split(r"[|\n\r｜]+", url_str) if u.strip()]
                    for i, u in enumerate(urls):
                        db.add(OrderUrl(order_id=existing.id, url=u.strip(), level="高", sort_no=i))
                new_urls = sorted(u.strip() for u in re.split(r"[|\n\r｜]+", url_str) if u.strip())
                changed = []
                if before_urls != new_urls:
                    changed.append("url")
                new_vals = {"qty": qty, "province": province, "city": r.get("city"),
                            "excl_province": r.get("excl_province"), "excl_city": r.get("excl_city"),
                            "age_min": _int(r.get("age_min")), "age_max": _int(r.get("age_max")), "pv": _int(r.get("pv"))}
                for f in ("qty", "province", "city", "excl_province", "excl_city", "age_min", "age_max", "pv"):
                    if _norm(before_vals.get(f)) != _norm(new_vals.get(f)):
                        changed.append(f)
                existing.change_fields_json = json.dumps(changed, ensure_ascii=False) if changed else None
                updated += 1
                updated_nos.append(existing.order_no)
            else:
                order_no = r.get("order_no") or _next_order_no(db, order_date)
                raw_task_id = (r.get("task_id") or "").strip()
                if raw_task_id:
                    task_id = raw_task_id
                elif up.name == "新":
                    task_id = _gen_new_task_id(db, order_date, ch.name if ch else "", op.name if op else "",
                                               province, r.get("city"), r.get("excl_city")) or order_no
                else:
                    task_id = "" if (up.name == "牛") else order_no
                o = Order(order_no=order_no, customer_id=cust.id, upstream_id=up.id,
                          channel_id=ch.id if ch else None, operator_id=op.id if op else None, task_name=task_name,
                          task_id=task_id, duration=duration,
                          qty=qty, province=province, city=r.get("city"),
                          excl_province=r.get("excl_province"), excl_city=r.get("excl_city"),
                          age_min=_int(r.get("age_min")), age_max=_int(r.get("age_max")), pv=_int(r.get("pv")),
                          start_date=start_date, end_date=end_date, status=import_status,
                          stop_date=date.today() if import_status == "已停" else None,
                          order_date=order_date, price=_float(r.get("price")),
                          secondary_agent=r.get("secondary_agent"), platform=r.get("platform"),
                          tpl_id=tpl.id if tpl else None,
                           group_name=r.get("group_name"), export_filename=r.get("export_filename"),
                           add_name=add_name, check_collision=check_collision)
                db.add(o)
                db.flush()
                o.template_id = match_template(db, o)
                if url_str:
                    urls = [u for u in re.split(r"[|\n\r｜]+", url_str) if u.strip()]
                    for i, u in enumerate(urls):
                        db.add(OrderUrl(order_id=o.id, url=u.strip(), level="高", sort_no=i))
                imported += 1
                imported_nos.append(order_no)
        except Exception as e:  # noqa
            errors.append({"row": idx + 2, "reason": str(e)})
    _op_log(db, user, "import_orders", f"导入 {imported} 条，更新 {updated} 条，失败 {len(errors)} 条",
            None, {"imported": imported, "updated": updated, "failed": len(errors), "order_nos": imported_nos, "updated_nos": updated_nos})
    db.commit()
    return {"code": 0, "data": {"imported": imported, "updated": updated, "errors": errors},
            "msg": f"导入 {imported} 条，更新 {updated} 条，失败 {len(errors)} 条"}


@router.post("/import-receipt")
async def import_receipt(file: UploadFile = File(...), party: str = Form("牛"),
                         db: Session = Depends(get_db), _=Depends(get_current_user)):
    """导入回执：按甲方+任务名匹配订单，更新工单号；未匹配或多匹配的提示错误"""
    fn = file.filename or ""
    content = await file.read()
    if not fn.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=422, detail="仅支持 xlsx / xls 文件")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    header = None
    rows = []
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(v).strip() if v is not None else "" for v in values]
            continue
        rows.append(values)
    wb.close()

    def _find(*names):
        for i, h in enumerate(header):
            hn = re.sub(r"\s+", "", str(h)).lower()
            for n in names:
                if re.sub(r"\s+", "", str(n)).lower() in hn:
                    return i
        return -1

    name_col = _find("任务名")
    id_col = _find("任务ID", "工单号", "甲方订单号", "订单号", "id")
    op_col = _find("运营商")
    ch_col = _find("渠道", "类型")
    if name_col < 0:
        raise HTTPException(status_code=422, detail="回执文件缺少「任务名」列")
    if id_col < 0:
        raise HTTPException(status_code=422, detail="回执文件缺少「任务ID/工单号/id」列")
    if op_col < 0:
        raise HTTPException(status_code=422, detail="回执文件缺少「运营商」列")
    if ch_col < 0:
        raise HTTPException(status_code=422, detail="回执文件缺少「渠道/类型」列")

    # 甲方（来自前端选择）
    up = db.query(Customer).filter(Customer.ctype == "upstream",
                                   or_(Customer.name == party, Customer.code == party)).first() if party else None
    up_id = up.id if up else None

    updated = 0
    updated_list = []
    errors = []
    for r in rows:
        if r is None:
            continue
        task_name = str(r[name_col]).strip() if name_col < len(r) and r[name_col] is not None else ""
        new_task_id = str(r[id_col]).strip() if id_col < len(r) and r[id_col] is not None else ""
        operator_val = str(r[op_col]).strip() if op_col < len(r) and r[op_col] is not None else ""
        channel_val = str(r[ch_col]).strip() if ch_col < len(r) and r[ch_col] is not None else ""
        if not task_name:
            continue
        op = db.query(Operator).filter(Operator.name == operator_val).first() if operator_val else None
        ch = db.query(Channel).filter(Channel.name == channel_val).first() if channel_val else None
        orders = db.query(Order).filter(
            Order.task_name == task_name,
            Order.upstream_id == up_id,
            Order.operator_id == (op.id if op else None),
            Order.channel_id == (ch.id if ch else None),
            Order.status != "已停",
            or_(Order.task_id.is_(None), Order.task_id == ""),
        ).all()
        if len(orders) == 0:
            errors.append({"task_name": task_name, "reason": "未找到匹配订单（甲方+任务名+运营商+渠道 且工单号为空、非已停）"})
            continue
        if len(orders) > 1:
            errors.append({"task_name": task_name, "reason": f"匹配到 {len(orders)} 个订单"})
            continue
        if not new_task_id:
            errors.append({"task_name": task_name, "reason": "工单号为空"})
            continue
        orders[0].task_id = new_task_id
        orders[0].updated_at = datetime.now()
        updated += 1
        updated_list.append({
            "party": party or "",
            "task_id": new_task_id,
            "task_name": task_name,
            "operator": operator_val,
            "channel": channel_val,
        })
    db.commit()
    return {"code": 0, "data": {"updated": updated, "updated_list": updated_list, "errors": errors},
            "msg": f"更新 {updated} 条，失败 {len(errors)} 条"}


def _next_batch_no(db, order_date):
    """提单批次 = 日期YYYYMMDD + 当日生成次数序号（01 起）"""
    prefix = order_date.strftime("%Y%m%d") + "-"
    max_seq = 0
    for (bn,) in db.query(TidabiaoHistory.batch_no).filter(TidabiaoHistory.batch_no.like(prefix + "%")).all():
        m = re.match(r"^\d{8}-(\d{2})$", bn or "")
        if m:
            max_seq = max(max_seq, int(m.group(1)))
    return f"{prefix}{max_seq + 1:02d}"


def _history_from_order(db, o: Order, batch_no: str):
    """把订单快照成提单历史（全文本，无关联）"""
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first()
    cust = db.query(Customer).filter(Customer.id == o.customer_id).first()
    ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
    tpl = db.query(Template).filter(Template.id == o.tpl_id).first() if o.tpl_id else None
    ch_name = ch.name if ch else ""
    operator = _operator_name(db, o)
    return TidabiaoHistory(
        batch_no=batch_no,
        order_id=o.id,
        order_no=o.order_no or "",
        status=o.status or "",
        order_date=str(o.order_date) if o.order_date else "",
        upstream=up.name if up else "",
        customer=(f"{cust.code} {cust.name}" if cust else ""),
        secondary_agent=o.secondary_agent or "",
        channel=ch_name,
        operator=operator,
        task_name=o.task_name or "",
        task_id=o.task_id or "",
        url="\n".join(u.url for u in o.urls),
        qty=str(o.qty) if o.qty is not None else "",
        duration=o.duration or "",
        age_min=str(o.age_min) if o.age_min is not None else "",
        age_max=str(o.age_max) if o.age_max is not None else "",
        pv=str(o.pv) if o.pv is not None else "",
        province=o.province or "",
        city=o.city or "",
        excl_province=o.excl_province or "",
        excl_city=o.excl_city or "",
        start_date=str(o.start_date) if o.start_date else "",
        end_date=str(o.end_date) if o.end_date else "",
        price=str(o.price) if o.price is not None else "",
        platform=o.platform or "",
        group_name=o.group_name or "",
        tpl_code=tpl.code if tpl else "",
        add_name="是" if o.add_name else "否",
        check_collision="是" if o.check_collision else "否",
    )


def _history_dict(h: TidabiaoHistory) -> dict:
    return {
        "id": h.id, "batch_no": h.batch_no or "", "order_no": h.order_no or "",
        "status": h.status or "", "order_date": h.order_date or "",
        "upstream": h.upstream or "", "customer": h.customer or "",
        "secondary_agent": h.secondary_agent or "", "channel": h.channel or "",
        "operator": h.operator or "", "task_name": h.task_name or "",
        "task_id": h.task_id or "", "url": h.url or "", "qty": h.qty or "",
        "duration": h.duration or "", "age_min": h.age_min or "", "age_max": h.age_max or "",
        "pv": h.pv or "", "province": h.province or "", "city": h.city or "",
        "excl_province": h.excl_province or "", "excl_city": h.excl_city or "",
        "start_date": h.start_date or "", "end_date": h.end_date or "",
        "price": h.price or "", "platform": h.platform or "", "group_name": h.group_name or "",
        "tpl_code": h.tpl_code or "", "add_name": h.add_name or "",
        "created_at": fmt_dt(h.created_at),
    }


@router.get("/tidabiao-history")
def list_tidabiao_history(db: Session = Depends(get_db), _=Depends(get_current_user),
                          batch_no: str = "", status: str = "", upstream: str = "",
                          customer: str = "", channel: str = "", task_name: str = "",
                          task_id: str = "", order_no: str = "", order_date: str = "",
                          q: str = "", page: int = 1, per_page: int = 10):
    """提单历史（全文本快照）"""
    qy = db.query(TidabiaoHistory)
    if batch_no:
        qy = qy.filter(TidabiaoHistory.batch_no.like(f"%{batch_no}%"))
    if status:
        qy = qy.filter(TidabiaoHistory.status == status)
    if upstream:
        qy = qy.filter(TidabiaoHistory.upstream.like(f"%{upstream}%"))
    if customer:
        qy = qy.filter(TidabiaoHistory.customer.like(f"%{customer}%"))
    if channel:
        qy = qy.filter(TidabiaoHistory.channel.like(f"%{channel}%"))
    if task_name:
        qy = qy.filter(TidabiaoHistory.task_name.like(f"%{task_name}%"))
    if task_id:
        qy = qy.filter(TidabiaoHistory.task_id.like(f"%{task_id}%"))
    if order_no:
        qy = qy.filter(TidabiaoHistory.order_no.like(f"%{order_no}%"))
    if order_date:
        qy = qy.filter(TidabiaoHistory.order_date == order_date)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(TidabiaoHistory.order_no.like(like), TidabiaoHistory.task_name.like(like),
                           TidabiaoHistory.upstream.like(like), TidabiaoHistory.customer.like(like),
                           TidabiaoHistory.channel.like(like)))
    total, rows = paginate(qy.order_by(TidabiaoHistory.id.desc()), page, per_page)
    return ok_page([_history_dict(h) for h in rows], total)


@router.post("/generate-tidabiao")
def generate_tidabiao(body: GenerateBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    from urllib.parse import quote
    from ..tidabiao import generate_tidabiao_all
    # 更新日期 → MMDD 文件名 + date 过滤
    if body.date:
        mmdd = body.date.replace("-", "")[4:]
        try:
            date_obj = datetime.strptime(body.date, "%Y-%m-%d").date()
        except ValueError:
            date_obj = date.today()
    else:
        mmdd = datetime.now().strftime("%m%d")
        date_obj = date.today()
    # 按勾选订单生成；勾选时校验其更新日期与所选日期一致
    if body.order_ids:
        orders = db.query(Order).filter(Order.id.in_(body.order_ids)).all()
        if not orders:
            raise HTTPException(status_code=404, detail="没有找到勾选的订单")
        mismatch = [o for o in orders if o.order_date != date_obj]
        if mismatch:
            shown = "、".join(f"{o.task_name or o.order_no}({o.order_date})" for o in mismatch[:10])
            more = f" 等共{len(mismatch)}条" if len(mismatch) > 10 else ""
            raise HTTPException(status_code=400, detail=f"勾选订单的更新日期与所选日期（{body.date}）不一致：{shown}{more}")
    else:
        orders = db.query(Order).filter(Order.order_date == date_obj).all()
    if not orders:
        raise HTTPException(status_code=404, detail=f"更新日期 {body.date or date_obj.strftime('%Y-%m-%d')} 没有可生成的订单")
    try:
        buf, manifest, processed = generate_tidabiao_all(db, orders, mmdd)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not manifest:
        raise HTTPException(status_code=404, detail="没有可生成的订单（仅未提/改单/待停状态会生成），请检查订单状态或模版")
    filename = f"提单表-{mmdd}.zip"
    return Response(content=buf.getvalue(),
                    media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.post("/confirm-tidabiao")
def confirm_tidabiao(body: GenerateBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """提单确认：按提单日期更新订单状态并写入提单历史"""
    if body.date:
        try:
            date_obj = datetime.strptime(body.date, "%Y-%m-%d").date()
        except ValueError:
            date_obj = date.today()
    else:
        date_obj = date.today()
    if body.order_ids:
        orders = db.query(Order).filter(Order.id.in_(body.order_ids)).all()
        if not orders:
            raise HTTPException(status_code=404, detail="没有找到勾选的订单")
        mismatch = [o for o in orders if o.order_date != date_obj]
        if mismatch:
            shown = "、".join(f"{o.task_name or o.order_no}({o.order_date})" for o in mismatch[:10])
            more = f" 等共{len(mismatch)}条" if len(mismatch) > 10 else ""
            raise HTTPException(status_code=400, detail=f"勾选订单的更新日期与提单日期（{body.date}）不一致：{shown}{more}")
    else:
        orders = db.query(Order).filter(Order.order_date == date_obj).all()
    if not orders:
        raise HTTPException(status_code=404, detail=f"提单日期 {body.date or date_obj.strftime('%Y-%m-%d')} 没有可确认的订单")
    batch_no = _next_batch_no(db, date_obj)
    before = [{"order_no": o.order_no, "task_name": o.task_name, "status": o.status} for o in orders]
    for o in orders:
        if o.status in ("未提", "改单"):
            o.status = "在执"
        elif o.status == "待停":
            o.status = "已停"
        o.batch_no = batch_no
        o.updated_at = datetime.now()
        db.add(_history_from_order(db, o, batch_no))
    after = [{"order_no": o.order_no, "task_name": o.task_name, "status": o.status, "batch_no": o.batch_no} for o in orders]
    _op_log(db, user, "confirm_tidabiao", f"{batch_no}（{date_obj.strftime('%Y-%m-%d')}）", before, after)
    db.commit()
    return {"code": 0, "data": {"confirmed": len(orders), "batch_no": batch_no},
            "msg": f"提单确认完成：{len(orders)} 单（批次 {batch_no}）"}


def _int(v):
    try:
        return int(float(v)) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _float(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _duration_days(start_date, end_date):
    """时长 = 截止日期 - 开始日期（天数）"""
    if not start_date or not end_date:
        return ""
    try:
        s = datetime.strptime(str(start_date), "%Y-%m-%d").date()
        e = datetime.strptime(str(end_date), "%Y-%m-%d").date()
        return str((e - s).days)
    except (ValueError, TypeError):
        return ""


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    year = datetime.now().year
    if re.fullmatch(r"\d{4}", s):
        s = f"{year}{s}"
    elif re.fullmatch(r"\d{1,2}-\d{1,2}", s):
        s = f"{year}-{s}"
    elif re.fullmatch(r"\d{1,2}/\d{1,2}", s):
        s = f"{year}/{s}"
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _norm_header(h):
    return re.sub(r"\s+", "", str(h)).lower()


def _map_row(header, values):
    idx = {_norm_header(h): i for i, h in enumerate(header)}
    row = {}
    for name, field in ORDER_HEADER_MAP.items():
        key = _norm_header(name)
        if key in idx and idx[key] < len(values):
            v = values[idx[key]]
            if v is None:
                row[field] = None
            elif isinstance(v, date):
                row[field] = v.strftime("%Y-%m-%d")
            else:
                row[field] = str(v).strip()
    return row


def _parse_csv(content):
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    if not header:
        return []
    return [_map_row(header, values) for values in reader]


def _parse_xlsx(content):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows, header = [], None
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = values
            continue
        rows.append(_map_row(header, values))
    wb.close()
    return rows
