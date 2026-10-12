"""模版管理路由：模版文件（模版类型/模版编号 + 文件上传下载预览）"""
import io
import re
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user
from ..models import Template
from ..schemas import TemplateBody
from ..utils import fmt_dt

router = APIRouter(prefix="/api/templates", tags=["templates"])


def _next_code(db):
    max_n = 0
    for (code,) in db.query(Template.code).all():
        m = re.match(r"^MB-(\d+)$", code or "")
        if m:
            max_n = max(max_n, int(m.group(1)))
    return f"MB-{max_n + 1:03d}"


def _dump(t: Template) -> dict:
    return {
        "id": t.id, "ttype": t.ttype or "", "tpl_type": t.tpl_type or "其他", "code": t.code,
        "name": t.name or "", "file_name": t.name or "",
        "file_type": t.file_type or "",
        "has_file": bool(t.file_data),
        "description": t.description or "",
        "status": t.status,
        "created_at": fmt_dt(t.created_at), "updated_at": fmt_dt(t.updated_at),
    }


@router.get("")
def list_templates(db: Session = Depends(get_db), _=Depends(get_current_user),
                   q: str = ""):
    query = db.query(Template)
    if q:
        like = f"%{q}%"
        query = query.filter((Template.ttype.like(like)) | (Template.code.like(like)) | (Template.name.like(like)))
    rows = query.order_by(Template.id).all()
    return {"code": 0, "data": [_dump(t) for t in rows], "msg": "ok"}


@router.post("")
def create_template(body: TemplateBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    code = (body.code or "").strip() or _next_code(db)
    if db.query(Template).filter(Template.code == code).first():
        raise HTTPException(status_code=409, detail="模版编号已存在")
    t = Template(ttype=body.ttype, tpl_type=body.tpl_type or "其他", code=code, name="", description=body.description or "", status=body.status)
    db.add(t)
    db.commit()
    return {"code": 0, "data": {"id": t.id, "code": code}, "msg": "已创建"}


@router.put("/{tid}")
def update_template(tid: int, body: TemplateBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    t.ttype = body.ttype
    t.tpl_type = body.tpl_type or "其他"
    t.description = body.description or ""
    t.status = body.status
    db.commit()
    return {"code": 0, "data": None, "msg": "已保存"}


@router.delete("/{tid}")
def delete_template(tid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    db.delete(t)
    db.commit()
    return {"code": 0, "data": None, "msg": "已删除"}


@router.post("/{tid}/upload")
async def upload_template(tid: int, file: UploadFile = File(...),
                          db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    content = await file.read()
    fn = file.filename or ""
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    t.file_data = content
    t.file_type = ext
    if fn:
        t.name = fn
    db.commit()
    return {"code": 0, "data": {"filename": fn}, "msg": "模版文件已上传"}


@router.delete("/{tid}/file")
def delete_template_file(tid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    t.file_data = None
    t.file_type = ""
    t.name = ""
    db.commit()
    return {"code": 0, "data": None, "msg": "模版文件已删除"}


@router.get("/{tid}/preview")
def preview_template(tid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    if not t.file_data:
        raise HTTPException(status_code=404, detail="该模版未上传文件")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(bytes(t.file_data)), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({"name": ws.title, "html": _ws_to_html(ws)})
    wb.close()
    return {"code": 0, "data": {"filename": t.name or t.code, "sheets": sheets}, "msg": "ok"}


def _color_hex(color):
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and rgb:
        rgb = rgb.strip().upper()
        if len(rgb) == 8 and rgb.startswith("00"):
            rgb = rgb[2:]
        if len(rgb) == 6:
            try:
                int(rgb, 16)
                return f"#{rgb}"
            except ValueError:
                return None
    return None


def _ws_to_html(ws, max_rows=200, max_cols=60):
    """把工作表转成带样式的 HTML 表格，尽量还原 Excel 外观"""
    from openpyxl.utils import get_column_letter
    from html import escape

    max_row = min(ws.max_row or 1, max_rows)
    max_col = min(ws.max_column or 1, max_cols)

    # 合并单元格：记录每个左上角单元格的 span，以及被覆盖的单元格
    span_map = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row > max_row or rng.min_col > max_col:
            continue
        r0, r1 = rng.min_row, min(rng.max_row, max_row)
        c0, c1 = rng.min_col, min(rng.max_col, max_col)
        span_map[(r0, c0)] = (r1 - r0 + 1, c1 - c0 + 1)
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) != (r0, c0):
                    covered.add((r, c))

    # 列宽
    cols_html = []
    for c in range(1, max_col + 1):
        w = ws.column_dimensions[get_column_letter(c)].width
        cols_html.append(f'<col style="{("width:%dpx" % round(w * 7)) if w else ""}">')

    rows_html = []
    for r in range(1, max_row + 1):
        rd = ws.row_dimensions[r]
        rh = rd.height if rd and rd.height else None
        row_style = f' style="height:{round(rh)}px"' if rh else ""
        tds = []
        for c in range(1, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            val = "" if cell.value is None else str(cell.value)
            span = span_map.get((r, c))
            span_attrs = ""
            if span and (span[0] > 1 or span[1] > 1):
                span_attrs = f' rowspan="{span[0]}" colspan="{span[1]}"'
            css = []
            f = cell.font
            if f and f.bold:
                css.append("font-weight:bold")
            if f and f.italic:
                css.append("font-style:italic")
            if f and f.underline and str(f.underline) != "none":
                css.append("text-decoration:underline")
            if f and f.size:
                css.append(f"font-size:{f.size}pt")
            if f and f.name:
                css.append(f"font-family:'{f.name}'")
            fg = _color_hex(f.color) if f else None
            if fg:
                css.append(f"color:{fg}")
            if cell.fill and cell.fill.patternType == "solid":
                bg = _color_hex(cell.fill.fgColor)
                if bg:
                    css.append(f"background-color:{bg}")
            if cell.alignment:
                a = cell.alignment
                if a.horizontal:
                    css.append(f"text-align:{a.horizontal}")
                if a.vertical:
                    css.append(f"vertical-align:{a.vertical}")
                if a.wrap_text:
                    css.append("white-space:pre-wrap")
            b = cell.border
            if b:
                for side in ("left", "right", "top", "bottom"):
                    bd = getattr(b, side)
                    if bd and bd.style:
                        col = _color_hex(bd.color) or "#000"
                        css.append(f"border-{side}:1px solid {col}")
            style = f' style="{";".join(css)}"' if css else ""
            tds.append(f"<td{span_attrs}{style}>{escape(val)}</td>")
        rows_html.append(f"<tr{row_style}>{''.join(tds)}</tr>")

    return f'<table class="excel-preview"><colgroup>{"".join(cols_html)}</colgroup><tbody>{"".join(rows_html)}</tbody></table>'


@router.get("/{tid}/download")
def download_template(tid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    t = db.query(Template).filter(Template.id == tid).first()
    if not t:
        raise HTTPException(status_code=404, detail="模版不存在")
    if not t.file_data:
        raise HTTPException(status_code=404, detail="该模版未上传文件")
    ext = t.file_type or "xlsx"
    filename = t.name or t.code
    if "." not in filename:
        filename = f"{filename}.{ext}"
    media = {
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xls": "application/vnd.ms-excel",
    }.get(ext, "application/octet-stream")
    return Response(content=bytes(t.file_data), media_type=media,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
