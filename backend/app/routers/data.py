"""数据路由：日活数据/公积金/账单/预警"""
import csv
import io
import json
import re
import threading
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import List

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy import func, insert, or_, update, bindparam
from sqlalchemy.orm import Session

from ..database import get_db, SessionLocal
from ..deps import get_current_user, require_admin
from ..models import (DailyData, Fund, Bill, Alert, Customer, Channel, Category,
                      Order, SysConfig, Operator, SourceFile, Template, WashName, Platform,
                      CustomerRecharge, CustomerPrice, Notification, Product, ProductSale)
from ..pagination import paginate, ok_page
from ..excel import xlsx_response
from ..utils import fmt_dt, fmt_date
from ..schemas import BatchStopBody, GenerateBody, ExtractBody

router = APIRouter(prefix="/api", tags=["data"])

_DISTRIBUTE_TASKS = {}

# 甲方账单模版编号（在模版中心上传对应模版文件）
UPSTREAM_BILL_TPL_CODE = "MB-032"


def _balance_suffix(balance):
    """账单文件名余额后缀：正/负"""
    b = round(float(balance or 0), 2)
    return f"（正{b:g}）" if b >= 0 else f"（负{abs(b):g}）"


class DistributeCancelled(Exception):
    pass

FUND_HEADER_MAP = {
    "手机号": "phone", "姓名": "name", "身份证号": "id_card", "性别": "gender",
    "省份": "province", "地市": "city", "单位名称": "company", "单位性质": "company_type",
    "缴存基数": "base", "缴存比例": "ratio", "月缴存额": "monthly", "账户余额": "balance",
    "缴存状态": "deposit_status", "开户日期": "open_date", "缴至年月": "pay_to",
    "运营商": "operator", "来源文件": "source_file",
}


# ============ 日活数据 ============
@router.get("/daily-data")
def list_daily(db: Session = Depends(get_db), _=Depends(get_current_user),
               q: str = "", phone: str = "", upstream: str = "", task_id: str = "",
               task_name: str = "", province: str = "", city: str = "",
               operator: str = "", cat1: str = "", cat2: str = "", platform: str = "",
                source_file: str = "", customer: str = "", secondary_agent: str = "",
                channel: str = "", group_name: str = "", biz_date: str = "", name: str = "",
                start_date: str = "", end_date: str = "", platform_empty: str = "",
                page: int = 1, per_page: int = 10):
    qy = db.query(DailyData)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(DailyData.phone.like(like), DailyData.task_name.like(like),
                           DailyData.upstream.like(like), DailyData.task_id.like(like)))
    if phone:
        qy = qy.filter(DailyData.phone.like(f"%{phone}%"))
    if upstream:
        qy = qy.filter(DailyData.upstream.like(f"%{upstream}%"))
    if task_id:
        qy = qy.filter(DailyData.task_id.like(f"%{task_id}%"))
    if task_name:
        qy = qy.filter(DailyData.task_name.like(f"%{task_name}%"))
    if province:
        provs = [p for p in province.split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[DailyData.province.like(f"%{p}%") for p in provs]))
    if city:
        cities = [c for c in city.split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[DailyData.city.like(f"%{c}%") for c in cities]))
    if operator:
        qy = qy.filter(DailyData.operator.like(f"%{operator}%"))
    if cat1:
        qy = qy.filter(DailyData.cat1.like(f"%{cat1}%"))
    if cat2:
        qy = qy.filter(DailyData.cat2.like(f"%{cat2}%"))
    if platform:
        qy = qy.filter(DailyData.platform.like(f"%{platform}%"))
    if platform_empty in ("1", "有", "是", "true", "True", "yes", "YES"):
        qy = qy.filter(or_(DailyData.platform.is_(None), DailyData.platform == ""))
    if source_file:
        qy = qy.filter(DailyData.source_file.like(f"%{source_file}%"))
    if customer:
        qy = qy.filter(DailyData.customer.like(f"%{customer}%"))
    if secondary_agent:
        qy = qy.filter(DailyData.secondary_agent.like(f"%{secondary_agent}%"))
    if channel:
        qy = qy.filter(DailyData.channel.like(f"%{channel}%"))
    if group_name:
        qy = qy.filter(DailyData.group_name.like(f"%{group_name}%"))
    if name:
        qy = qy.filter(DailyData.name.like(f"%{name}%"))
    if biz_date:
        qy = qy.filter(DailyData.biz_date == biz_date)
    if start_date:
        sd = _parse_biz_date(start_date)
        if sd:
            qy = qy.filter(DailyData.biz_date >= sd)
    if end_date:
        ed = _parse_biz_date(end_date)
        if ed:
            qy = qy.filter(DailyData.biz_date <= ed)
    total, rows = paginate(qy.order_by(DailyData.id.desc()), page, per_page)
    return ok_page([_daily_dict(d) for d in rows], total)


@router.put("/daily-data/{did}")
def update_daily(did: int, body: dict, db: Session = Depends(get_db), _=Depends(get_current_user)):
    d = db.query(DailyData).filter(DailyData.id == did).first()
    if not d:
        raise HTTPException(status_code=404, detail="日活数据不存在")
    editable = ["biz_date", "upstream", "task_id", "task_name", "phone", "name",
                "province", "city", "operator", "cat1", "cat2", "platform",
                "customer", "secondary_agent", "channel"]
    for field in editable:
        if field not in body:
            continue
        val = body[field]
        if field == "biz_date":
            parsed = _parse_biz_date(val) if val else None
            if parsed:
                d.biz_date = parsed
        elif field == "phone":
            pv = str(val or "").strip()
            if pv:
                d.phone = pv
        else:
            setattr(d, field, (str(val).strip() if val is not None and str(val).strip() else None))
    d.updated_at = datetime.now()
    db.commit()
    return {"code": 0, "data": None, "msg": "已更新"}


def _daily_dict(d: DailyData):
    return {
        "id": d.id, "biz_date": fmt_date(d.biz_date) if d.biz_date else None,
        "upstream": d.upstream or "", "task_id": d.task_id or "",
        "task_name": d.task_name or "", "phone": d.phone or "",
        "name": d.name or "",
        "province": d.province or "", "city": d.city or "",
        "operator": d.operator or "",
        "cat1": d.cat1 or "", "cat2": d.cat2 or "",
        "platform": d.platform or "",
        "customer": d.customer or "", "secondary_agent": d.secondary_agent or "",
        "channel": d.channel or "", "group_name": d.group_name or "",
        "source_file": d.source_file or "",
        "source_file_id": d.source_file_id,
        "created_at": fmt_dt(d.created_at), "updated_at": fmt_dt(d.updated_at),
    }


DAILY_EXPORT_COLS = [
    ("biz_date", "数据日期", lambda d: fmt_date(d.biz_date) if d.biz_date else ""),
    ("upstream", "甲方", lambda d: d.upstream or ""),
    ("task_id", "任务id", lambda d: d.task_id or ""),
    ("task_name", "任务名", lambda d: d.task_name or ""),
    ("phone", "手机号", lambda d: d.phone or ""),
    ("name", "姓名", lambda d: d.name or ""),
    ("province", "省", lambda d: d.province or ""),
    ("city", "市", lambda d: d.city or ""),
    ("operator", "运营商", lambda d: d.operator or ""),
    ("cat1", "一级品类", lambda d: d.cat1 or ""),
    ("cat2", "二级品类", lambda d: d.cat2 or ""),
    ("platform", "平台", lambda d: d.platform or ""),
    ("customer", "一级代理", lambda d: d.customer or ""),
    ("secondary_agent", "二级代理", lambda d: d.secondary_agent or ""),
    ("channel", "渠道", lambda d: d.channel or ""),
    ("group_name", "小组", lambda d: d.group_name or ""),
    ("source_file", "来源文件名", lambda d: d.source_file or ""),
    ("created_at", "创建时间", lambda d: fmt_dt(d.created_at)),
    ("updated_at", "更新时间", lambda d: fmt_dt(d.updated_at)),
]


@router.get("/daily-data/export")
def export_daily(db: Session = Depends(get_db), _=Depends(get_current_user),
                 q: str = "", phone: str = "", upstream: str = "", task_id: str = "",
                 task_name: str = "", province: str = "", city: str = "",
                 operator: str = "", cat1: str = "", cat2: str = "", platform: str = "",
                 source_file: str = "", customer: str = "", secondary_agent: str = "",
                 channel: str = "", group_name: str = "", biz_date: str = "", name: str = "", cols: str = "",
                 start_date: str = "", end_date: str = "", platform_empty: str = ""):
    qy = db.query(DailyData)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(DailyData.phone.like(like), DailyData.task_name.like(like),
                           DailyData.upstream.like(like), DailyData.task_id.like(like)))
    if phone:
        qy = qy.filter(DailyData.phone.like(f"%{phone}%"))
    if upstream:
        qy = qy.filter(DailyData.upstream.like(f"%{upstream}%"))
    if task_id:
        qy = qy.filter(DailyData.task_id.like(f"%{task_id}%"))
    if task_name:
        qy = qy.filter(DailyData.task_name.like(f"%{task_name}%"))
    if province:
        provs = [p for p in province.split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[DailyData.province.like(f"%{p}%") for p in provs]))
    if city:
        cities = [c for c in city.split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[DailyData.city.like(f"%{c}%") for c in cities]))
    if operator:
        qy = qy.filter(DailyData.operator.like(f"%{operator}%"))
    if cat1:
        qy = qy.filter(DailyData.cat1.like(f"%{cat1}%"))
    if cat2:
        qy = qy.filter(DailyData.cat2.like(f"%{cat2}%"))
    if platform:
        qy = qy.filter(DailyData.platform.like(f"%{platform}%"))
    if platform_empty in ("1", "有", "是", "true", "True", "yes", "YES"):
        qy = qy.filter(or_(DailyData.platform.is_(None), DailyData.platform == ""))
    if source_file:
        qy = qy.filter(DailyData.source_file.like(f"%{source_file}%"))
    if customer:
        qy = qy.filter(DailyData.customer.like(f"%{customer}%"))
    if secondary_agent:
        qy = qy.filter(DailyData.secondary_agent.like(f"%{secondary_agent}%"))
    if channel:
        qy = qy.filter(DailyData.channel.like(f"%{channel}%"))
    if group_name:
        qy = qy.filter(DailyData.group_name.like(f"%{group_name}%"))
    if name:
        qy = qy.filter(DailyData.name.like(f"%{name}%"))
    if biz_date:
        qy = qy.filter(DailyData.biz_date == biz_date)
    if start_date:
        sd = _parse_biz_date(start_date)
        if sd:
            qy = qy.filter(DailyData.biz_date >= sd)
    if end_date:
        ed = _parse_biz_date(end_date)
        if ed:
            qy = qy.filter(DailyData.biz_date <= ed)
    rows = qy.order_by(DailyData.id.desc())
    selected = [c for c in (cols or "").split(",") if c.strip()]
    spec = [(k, l, fn) for (k, l, fn) in DAILY_EXPORT_COLS if not selected or k in selected]
    headers = [l for _, l, _ in spec]

    def gen():
        for d in rows.yield_per(2000):
            yield [fn(d) for _, _, fn in spec]

    return xlsx_response(headers, gen(), "日活数据.xlsx", "日活数据")


@router.get("/daily-data/export-csv")
def export_daily_csv(db: Session = Depends(get_db), _=Depends(get_current_user),
                     q: str = "", phone: str = "", upstream: str = "", task_id: str = "",
                     task_name: str = "", province: str = "", city: str = "",
                     operator: str = "", cat1: str = "", cat2: str = "", platform: str = "",
                     source_file: str = "", customer: str = "", secondary_agent: str = "",
                     channel: str = "", group_name: str = "", biz_date: str = "", name: str = "",
                     start_date: str = "", end_date: str = "", cols: str = ""):
    from urllib.parse import quote
    from fastapi.responses import StreamingResponse

    f = {k: v for k, v in {
        "q": q, "phone": phone, "upstream": upstream, "task_id": task_id,
        "task_name": task_name, "province": province, "city": city, "operator": operator,
        "cat1": cat1, "cat2": cat2, "platform": platform, "source_file": source_file,
        "customer": customer, "secondary_agent": secondary_agent, "channel": channel,
        "group_name": group_name, "biz_date": biz_date, "name": name,
        "start_date": start_date, "end_date": end_date,
    }.items() if v}
    qy = _apply_daily_filters(db.query(DailyData), f).order_by(DailyData.id.desc())
    selected = [c for c in (cols or "").split(",") if c.strip()]
    spec = [(k, l, fn) for (k, l, fn) in DAILY_EXPORT_COLS if not selected or k in selected]
    headers = [l for _, l, _ in spec]

    def _esc(v):
        s = str(v) if v is not None else ""
        if any(ch in s for ch in (',', '"', '\n', '\r')):
            return '"' + s.replace('"', '""') + '"'
        return s

    def gen():
        yield "\ufeff" + ",".join(_esc(h) for h in headers) + "\r\n"
        for d in qy.yield_per(5000):
            yield ",".join(_esc(fn(d)) for _, _, fn in spec) + "\r\n"

    return StreamingResponse(gen(), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('日活数据.csv')}"})


def _parse_biz_date(v):
    if not v:
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _order_info_by_task_id(db, task_id):
    """按工单号关联订单，返回 甲方/任务名/运营商/一级代理/二级代理/渠道/平台"""
    if not task_id:
        return {}
    o = db.query(Order).filter(Order.task_id == task_id).first()
    if not o:
        return {}
    up = db.query(Customer).filter(Customer.id == o.upstream_id).first()
    cust = db.query(Customer).filter(Customer.id == o.customer_id).first()
    op = db.query(Operator).filter(Operator.id == o.operator_id).first()
    ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
    return {
        "upstream": up.name if up else "",
        "task_name": o.task_name or "",
        "operator": op.name if op else "",
        "customer": (f"{cust.code} {cust.name}" if cust else ""),
        "secondary_agent": o.secondary_agent or "",
        "channel": ch.name if ch else "",
        "platform": o.platform or "",
        "group_name": o.group_name or "",
    }


def _wash_name_by_phone(db, phone):
    """按手机号从洗名库取姓名"""
    if not phone:
        return ""
    w = db.query(WashName).filter(WashName.phone == phone).first()
    return w.name if w and w.name else ""


def _build_import_ctx(db):
    """预加载导入所需的参考数据，避免逐行 N+1 查询"""
    wash = {w.phone: w.name for w in db.query(WashName).filter(WashName.name.isnot(None), WashName.name != "").all()}
    platforms = {p.name: p.cat_id for p in db.query(Platform).all()}
    cats = {c.id: (c.name, c.parent_id) for c in db.query(Category).all()}
    customers = {c.id: c for c in db.query(Customer).all()}
    operators = {op.id: op.name for op in db.query(Operator).all()}
    channels = {ch.id: ch.name for ch in db.query(Channel).all()}
    orders = {o.task_id: o for o in db.query(Order).filter(Order.task_id.isnot(None)).all()}
    return {"wash": wash, "platforms": platforms, "cats": cats,
            "customers": customers, "operators": operators, "channels": channels, "orders": orders}


def _ctx_order_info(ctx, task_id):
    if not task_id:
        return {}
    o = ctx["orders"].get(task_id)
    if not o:
        return {}
    up = ctx["customers"].get(o.upstream_id)
    cust = ctx["customers"].get(o.customer_id)
    return {
        "upstream": up.name if up else "",
        "task_name": o.task_name or "",
        "operator": ctx["operators"].get(o.operator_id, ""),
        "customer": (f"{cust.code} {cust.name}" if cust else ""),
        "secondary_agent": o.secondary_agent or "",
        "channel": ctx["channels"].get(o.channel_id, ""),
        "platform": o.platform or "",
        "group_name": o.group_name or "",
    }


def _ctx_wash(ctx, phone):
    if not phone:
        return ""
    return ctx["wash"].get(phone, "")


def _ctx_cat(ctx, platform_name):
    if not platform_name:
        return "", ""
    cat_id = ctx["platforms"].get(platform_name)
    if not cat_id:
        return "", ""
    cat2_name, parent_id = ctx["cats"].get(cat_id, (None, None))
    if not cat2_name:
        return "", ""
    cat1_name = ctx["cats"].get(parent_id, (None, None))[0] if parent_id else None
    return cat1_name or "", cat2_name or ""


def _cat_by_platform(db, platform_name):
    """按平台名称关联品类管理，返回 (一级品类, 二级品类)"""
    if not platform_name:
        return "", ""
    p = db.query(Platform).filter(Platform.name == platform_name).first()
    if not p:
        return "", ""
    cat2 = db.query(Category).filter(Category.id == p.cat_id).first()
    if not cat2:
        return "", ""
    cat1 = db.query(Category).filter(Category.id == cat2.parent_id).first()
    return (cat1.name if cat1 else ""), (cat2.name or "")


def _is_header_row(values):
    """第一行是否为表头：含「任务id/手机号/运营商」等文字则视为表头"""
    if not values:
        return False
    joined = "".join(str(v) for v in values if v is not None).lower()
    return any(k in joined for k in ("任务id", "任务名", "手机号", "运营商", "省", "市"))


def _import_niu_excel(db, content, biz_date, source_file, source_file_id, errors, seen, ctx=None):
    """解析牛的数据 excel（第一行是表头）：任务id/任务名/手机号/省/市/运营商"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    n = 0
    deduped = 0
    header = None
    batch = []
    now = datetime.now()
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(v).strip() if v is not None else "" for v in values]
            continue
        if not any((str(v) if v is not None else "").strip() for v in values):
            continue
        def col(*names):
            for i, h in enumerate(header):
                hn = re.sub(r"\s+", "", h).lower()
                for nm in names:
                    if re.sub(r"\s+", "", nm).lower() in hn:
                        return i
            return -1
        tid = col("任务id")
        tnm = col("任务名")
        pn = col("手机号")
        prov = col("省")
        cty = col("市")
        op = col("运营商")
        task_id = str(values[tid]).strip() if tid >= 0 and tid < len(values) and values[tid] is not None else ""
        task_name = str(values[tnm]).strip() if tnm >= 0 and tnm < len(values) and values[tnm] is not None else ""
        phone = str(values[pn]).strip() if pn >= 0 and pn < len(values) and values[pn] is not None else ""
        province = str(values[prov]).strip() if prov >= 0 and prov < len(values) and values[prov] is not None else ""
        city = str(values[cty]).strip() if cty >= 0 and cty < len(values) and values[cty] is not None else ""
        operator = str(values[op]).strip() if op >= 0 and op < len(values) and values[op] is not None else ""
        if not phone:
            continue
        key = (task_id, phone)
        if key in seen:
            deduped += 1
            continue
        seen.add(key)
        info = _ctx_order_info(ctx, task_id) if ctx else _order_info_by_task_id(db, task_id)
        cat1, cat2 = _ctx_cat(ctx, info.get("platform") or "") if ctx else _cat_by_platform(db, info.get("platform") or "")
        batch.append({
            "biz_date": biz_date,
            "upstream": info.get("upstream") or "",
            "task_id": task_id,
            "task_name": info.get("task_name") or task_name or "",
            "phone": phone,
            "name": _ctx_wash(ctx, phone) if ctx else _wash_name_by_phone(db, phone),
            "province": province,
            "city": city,
            "operator": operator or info.get("operator") or "",
            "cat1": cat1,
            "cat2": cat2,
            "platform": info.get("platform") or "",
            "customer": info.get("customer") or "",
            "secondary_agent": info.get("secondary_agent") or "",
            "channel": info.get("channel") or "",
            "group_name": info.get("group_name") or "",
            "source_file": source_file,
            "source_file_id": source_file_id,
            "created_at": now,
            "updated_at": now,
        })
        n += 1
        if len(batch) >= 5000:
            db.execute(insert(DailyData), batch)
            batch = []
    if batch:
        db.execute(insert(DailyData), batch)
    wb.close()
    return n, deduped


def _import_new_excel(db, content, biz_date, source_file, source_file_id, errors, seen, ctx=None):
    """解析新的数据 excel（第一行是数据，无表头）：手机号/甲方工单号/平台/省/市"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    n = 0
    deduped = 0
    batch = []
    now = datetime.now()
    for values in ws.iter_rows(values_only=True):
        if not any((str(v) if v is not None else "").strip() for v in values):
            continue
        vals = [str(v).strip() if v is not None else "" for v in values]
        phone = vals[0] if len(vals) > 0 else ""
        task_id = vals[1] if len(vals) > 1 else ""
        province = vals[3] if len(vals) > 3 else ""
        city = vals[4] if len(vals) > 4 else ""
        if not phone:
            continue
        key = (task_id, phone)
        if key in seen:
            deduped += 1
            continue
        seen.add(key)
        info = _ctx_order_info(ctx, task_id) if ctx else _order_info_by_task_id(db, task_id)
        cat1, cat2 = _ctx_cat(ctx, info.get("platform") or "") if ctx else _cat_by_platform(db, info.get("platform") or "")
        batch.append({
            "biz_date": biz_date,
            "upstream": info.get("upstream") or "",
            "task_id": task_id,
            "task_name": info.get("task_name") or "",
            "phone": phone,
            "name": _ctx_wash(ctx, phone) if ctx else _wash_name_by_phone(db, phone),
            "province": province,
            "city": city,
            "operator": info.get("operator") or "",
            "cat1": cat1,
            "cat2": cat2,
            "platform": info.get("platform") or "",
            "customer": info.get("customer") or "",
            "secondary_agent": info.get("secondary_agent") or "",
            "channel": info.get("channel") or "",
            "group_name": info.get("group_name") or "",
            "source_file": source_file,
            "source_file_id": source_file_id,
            "created_at": now,
            "updated_at": now,
        })
        n += 1
        if len(batch) >= 5000:
            db.execute(insert(DailyData), batch)
            batch = []
    if batch:
        db.execute(insert(DailyData), batch)
    wb.close()
    return n, deduped


def _import_one_daily_file(db, fn, content, date_obj, source_file_id, errors, seen, ctx=None):
    """处理单个文件（zip 或 excel），返回 (导入条数, 去重条数)"""
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    if ext not in ("zip", "xlsx", "xls"):
        errors.append(f"{fn}: 仅支持 zip / xlsx / xls 文件")
        return 0, 0
    n = 0
    deduped = 0
    if ext == "zip":
        import zipfile
        try:
            zf = zipfile.ZipFile(io.BytesIO(content))
            for name in zf.namelist():
                if name.lower().endswith((".xlsx", ".xls")):
                    try:
                        nn, dd = _import_niu_excel(db, zf.read(name), date_obj, fn, source_file_id, errors, seen, ctx)
                        n += nn
                        deduped += dd
                    except Exception as e:  # noqa
                        errors.append(f"{name}: {e}")
            zf.close()
        except Exception as e:  # noqa
            errors.append(f"{fn}: 解压失败 {e}")
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        first = None
        for values in ws.iter_rows(values_only=True):
            first = values
            break
        wb.close()
        if _is_header_row(first):
            nn, dd = _import_niu_excel(db, content, date_obj, fn, source_file_id, errors, seen, ctx)
        else:
            nn, dd = _import_new_excel(db, content, date_obj, fn, source_file_id, errors, seen, ctx)
        n += nn
        deduped += dd
    return n, deduped


def _extract_task_ids_from_excel(content):
    """解析 excel，返回所有（有手机号行的）任务id列表（可能含空串）"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    tids = []
    header = None
    header_mode = False
    tid_i = phone_i = -1

    def find_idx(names):
        for i, h in enumerate(header or []):
            hn = re.sub(r"\s+", "", h).lower()
            for nm in names:
                if re.sub(r"\s+", "", nm).lower() in hn:
                    return i
        return -1

    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(v).strip() if v is not None else "" for v in values]
            if _is_header_row(values):
                header_mode = True
                tid_i = find_idx(["任务id"])
                phone_i = find_idx(["手机号"])
                continue  # 表头，跳过
        if not any((str(v) if v is not None else "").strip() for v in values):
            continue
        vals = [str(v).strip() if v is not None else "" for v in values]
        if header_mode:
            phone = vals[phone_i] if phone_i >= 0 and phone_i < len(vals) else ""
            tid = vals[tid_i] if tid_i >= 0 and tid_i < len(vals) else ""
        else:
            phone = vals[0] if len(vals) > 0 else ""
            tid = vals[1] if len(vals) > 1 else ""
        if phone:
            tids.append(tid)
    wb.close()
    return tids


def _extract_task_ids(fn, content):
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    tids = []
    if ext == "zip":
        import zipfile
        zf = zipfile.ZipFile(io.BytesIO(content))
        for name in zf.namelist():
            if name.lower().endswith((".xlsx", ".xls")):
                tids += _extract_task_ids_from_excel(zf.read(name))
        zf.close()
    elif ext in ("xlsx", "xls"):
        tids += _extract_task_ids_from_excel(content)
    return tids


def _count_unmatched(db, file_list):
    """统计未匹配到订单的日活数据条数，返回 (条数, 未匹配工单明细列表)"""
    order_tids = {r[0] for r in db.query(Order.task_id).all() if r[0]}
    total = 0
    cnt = defaultdict(int)
    for fn, content in file_list:
        for tid in _extract_task_ids(fn, content):
            if tid not in order_tids:
                total += 1
                if tid:
                    cnt[tid] += 1
    items = [{"task_id": tid, "task_name": "", "channel": "", "reason": "未匹配到订单（系统中无此工单号）",
              "qty": cnt[tid]} for tid in sorted(cnt)]
    return total, items


def _count_stopped_abnormal(db, file_list, date_obj):
    """统计关联到已停订单、且数据日期已过截止日期 T+1（末次出数日期 end_date+1）仍在出数的异常数据"""
    stopped = {o.task_id: o for o in db.query(Order).filter(Order.status == "已停").all() if o.task_id}
    chans = {c.id: c.name for c in db.query(Channel).all()}
    total = 0
    tids = set()
    cnt = defaultdict(int)
    for fn, content in file_list:
        for tid in _extract_task_ids(fn, content):
            o = stopped.get(tid)
            if o and o.end_date and date_obj > o.end_date + timedelta(days=1):
                total += 1
                tids.add(tid)
                cnt[tid] += 1
    items = []
    for tid in sorted(tids):
        o = stopped.get(tid)
        end_str = fmt_date(o.end_date) if (o and o.end_date) else ""
        items.append({"task_id": tid, "task_name": (o.task_name or "") if o else "",
                      "channel": chans.get(o.channel_id, "") if o else "",
                      "reason": f"已停订单，数据日期已过末次出数日（截止日期 {end_str}）" if end_str else "已停订单，数据日期已过末次出数日",
                      "qty": cnt[tid]})
    return total, items


def _count_inactive_orders(db, date_obj):
    """统计在执订单中，数据日期 >= 订单更新日期+2 但当天无数据的工单（明细列表）"""
    has_data = {r[0] for r in db.query(DailyData.task_id).filter(DailyData.biz_date == date_obj).all() if r[0]}
    chans = {c.id: c.name for c in db.query(Channel).all()}
    items = []
    for o in db.query(Order).filter(Order.status == "在执").all():
        if o.task_id and o.order_date and date_obj >= o.order_date + timedelta(days=2) and o.task_id not in has_data:
            items.append({"task_id": o.task_id, "task_name": o.task_name or "",
                          "channel": chans.get(o.channel_id, ""),
                          "reason": f"在执订单，更新日期 {o.order_date} +2 天后仍未出数据"})
    items.sort(key=lambda x: x["task_id"])
    return items


@router.post("/daily-data/import")
async def import_daily(files: List[UploadFile] = File(...), biz_date: str = Form(""),
                       confirm: bool = Form(False), db: Session = Depends(get_db), _=Depends(get_current_user)):
    date_obj = _parse_biz_date(biz_date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="请选择数据日期")

    # 读取所有文件
    file_list = []
    for file in files:
        fn = file.filename or ""
        content = await file.read()
        if content:
            file_list.append((fn, content))

    # 未匹配订单 或 已停订单异常数据时，弹出提示确认
    if not confirm:
        unmatched, unmatched_items = _count_unmatched(db, file_list)
        stopped_abnormal, stopped_items = _count_stopped_abnormal(db, file_list, date_obj)
        if unmatched > 0 or stopped_abnormal > 0:
            return {"code": 0, "data": {"needs_confirm": True, "unmatched": unmatched, "unmatched_items": unmatched_items,
                                        "stopped_abnormal": stopped_abnormal, "stopped_items": stopped_items},
                    "msg": f"有 {unmatched} 条数据未匹配到订单，{stopped_abnormal} 条关联到已停订单"}

    # 当天已有 (任务id, 手机号)，用于同任务内去重
    seen = {(r.task_id, r.phone) for r in db.query(DailyData).filter(DailyData.biz_date == date_obj).all()}

    # 预加载参考数据，避免逐行 N+1 查询
    ctx = _build_import_ctx(db)

    imported = 0
    deduped = 0
    errors = []
    for fn, content in file_list:
        ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
        if ext not in ("zip", "xlsx", "xls"):
            errors.append(f"{fn}: 仅支持 zip / xlsx / xls 文件")
            continue
        # 上传到源文件管理表
        sf = SourceFile(filename=fn, file_data=content, file_type=ext,
                        category="进货", biz_date=date_obj.strftime("%Y-%m-%d"), size=len(content))
        db.add(sf)
        db.flush()
        n, dd = _import_one_daily_file(db, fn, content, date_obj, sf.id, errors, seen, ctx)
        imported += n
        deduped += dd
    db.commit()
    inactive_items = _count_inactive_orders(db, date_obj)
    return {"code": 0, "data": {"imported": imported, "deduped": deduped, "errors": errors, "inactive_items": inactive_items},
            "msg": f"导入 {imported} 条，去重 {deduped} 条，失败 {len(errors)} 条"}


# ============ 日活数据：LM 结构 CSV 原样导入 ============
_CSV_FIELD_ALIASES = {
    "数据日期": "biz_date", "日期": "biz_date", "biz_date": "biz_date",
    "甲方": "upstream", "upstream": "upstream",
    "任务id": "task_id", "任务ID": "task_id", "工单号": "task_id", "task_id": "task_id",
    "任务名": "task_name", "task_name": "task_name",
    "手机号": "phone", "联系方式": "phone", "phone": "phone",
    "姓名": "name", "name": "name",
    "省": "province", "省份": "province", "province": "province",
    "市": "city", "城市": "city", "city": "city",
    "运营商": "operator", "operator": "operator",
    "一级品类": "cat1", "cat1": "cat1",
    "二级品类": "cat2", "cat2": "cat2",
    "平台": "platform", "platform": "platform",
    "一级代理": "customer", "customer": "customer",
    "二级代理": "secondary_agent", "secondary_agent": "secondary_agent",
    "渠道": "channel", "channel": "channel",
    "来源文件名": "source_file", "source_file": "source_file",
}


def _norm_csv_header(h):
    return re.sub(r"[\s\u3000]+", "", str(h or "")).strip().lstrip("\ufeff")


def _import_daily_csv_stream(db, fh, filename, default_date, stats):
    """流式解析 LM 结构 CSV，按行原样写入 daily_data。

    规则：不关联订单、不查洗名库、不做品类映射（原样入库）；
    按（数据日期+任务id+手机号）去重（含库内已有行，可重复导入不重复计数）。
    """
    import csv as _csv
    reader = _csv.reader(fh)
    header = next(reader, None)
    if header is None:
        return
    idx = {}
    for i, h in enumerate(header):
        key = _CSV_FIELD_ALIASES.get(_norm_csv_header(h))
        if key and key not in idx:
            idx[key] = i
    if "phone" not in idx:
        stats["errors"].append(f"{filename}: 缺少「手机号」列")
        return

    batch = []
    now = datetime.now()
    for values in reader:
        if not any((str(v).strip() if v is not None else "") for v in values):
            continue

        def g(field):
            i = idx.get(field)
            if i is None or i >= len(values):
                return ""
            v = values[i]
            return "" if v is None else str(v).strip()

        phone = g("phone")
        if not phone:
            continue
        d = _parse_biz_date(g("biz_date")) or default_date
        if not d:
            stats["skipped"] += 1
            continue
        ds = str(d)
        # 首次遇到某日期时，把库里该日期已有行按去重键载入
        if ds not in stats["loaded_dates"]:
            stats["loaded_dates"].add(ds)
            for tid, ph, pf, cu, ch, tn in db.query(
                    DailyData.task_id, DailyData.phone, DailyData.platform,
                    DailyData.customer, DailyData.channel, DailyData.task_name).filter(
                    DailyData.biz_date == d).all():
                stats["seen"].add((ds, tid or "", ph or "", pf or "", cu or "", ch or "", tn or ""))

        task_id = g("task_id")
        # 去重键：任务id 为空时（历史日活文件），用平台/一级代理/渠道/任务名 补充区分，避免误吞同日同号的不同行
        key = (ds, task_id, phone, g("platform"), g("customer"), g("channel"), g("task_name"))
        if key in stats["seen"]:
            stats["deduped"] += 1
            continue
        stats["seen"].add(key)
        batch.append({
            "biz_date": d, "upstream": g("upstream"), "task_id": task_id,
            "task_name": g("task_name"), "phone": phone, "name": g("name"),
            "province": g("province"), "city": g("city"), "operator": g("operator"),
            "cat1": g("cat1"), "cat2": g("cat2"), "platform": g("platform"),
            "customer": g("customer"), "secondary_agent": g("secondary_agent"),
            "channel": g("channel"), "source_file": g("source_file") or filename,
            "source_file_id": None, "created_at": now, "updated_at": now,
        })
        if len(batch) >= 5000:
            db.execute(insert(DailyData), batch)
            db.commit()
            stats["imported"] += len(batch)
            batch.clear()
    if batch:
        db.execute(insert(DailyData), batch)
        db.commit()
        stats["imported"] += len(batch)
        batch.clear()


@router.post("/daily-data/import-csv")
async def import_daily_csv(files: List[UploadFile] = File(...), biz_date: str = Form(""),
                           db: Session = Depends(get_db), _=Depends(require_admin)):
    """原样导入 LM 结构 CSV（表头：数据日期/甲方/任务id/任务名/手机号/姓名/省/市/运营商/…）。

    与 /daily-data/import 的区别：不做订单关联、不查洗名库，字段按 CSV 原样入库。
    数据日期以每行的「数据日期」为准，缺省用表单 biz_date 兜底。
    """
    default_date = _parse_biz_date(biz_date)
    stats = {"imported": 0, "deduped": 0, "skipped": 0, "errors": [],
             "seen": set(), "loaded_dates": set()}
    seen_files = 0
    for file in files:
        fn = file.filename or ""
        if not fn.lower().endswith(".csv"):
            stats["errors"].append(f"{fn}: 仅支持 csv 文件")
            continue
        seen_files += 1
        try:
            raw = await file.read()
            text = io.StringIO(raw.decode("utf-8-sig", errors="ignore"))
            _import_daily_csv_stream(db, text, fn, default_date, stats)
            text.close()
            del raw
        except Exception as e:  # noqa
            stats["errors"].append(f"{fn}: {e}")
    if seen_files == 0 and not stats["errors"]:
        raise HTTPException(status_code=422, detail="请选择 csv 文件")
    return {"code": 0, "data": {"imported": stats["imported"], "deduped": stats["deduped"],
                                 "skipped": stats["skipped"], "errors": stats["errors"]},
            "msg": f"CSV 原样导入 {stats['imported']} 条，去重 {stats['deduped']} 条，跳过 {stats['skipped']} 条"}


@router.post("/daily-data/batch-delete")
def batch_delete_daily(body: BatchStopBody, db: Session = Depends(get_db), _=Depends(require_admin)):
    rows = db.query(DailyData).filter(DailyData.id.in_(body.ids)).all()
    for r in rows:
        db.delete(r)
    db.commit()
    return {"code": 0, "data": {"deleted": len(rows)}, "msg": f"已删除 {len(rows)} 条"}


@router.post("/daily-data/delete-by-date")
def delete_daily_by_date(body: GenerateBody, db: Session = Depends(get_db), _=Depends(require_admin)):
    d = _parse_biz_date(body.date)
    if not d:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    n = db.query(DailyData).filter(DailyData.biz_date == d).count()
    db.query(DailyData).filter(DailyData.biz_date == d).delete()
    db.commit()
    return {"code": 0, "data": {"deleted": n}, "msg": f"已删除 {body.date} 的 {n} 条日活数据"}


@router.post("/daily-data/backfill-category")
def backfill_daily_category(db: Session = Depends(get_db), _=Depends(get_current_user)):
    """按平台重新补充缺失的一级/二级品类（品类管理补全平台映射后使用）"""
    missing_cond = or_(DailyData.cat1.is_(None), DailyData.cat1 == "",
                       DailyData.cat2.is_(None), DailyData.cat2 == "")
    platforms = sorted({p for p, in db.query(DailyData.platform).filter(
        DailyData.platform.isnot(None), DailyData.platform != "", missing_cond).distinct().all()})
    updated = 0
    for p in platforms:
        cat1, cat2 = _cat_by_platform(db, p)
        if not cat1 and not cat2:
            continue
        n1 = n2 = 0
        if cat1:
            n1 = db.query(DailyData).filter(DailyData.platform == p,
                                            or_(DailyData.cat1.is_(None), DailyData.cat1 == "")).update(
                {"cat1": cat1}, synchronize_session=False)
        if cat2:
            n2 = db.query(DailyData).filter(DailyData.platform == p,
                                            or_(DailyData.cat2.is_(None), DailyData.cat2 == "")).update(
                {"cat2": cat2}, synchronize_session=False)
        updated += max(n1, n2)
    db.commit()
    return {"code": 0, "data": {"updated": updated}, "msg": f"已补充 {updated} 条品类信息"}


@router.delete("/daily-data")
def clear_daily(db: Session = Depends(get_db), _=Depends(get_current_user)):
    n = db.query(DailyData).count()
    db.query(DailyData).delete()
    db.commit()
    return {"code": 0, "data": {"cleared": n}, "msg": f"已清空 {n} 条日活数据"}


def _phone_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


@router.get("/daily-data/wash-export")
def wash_export(date: str = "", db: Session = Depends(get_db), _=Depends(get_current_user)):
    """洗名：导出所选日期下，姓名为空且关联订单「是否加名=是」的手机号（Excel）"""
    date_obj = _parse_biz_date(date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    add_name_tids = [o.task_id for o in db.query(Order).filter(Order.add_name.is_(True)).all() if o.task_id]
    qy = db.query(DailyData).filter(
        DailyData.biz_date == date_obj,
        or_(DailyData.name.is_(None), DailyData.name == ""),
    )
    if add_name_tids:
        qy = qy.filter(DailyData.task_id.in_(add_name_tids))
    else:
        qy = qy.filter(DailyData.id < 0)
    rows = qy.order_by(DailyData.id).all()
    phones = list(dict.fromkeys(r.phone for r in rows if r.phone))
    return xlsx_response(["手机号", "姓名"], [[p, ""] for p in phones], f"洗名-{date}.xlsx", "洗名")


@router.get("/daily-data/wash-stats")
def wash_stats(date: str = "", db: Session = Depends(get_db), _=Depends(get_current_user)):
    """洗名导出统计：一共要导出、已匹配洗名库、需要去小逸洗名"""
    date_obj = _parse_biz_date(date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    add_name_tids = [o.task_id for o in db.query(Order).filter(Order.add_name.is_(True)).all() if o.task_id]
    qy = db.query(DailyData).filter(DailyData.biz_date == date_obj)
    if add_name_tids:
        qy = qy.filter(DailyData.task_id.in_(add_name_tids))
    else:
        qy = qy.filter(DailyData.id < 0)
    total = qy.count()
    matched = qy.filter(DailyData.name.isnot(None), DailyData.name != "").count()
    need_wash = total - matched
    return {"code": 0, "data": {"total": total, "matched": matched, "need_wash": need_wash}, "msg": "ok"}


def _check_repeat_days(db):
    """工单检查重复天数（配置中心可配置，默认 30）"""
    cfg = db.query(SysConfig).filter(SysConfig.key == "check_repeat_days").first()
    try:
        return int(cfg.value) if cfg and cfg.value else 30
    except (TypeError, ValueError):
        return 30


def _region_violation(o, dp, dc):
    """日活记录地区是否超出订单要求地区（违规）：True=违规"""
    provs = [p for p in re.split(r"[|｜,，;；]", o.province or "") if p.strip()]
    cities = [c for c in re.split(r"[|｜,，;；]", o.city or "") if c.strip()]
    excl_provs = [p for p in re.split(r"[|｜,，;；]", o.excl_province or "") if p.strip()]
    excl_cities = [c for c in re.split(r"[|｜,，;；]", o.excl_city or "") if c.strip()]
    dp = (dp or "").strip()
    dc = (dc or "").strip()
    if dp and dp in excl_provs:
        return True
    if dc and dc in excl_cities:
        return True
    if cities:
        return not dc or dc not in cities
    if provs and "全国" not in provs:
        return not dp or dp not in provs
    return False


def _check_daily_result(db, date_obj, upstream=None):
    """工单检查：统计所选甲方+日期的在执工单数据量（按数据量降序），未到出数日(T+2)的工单不检查"""
    up_id = None
    if upstream:
        up = db.query(Customer).filter(Customer.ctype == "upstream", Customer.status == 1,
                                        or_(Customer.name == upstream, Customer.code == upstream)).first()
        up_id = up.id if up else None
    orders_q = db.query(Order).filter(Order.status == "在执")
    if up_id:
        orders_q = orders_q.filter(Order.upstream_id == up_id)
    orders = orders_q.all()
    chans = {c.id: c.name for c in db.query(Channel).all()}
    ops = {o.id: o.name for o in db.query(Operator).all()}
    ups = {c.id: c.name for c in db.query(Customer).filter(Customer.ctype == "upstream", Customer.status == 1).all()}

    def _day_counts(day):
        c = {}
        for tid, cnt in db.query(DailyData.task_id, func.count(DailyData.id)) \
                .filter(DailyData.biz_date == day, DailyData.task_id.isnot(None)) \
                .group_by(DailyData.task_id).all():
            c[tid] = int(cnt)
        return c

    counts = _day_counts(date_obj)
    counts_prev1 = _day_counts(date_obj - timedelta(days=1))
    counts_prev2 = _day_counts(date_obj - timedelta(days=2))
    # 当天各地区记录：task_id -> [(province, city)]，用于违规地区检查
    today_regions = defaultdict(list)
    for tid, province, city in db.query(DailyData.task_id, DailyData.province, DailyData.city) \
            .filter(DailyData.biz_date == date_obj, DailyData.task_id.isnot(None),
                    DailyData.task_id != "").all():
        today_regions[tid].append((province or "", city or ""))
    # 重复出现检查：同一工单号下，指定日期的手机号在前 N 天内是否重复出现（按手机号、同工单号比对）
    repeat_days = _check_repeat_days(db)
    start = date_obj - timedelta(days=repeat_days)
    today_phones = defaultdict(set)
    for tid, phone in db.query(DailyData.task_id, DailyData.phone) \
            .filter(DailyData.biz_date == date_obj, DailyData.task_id.isnot(None),
                    DailyData.task_id != "", DailyData.phone.isnot(None), DailyData.phone != "").all():
        today_phones[tid].add(phone)
    # 前 N 天内：工单号 -> 手机号 -> 出现日期集合
    history_map = defaultdict(lambda: defaultdict(set))
    for tid, phone, bd in db.query(DailyData.task_id, DailyData.phone, DailyData.biz_date) \
            .filter(DailyData.biz_date >= start, DailyData.biz_date < date_obj,
                    DailyData.task_id.isnot(None), DailyData.task_id != "",
                    DailyData.phone.isnot(None), DailyData.phone != "").all():
        history_map[tid][phone].add(str(bd))
    result = []
    for o in orders:
        # 未到达出数日（提单日+2天）的工单不检查
        if o.order_date and date_obj < o.order_date + timedelta(days=2):
            continue
        tid = o.task_id or ""
        phones = today_phones.get(tid, set())
        hist = history_map.get(tid, {})
        dup_phones = {p for p in phones if p in hist}
        repeat_dates = sorted({bd for p in dup_phones for bd in hist.get(p, set())})
        repeat_detail = [{"phone": p, "dates": sorted(hist.get(p, set()))} for p in sorted(dup_phones)]
        count = counts.get(o.task_id, 0)
        count_prev1 = counts_prev1.get(o.task_id, 0)
        count_prev2 = counts_prev2.get(o.task_id, 0)
        # 未到出数日（提单日+2天）的前一天/前两天显示为空（/）
        if o.order_date:
            first_output = o.order_date + timedelta(days=2)
            if date_obj - timedelta(days=1) < first_output:
                count_prev1 = None
            if date_obj - timedelta(days=2) < first_output:
                count_prev2 = None
        over_qty = max(count - (o.qty or 0), 0)
        region_violation = sum(1 for dp, dc in today_regions.get(tid, []) if _region_violation(o, dp, dc))
        result.append({
            "task_id": tid,
            "task_name": o.task_name or "",
            "order_no": o.order_no or "",
            "upstream": ups.get(o.upstream_id, ""),
            "channel": chans.get(o.channel_id, ""),
            "operator": ops.get(o.operator_id, ""),
            "count": count,
            "count_prev1": count_prev1,
            "count_prev2": count_prev2,
            "over_qty": over_qty,
            "region_violation": region_violation,
            "repeat_count": len(dup_phones),
            "repeat_dates": repeat_dates,
            "repeat_detail": repeat_detail,
        })
    result.sort(key=lambda x: (-x["count"], x["task_id"]))
    return result


@router.get("/daily-data/check")
def check_daily(date: str = "", upstream: str = "", db: Session = Depends(get_db), _=Depends(get_current_user)):
    """工单检查：按甲方+日期统计在执工单数据量（按数据量降序）"""
    date_obj = _parse_biz_date(date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    return {"code": 0, "data": _check_daily_result(db, date_obj, upstream or None), "msg": "ok"}


@router.get("/daily-data/check-export")
def check_export(date: str = "", upstream: str = "", db: Session = Depends(get_db), _=Depends(get_current_user)):
    """工单检查结果导出（Excel，带样式）"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from urllib.parse import quote
    from fastapi.responses import Response

    date_obj = _parse_biz_date(date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    result = _check_daily_result(db, date_obj, upstream or None)
    headers = ["工单号", "任务名", "甲方", "渠道", "运营商", "当天", "前两天", "前一天", "超出要求数量", "违规地区", "重复数量", "重复手机号与日期"]
    rows = [[r["task_id"], r["task_name"], r["upstream"], r["channel"], r["operator"], r["count"],
             r["count_prev2"], r["count_prev1"], r["over_qty"], r["region_violation"], r["repeat_count"],
             ("；".join(f"{d['phone']}（{'、'.join(d['dates'])}）" for d in r["repeat_detail"]) if r["repeat_detail"] else "无")] for r in result]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工单检查"
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = center; c.border = border
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = border
            c.alignment = right if j in (6, 7, 8, 9, 10, 11) else center

    def disp_len(v):
        s = str(v) if v is not None else ""
        return sum(2 if ord(ch) > 127 else 1 for ch in s)
    for j, h in enumerate(headers, 1):
        max_len = disp_len(h)
        for row in rows:
            max_len = max(max_len, disp_len(row[j - 1]))
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max_len + 4

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"工单检查-{date}.xlsx"
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.post("/daily-data/export-warn")
def export_warn(body: dict, db: Session = Depends(get_db), _=Depends(get_current_user)):
    """导出工单异常明细（Excel，表头带背景色、列宽自适应）"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from urllib.parse import quote
    from fastapi.responses import Response

    items = body.get("items") or []
    headers = ["工单号", "任务名", "数量", "原因"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工单异常"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = center
    for i, item in enumerate(items, 2):
        vals = [item.get("task_id", ""), item.get("task_name", ""), item.get("qty", ""), item.get("reason", "")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = border
            if j == 3:
                c.alignment = center

    def disp_len(v):
        s = str(v) if v is not None else ""
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    for j, h in enumerate(headers, 1):
        max_len = disp_len(h)
        for item in items:
            v = [item.get("task_id", ""), item.get("task_name", ""), item.get("qty", ""), item.get("reason", "")][j - 1]
            max_len = max(max_len, disp_len(v))
        ws.column_dimensions[get_column_letter(j)].width = min(max_len + 2, 60)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('工单异常.xlsx')}"})


def _parse_wash_file(fn, content):
    """解析洗名文件，返回 [(phone, name, province, city, operator)]"""
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    rows = []
    if ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header = None
        for values in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(v).strip() if v is not None else "" for v in values]
                continue
            if not any((str(v) if v is not None else "").strip() for v in values):
                continue
            def col(*names):
                for i, h in enumerate(header):
                    hn = re.sub(r"\s+", "", h).lower()
                    for nm in names:
                        if re.sub(r"\s+", "", nm).lower() in hn:
                            return i
                return -1
            pi = col("phone", "手机号")
            ni = col("name", "姓名")
            pri = col("province", "省")
            ci = col("city", "市")
            ii = col("isp", "运营商")
            phone = _phone_str(values[pi]) if pi >= 0 and pi < len(values) else ""
            name = str(values[ni]).strip() if ni >= 0 and ni < len(values) and values[ni] is not None else ""
            province = str(values[pri]).strip() if pri >= 0 and pri < len(values) and values[pri] is not None else ""
            city = str(values[ci]).strip() if ci >= 0 and ci < len(values) and values[ci] is not None else ""
            operator = str(values[ii]).strip() if ii >= 0 and ii < len(values) and values[ii] is not None else ""
            if phone:
                rows.append((phone, name, province, city, operator))
        wb.close()
    elif ext == "csv":
        import csv as csv_mod
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv_mod.reader(io.StringIO(text))
        header = None
        for line in reader:
            if header is None:
                header = [str(v).strip() for v in line]
                continue
            if not any(v.strip() for v in line):
                continue
            def col(*names):
                for i, h in enumerate(header):
                    hn = re.sub(r"\s+", "", h).lower()
                    for nm in names:
                        if re.sub(r"\s+", "", nm).lower() in hn:
                            return i
                return -1
            pi = col("phone", "手机号")
            ni = col("name", "姓名")
            pri = col("province", "省")
            ci = col("city", "市")
            ii = col("isp", "运营商")
            phone = _phone_str(line[pi]) if pi >= 0 and pi < len(line) else ""
            name = line[ni].strip() if ni >= 0 and ni < len(line) else ""
            province = line[pri].strip() if pri >= 0 and pri < len(line) else ""
            city = line[ci].strip() if ci >= 0 and ci < len(line) else ""
            operator = line[ii].strip() if ii >= 0 and ii < len(line) else ""
            if phone:
                rows.append((phone, name, province, city, operator))
    return rows


@router.post("/daily-data/import-wash")
async def import_wash(file: UploadFile = File(...), db: Session = Depends(get_db), _=Depends(get_current_user)):
    """导入洗名：按手机号更新姓名，同时更新到洗名库（手机号存在则更新，不存在则插入）"""
    fn = file.filename or ""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="文件为空")
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    if ext not in ("xlsx", "xls", "csv"):
        raise HTTPException(status_code=422, detail="仅支持 xlsx / xls / csv 文件")
    rows = _parse_wash_file(fn, content)
    if not rows:
        raise HTTPException(status_code=422, detail="未解析到洗名数据，请检查文件表头是否包含「手机号」列")
    rec_map = {}
    for phone, name, province, city, operator in rows:
        rec_map[phone] = {"name": name, "province": province, "city": city, "operator": operator}
    phones = list(rec_map.keys())
    if phones:
        # 批量更新日活数据姓名（按手机号，走 phone 索引）
        stmt = (update(DailyData.__table__)
                .where(DailyData.__table__.c.phone == bindparam("ph"))
                .values(name=bindparam("nm")))
        db.execute(stmt, [{"ph": p, "nm": rec_map[p]["name"]} for p in phones])
    for phone in phones:
        rec = rec_map[phone]
        # 更新洗名库（手机号存在则更新，不存在且有名则插入；无姓名不插入，避免清空已有姓名）
        w = db.query(WashName).filter(WashName.phone == phone).first()
        if w:
            if rec["name"]:
                w.name = rec["name"]
            if rec["province"]:
                w.province = rec["province"]
            if rec["city"]:
                w.city = rec["city"]
            if rec["operator"]:
                w.operator = rec["operator"]
        elif rec["name"]:
            db.add(WashName(phone=phone, name=rec["name"], province=rec["province"],
                            city=rec["city"], operator=rec["operator"]))
    db.commit()
    return {"code": 0, "data": {"updated": len(phones)}, "msg": f"已更新 {len(phones)} 个手机号的姓名，洗名库同步更新"}


# ============ 洗名库 ============

@router.get("/wash-names")
def list_wash_names(db: Session = Depends(get_db), _=Depends(get_current_user),
                    q: str = "", phone: str = "", name: str = "", name_empty: str = "",
                    province: str = "", city: str = "", operator: str = "",
                    page: int = 1, per_page: int = 10):
    qy = db.query(WashName)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(WashName.phone.like(like), WashName.name.like(like)))
    if phone:
        qy = qy.filter(WashName.phone.like(f"%{phone}%"))
    if name:
        qy = qy.filter(WashName.name.like(f"%{name}%"))
    if name_empty in ("1", "true", "yes", "on"):
        qy = qy.filter(or_(WashName.name.is_(None), WashName.name == ""))
    if province:
        provs = [p for p in province.split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[WashName.province.like(f"%{p}%") for p in provs]))
    if city:
        cities = [c for c in city.split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[WashName.city.like(f"%{c}%") for c in cities]))
    if operator:
        qy = qy.filter(WashName.operator.like(f"%{operator}%"))
    total, rows = paginate(qy.order_by(WashName.id.desc()), page, per_page)
    data = [{
        "id": w.id, "phone": w.phone or "", "name": w.name or "",
        "province": w.province or "", "city": w.city or "", "operator": w.operator or "",
        "created_at": fmt_dt(w.created_at), "updated_at": fmt_dt(w.updated_at),
    } for w in rows]
    return ok_page(data, total)


def _map_wash_row(header, values):
    row = {}
    for i, h in enumerate(header):
        if i >= len(values):
            continue
        hn = re.sub(r"\s+", "", str(h))
        v = values[i]
        sv = None if v is None else str(v).strip()
        if not sv:
            continue
        if "手机号" in hn:
            row["phone"] = sv
        elif "姓名" in hn:
            row["name"] = sv
        elif "省" in hn:
            row["province"] = sv
        elif "市" in hn:
            row["city"] = sv
        elif "运营商" in hn:
            row["operator"] = sv
    return row


def _parse_wash_csv(content):
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    if not header:
        return []
    return [_map_wash_row(header, values) for values in reader]


def _parse_wash_xlsx(content):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    header = None
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = values
            continue
        rows.append(_map_wash_row(header, values))
    wb.close()
    return rows


@router.post("/wash-names/import")
async def import_wash_names(file: UploadFile = File(...), db: Session = Depends(get_db), _=Depends(get_current_user)):
    fn = file.filename or ""
    content = await file.read()
    if fn.lower().endswith(".csv"):
        rows = _parse_wash_csv(content)
    elif fn.lower().endswith((".xlsx", ".xls")):
        rows = _parse_wash_xlsx(content)
    else:
        raise HTTPException(status_code=422, detail="仅支持 csv / xlsx 文件")
    imported = 0
    updated = 0
    for r in rows:
        phone = r.get("phone")
        if not phone:
            continue
        w = db.query(WashName).filter(WashName.phone == phone).first()
        if w:
            if r.get("name"):
                w.name = r["name"]
            if r.get("province"):
                w.province = r["province"]
            if r.get("city"):
                w.city = r["city"]
            if r.get("operator"):
                w.operator = r["operator"]
            updated += 1
        else:
            if not r.get("name"):
                continue
            db.add(WashName(phone=phone, name=r.get("name"), province=r.get("province"),
                            city=r.get("city"), operator=r.get("operator")))
            imported += 1
    db.commit()
    return {"code": 0, "data": {"imported": imported, "updated": updated},
            "msg": f"导入 {imported} 条，更新 {updated} 条"}


@router.post("/wash-names/batch-delete")
def batch_delete_wash_names(body: BatchStopBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    rows = db.query(WashName).filter(WashName.id.in_(body.ids)).all()
    for r in rows:
        db.delete(r)
    db.commit()
    return {"code": 0, "data": {"deleted": len(rows)}, "msg": f"已删除 {len(rows)} 条洗名数据"}


@router.delete("/wash-names")
def clear_wash_names(db: Session = Depends(get_db), _=Depends(get_current_user)):
    n = db.query(WashName).count()
    db.query(WashName).delete()
    db.commit()
    return {"code": 0, "data": {"cleared": n}, "msg": f"已清空 {n} 条洗名库数据"}


def _task_suffix(task_name):
    """任务名只保留尾缀（最后一个 - 之后的内容）"""
    tn = (task_name or "").strip()
    return tn.rsplit("-", 1)[-1] if "-" in tn else tn


def _strip_task_prefix(task_name, city, province):
    """去掉任务名中的「代理编号+代理名+地区」前缀，返回实际任务名"""
    tn = (task_name or "").strip()
    parts = [p for p in tn.split("-") if p.strip()]
    if not parts:
        return tn
    idx = 1 if parts[0].isdigit() else 0
    region = (city or "").strip() or (province or "").strip()
    if region:
        for i in range(idx, len(parts)):
            if parts[i] == region:
                idx = i + 1
                break
    return "-".join(parts[idx:]) if idx < len(parts) else tn


def _pinyin_key(s):
    """把中文转成拼音（用于排序）"""
    try:
        from pypinyin import lazy_pinyin
        return "".join(lazy_pinyin(str(s or "")))
    except Exception:  # noqa
        return str(s or "")


def _fill_mb030_file(db, tpl_bytes, records):
    """MB-030 产品出数模版：数据代码/本次申请平台/联系方式/客户名称/本次状态/近期申请平台A/B/C。
    数据代码 = LeMon_ + 二级品类编号 + 四位平台编号；数据按数据代码排序；数据单元格绿色背景+边框。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
    wb = load_workbook(io.BytesIO(tpl_bytes))
    ws = wb.active
    green = PatternFill(start_color="408110", end_color="408110", fill_type="solid")
    black_font = Font(color="000000")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    plat_by_name = {p.name: p for p in db.query(Platform).all() if p.name}
    cat_by_id = {c.id: c for c in db.query(Category).all()}

    def platform_codes(name):
        p = plat_by_name.get((name or "").strip())
        if not p:
            return "", ""
        cat = cat_by_id.get(p.cat_id)
        return (p.code or ""), (cat.code if cat else "")

    rows = []
    for rec in records:
        pcode, ccode = platform_codes(rec.get("platform") or "")
        data_code = "LeMon_" + ccode + pcode
        hist = [h for h in (rec.get("history") or []) if h and h != "无"]
        row = [
            data_code,
            rec.get("platform") or "",
            rec.get("url") or "",
            rec.get("name") or "-",
            "申请成功",
            hist[0] if len(hist) > 0 else "-",
            hist[1] if len(hist) > 1 else "-",
            hist[2] if len(hist) > 2 else "-",
        ]
        rows.append((data_code, row))
    rows.sort(key=lambda x: x[0])

    header_row = 1
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)

    r = header_row + 1
    for _, row in rows:
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.fill = green
            cell.border = border
            cell.alignment = center
            cell.font = black_font
        r += 1

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _fill_distribute_file(db, tpl_bytes, tpl_code, records):
    """用模版字节 + 日活数据记录（含手机号/省市/运营商/渠道/平台/任务名）生成单个文件；
    若存在「是否撞库」的记录，则在末尾追加「近期申请记录」列（该类记录填历史，无历史填「无」，其余为空）。
    tpl_code == MB-030 时走产品出数专用填充。"""
    if tpl_code == "MB-030":
        return _fill_mb030_file(db, tpl_bytes, records)
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from ..tidabiao import _header_field, _cell_value_split
    has_collision = any(rec.get("check_collision") for rec in records)
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    if tpl_bytes:
        wb = load_workbook(io.BytesIO(tpl_bytes))
        ws = wb.active
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_row = 1
        headers = [c.value for c in ws[header_row]]
        if all(h is None or str(h).strip() == "" for h in headers):
            header_row = 2
            headers = [c.value for c in ws[header_row]]
        colmap = [(j, _header_field(h) if h else None) for j, h in enumerate(headers, 1)]
        hist_col = len(headers) + 1
        if ws.max_row > header_row:
            ws.delete_rows(header_row + 1, ws.max_row - header_row)
        if has_collision:
            max_w = sum(2 if ord(ch) > 127 else 1 for ch in "近期申请记录")
            for rec in records:
                for line in (rec.get("history") or []):
                    max_w = max(max_w, sum(2 if ord(ch) > 127 else 1 for ch in str(line)))
            hc = ws.cell(row=header_row, column=hist_col, value="近期申请记录")
            hc.alignment = center
            hc.border = border
            hc.font = Font(bold=True)
            hc.fill = green
            ws.column_dimensions[get_column_letter(hist_col)].width = min(max_w + 2, 60)
        row = header_row + 1
        for rec in records:
            for j, m in colmap:
                if not m:
                    continue
                kind, key = m
                value = _cell_value_split(kind, key, rec, rec.get("url", ""))
                cell = ws.cell(row=row, column=j, value=value)
                cell.alignment = center
                cell.border = border
            if has_collision:
                hist = "\n".join(rec.get("history") or [])
                hcell = ws.cell(row=row, column=hist_col, value=hist)
                hcell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                hcell.border = border
            row += 1
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["手机号"] + (["近期申请记录"] if has_collision else []))
        for rec in records:
            ws.append([rec.get("url", "")] + (["\n".join(rec.get("history") or [])] if has_collision else []))
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _fill_bill_file(cust_label, date_str, bill, tpl_bytes=None):
    """生成单个代理的账单 excel；配置了账单模版则按模版填充，否则用默认格式"""
    if tpl_bytes:
        return _fill_bill_file_from_template(tpl_bytes, cust_label, date_str, bill)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "账单"
    qty = bill["qty"]
    sales = round(float(bill["sales"]), 2)
    unit = round(sales / qty, 4) if qty else 0
    rows = [
        ["代理", cust_label],
        ["业务日期", date_str],
        ["进货量", qty],
        ["单价(元/条)", unit],
        ["销售金额(元)", sales],
        ["余额", round(float(bill.get("balance") or 0), 2)],
    ]
    recharge_info = bill.get("recharge_info") or []
    if recharge_info:
        rows.append(["充值/补数", "；".join(recharge_info)])
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["任务名", "手机号数量", "金额(元)"])
    for g, gg in sorted(bill.get("groups", {}).items(),
                        key=lambda kv: _pinyin_key(kv[1].get("task_name") or kv[0])):
        ws.append([gg.get("task_name") or g, gg["qty"], round(float(gg["amount"]), 2)])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _copy_cell_style(src, dst):
    """把源单元格的字体/边框/填充/对齐/数字格式复制到目标单元格"""
    import copy as _copy
    if src.has_style:
        dst.font = _copy.copy(src.font)
        dst.border = _copy.copy(src.border)
        dst.fill = _copy.copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = _copy.copy(src.protection)
        dst.alignment = _copy.copy(src.alignment)


def _fill_bill_file_from_template(tpl_bytes, cust_label, date_str, bill):
    """按账单模版结构填充：
    汇总区（客户/业务日期/进货量/销售金额/余额）填右侧值单元格；
    明细表（任务名/数量/单价/是否洗名/金额）按表头定位列并填充数据行，
    当模版空白行数不足时自动插行并复制样式，不影响整体版式。
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(tpl_bytes))
    ws = wb.active

    qty = bill["qty"]
    sales = round(float(bill["sales"]), 2)
    balance = bill.get("balance")
    groups = sorted(bill.get("groups", {}).items(),
                    key=lambda kv: _pinyin_key(kv[1].get("task_name") or kv[0]))

    # ---- 1. 汇总区：标签 -> 右侧值单元格 ----
    def _summary_field(s):
        if s in ("客户", "客户名称", "代理", "代理名称", "客户名"):
            return "客户"
        if "日期" in s:
            return "日期"
        if "进货" in s:
            return "进货量"
        if "销售" in s:
            return "销售金额"
        if "余额" in s:
            return "余额"
        return None

    summary_vals = {
        "客户": cust_label,
        "日期": date_str,
        "进货量": qty,
        "销售金额": sales,
        "余额": (round(float(balance), 2) if balance is not None else ""),
    }
    # 值单元格：标签右侧的第一个合并区间左上角（如 E 列），无合并则用相邻列
    merged_by_row = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col > 1:
            merged_by_row.setdefault(rng.min_row, rng.min_col)
    balance_row = None
    balance_label_col = None
    balance_val_col = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            field = _summary_field(str(cell.value).strip())
            if field:
                target_col = merged_by_row.get(cell.row, cell.column + 1)
                ws.cell(row=cell.row, column=target_col, value=summary_vals[field])
                if field == "余额":
                    balance_row = cell.row
                    balance_label_col = cell.column
                    balance_val_col = target_col

    # ---- 余额下追加充值/补数信息 ----
    recharge_info = bill.get("recharge_info") or []
    if recharge_info and balance_row is not None:
        insert_at = balance_row + 1
        # 余额值单元格的合并列范围
        val_min = balance_val_col
        val_max = balance_val_col
        for rng in ws.merged_cells.ranges:
            if rng.min_row == balance_row and rng.min_col == balance_val_col:
                val_max = rng.max_col
                break
        # 记录插入点下方合并区间与行高（openpyxl insert_rows 不移动它们，需手动下移）
        shifted = [rng for rng in list(ws.merged_cells.ranges) if rng.min_row >= insert_at]
        heights = {}
        for r in range(insert_at, ws.max_row + 2):
            h = ws.row_dimensions[r].height
            if h is not None:
                heights[r] = h
        for rng in shifted:
            ws.unmerge_cells(str(rng))
        ws.insert_rows(insert_at)
        for rng in shifted:
            ws.merge_cells(start_row=rng.min_row + 1, start_column=rng.min_col,
                           end_row=rng.max_row + 1, end_column=rng.max_col)
        for r in sorted(heights.keys(), reverse=True):
            ws.row_dimensions[r + 1].height = heights[r]
        ws.row_dimensions[insert_at].height = None
        # 复制余额整行（值+样式），再填充新内容
        for c in range(1, ws.max_column + 1):
            src = ws.cell(row=balance_row, column=c)
            dst = ws.cell(row=insert_at, column=c)
            dst.value = src.value
            _copy_cell_style(src, dst)
        src_height = ws.row_dimensions[balance_row].height
        if src_height is not None:
            ws.row_dimensions[insert_at].height = src_height
        ws.cell(row=insert_at, column=balance_label_col, value="充值/补数")
        ws.cell(row=insert_at, column=val_min, value="；".join(recharge_info))
        if val_max > val_min:
            ws.merge_cells(start_row=insert_at, start_column=val_min,
                           end_row=insert_at, end_column=val_max)

    # ---- 2. 明细表：定位表头行与各列 ----
    header_row = None
    cols = {}
    for row in ws.iter_rows():
        rowvals = {}
        for cell in row:
            if cell.value is not None:
                rowvals[cell.column] = str(cell.value).strip()
        has_task = any(("任务" in v or "小组" in v or "工单" in v) for v in rowvals.values())
        has_qty = any("数量" in v for v in rowvals.values())
        has_amount = any("金额" in v for v in rowvals.values())
        if has_task and has_qty and has_amount:
            header_row = row[0].row
            for col, v in rowvals.items():
                if "任务" in v or "小组" in v:
                    cols["task"] = col
                elif "工单" in v or "任务id" in v or "订单编号" in v or "订单号" in v:
                    cols["task_id"] = col
                elif "类型" in v or "渠道" in v:
                    cols["channel"] = col
                elif "运营商" in v:
                    cols["operator"] = col
                elif "单价" in v or "价格" in v:
                    cols["price"] = col
                elif "洗名" in v or "加名" in v:
                    cols["name"] = col
                elif "数量" in v:
                    cols["qty"] = col
                elif "金额" in v:
                    cols["amount"] = col
            break

    if header_row is None:
        return _fill_bill_file(cust_label, date_str, bill, None)

    # 找合计行（表头下方第一个含「合计/总计」的行）
    total_row = None
    for r in range(header_row + 1, ws.max_row + 2):
        for c in range(1, ws.max_column + 2):
            v = ws.cell(row=r, column=c).value
            if v is not None and ("合计" in str(v) or "总计" in str(v)):
                total_row = r
                break
        if total_row:
            break

    data_start = header_row + 1
    data_end = (total_row - 1) if total_row else ws.max_row
    n_blank = max(0, data_end - data_start + 1)
    n_groups = len(groups)

    # 空白行不足时自动插行（插在合计行之前，或末尾）
    if n_groups > n_blank:
        insert_at = total_row if total_row else (data_end + 1)
        ws.insert_rows(insert_at, amount=n_groups - n_blank)

    # 参考样式行：优先第一条数据行（空白行自带边框样式），否则用表头行
    ref_row = data_start if n_blank >= 1 else header_row
    ref_cols = [c for c in cols.values() if c is not None]

    # 斑马纹：识别模版已有数据行中「填充行」的奇偶性与填充样式，用于还原新插入行的斑马纹
    import copy as _copy
    striped_fill = None
    striped_parity = None
    for i in range(n_blank):
        c = ws.cell(row=data_start + i, column=ref_cols[0])
        if c.fill and c.fill.patternType == 'solid':
            striped_fill = _copy.copy(c.fill)
            striped_parity = i % 2
            break

    for i, (g, gg) in enumerate(groups):
        r = data_start + i
        if "task" in cols:
            ws.cell(row=r, column=cols["task"], value=(gg.get("task_name") or g))
        if "task_id" in cols:
            ws.cell(row=r, column=cols["task_id"], value=(gg.get("order_no") or g))
        if "channel" in cols:
            ws.cell(row=r, column=cols["channel"], value=(gg.get("channel") or ""))
        if "operator" in cols:
            ws.cell(row=r, column=cols["operator"], value=(gg.get("operator") or ""))
        if "qty" in cols:
            ws.cell(row=r, column=cols["qty"], value=gg["qty"])
        if "price" in cols:
            ws.cell(row=r, column=cols["price"], value=(gg.get("unit") or 0))
        if "name" in cols:
            ws.cell(row=r, column=cols["name"], value=("/" if gg.get("is_product") else ("是" if gg.get("add_name") else "否")))
        if "amount" in cols:
            ws.cell(row=r, column=cols["amount"], value=round(float(gg["amount"]), 2))
        # 仅对「新插入」的行复制样式（原有空白行保留模版本来的样式）
        if i >= n_blank:
            for c in ref_cols:
                _copy_cell_style(ws.cell(row=ref_row, column=c), ws.cell(row=r, column=c))
            # 还原斑马纹：与模版填充行同奇偶的新行补上填充色
            if striped_fill is not None and (i % 2) == striped_parity:
                for c in ref_cols:
                    ws.cell(row=r, column=c).fill = _copy.copy(striped_fill)

    # 修正合计行的 SUM 公式范围（openpyxl 插行不会自动调整公式范围）
    if n_groups > n_blank and total_row:
        from openpyxl.utils import get_column_letter
        new_total_row = total_row + (n_groups - n_blank)
        end_row = data_start + n_groups - 1
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=new_total_row, column=c)
            if cell.value is not None and "SUM(" in str(cell.value):
                letter = get_column_letter(c)
                cell.value = f"=SUM({letter}{data_start}:{letter}{end_row})"

    # 任务名 / 订单编号列宽度自适应内容（中文按 2 个字符宽度计）
    from openpyxl.utils import get_column_letter as _gcl
    for key, attr in [("task", "task_name"), ("task_id", "order_no")]:
        if key not in cols:
            continue
        max_w = 0
        for i, (g, gg) in enumerate(groups):
            val = str(gg.get(attr) or g)
            w = sum(2 if ord(ch) > 127 else 1 for ch in val)
            max_w = max(max_w, w)
        if max_w > 0:
            ws.column_dimensions[_gcl(cols[key])].width = min(max_w + 2, 60)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _attach_history(db, jobs, date_obj):
    """给「是否撞库=是」的记录补充「近期申请记录」：手机号在日活历史中（早于本次数据日期）
    最近三次出现，只匹配相同一级品类且非当前平台的历史（日期+平台）。"""
    keys = set()
    for job in jobs:
        for rec in job["records"]:
            if rec.get("check_collision") and rec.get("url"):
                keys.add((rec["url"], rec.get("cat1") or "", rec.get("platform") or ""))
    keys = sorted(keys)
    lookup = {}
    for k in keys:
        lookup.setdefault((k[0], k[1]), []).append(k)
    history_map = {}
    CHUNK = 500
    phones = sorted({k[0] for k in keys})
    for i in range(0, len(phones), CHUNK):
        chunk_phones = phones[i:i + CHUNK]
        rows = db.query(DailyData.phone, DailyData.cat1, DailyData.biz_date, DailyData.platform).filter(
            DailyData.phone.in_(chunk_phones), DailyData.biz_date < date_obj
        ).order_by(DailyData.phone, DailyData.biz_date.desc()).all()
        for phone, cat1, bd, pf in rows:
            for key in lookup.get((phone, cat1 or ""), []):
                if key[2] and pf == key[2]:
                    continue
                lst = history_map.setdefault(key, [])
                if any(x[0] == bd for x in lst):
                    continue
                if len(lst) < 3:
                    lst.append((bd, pf or ""))
    for job in jobs:
        for rec in job["records"]:
            if not rec.get("check_collision"):
                continue
            key = (rec.get("url"), rec.get("cat1") or "", rec.get("platform") or "")
            lst = history_map.get(key, [])
            rec["history"] = [f"{bd.strftime('%Y%m%d')} {pf or '未知'}" for bd, pf in lst] or ["无"]


def _build_distribute_zip(db, jobs, date_str, progress_cb, should_cancel=None):
    import zipfile
    buf = io.BytesIO()
    used_names = set()
    folder_totals = {}
    for job in jobs:
        folder_totals[job["cust_code"]] = folder_totals.get(job["cust_code"], 0) + len(job["records"])
    folder_of = {cc: f"{cc}-{folder_totals[cc]}" for cc in folder_totals}
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, job in enumerate(jobs):
            if should_cancel and should_cancel():
                raise DistributeCancelled()
            parts = [date_str, job["cust_code"]]
            if job["group_name"]:
                parts.append(job["group_name"])
            parts.append(str(len(job["records"])))
            base = "-".join(parts)
            filename = base + ".xlsx"
            n = 1
            while filename in used_names:
                n += 1
                filename = f"{base}-{n}.xlsx"
            used_names.add(filename)
            folder = folder_of[job["cust_code"]]
            bio = _fill_distribute_file(db, job["tpl_bytes"], job.get("tpl_code") or "", job["records"])
            zf.writestr(f"{folder}/{filename}", bio.read())
            if progress_cb:
                progress_cb(i + 1)
    buf.seek(0)
    return buf


def _build_daily_bills(db, date_obj, customer_ids=None):
    """构建某日期各代理账单明细（不落库）：返回 {cust_code: {customer_id, cust_label, qty, sales, cost, balance, groups, recharge_info}}"""
    rows = db.query(DailyData.task_id, func.count(DailyData.id)).filter(
        DailyData.biz_date == date_obj).group_by(DailyData.task_id).all()
    tids = {tid for tid, _ in rows if tid}
    orders_by_tid = {o.task_id: o for o in db.query(Order).filter(
        Order.task_id.in_(tids)).all()} if tids else {}
    needed_cids = {o.customer_id for o in orders_by_tid.values()}
    needed_chids = {o.channel_id for o in orders_by_tid.values() if o.channel_id}
    needed_oids = {o.operator_id for o in orders_by_tid.values() if o.operator_id}
    cust_map = {c.id: c for c in db.query(Customer).filter(
        Customer.id.in_(needed_cids)).all()} if needed_cids else {}
    ch_map = {c.id: c for c in db.query(Channel).filter(
        Channel.id.in_(needed_chids)).all()} if needed_chids else {}
    op_map = {o.id: o for o in db.query(Operator).filter(
        Operator.id.in_(needed_oids)).all()} if needed_oids else {}
    # 预加载甲方价格（避免逐订单 N+1 查询）
    needed_upids = {o.upstream_id for o in orders_by_tid.values() if o.upstream_id}
    up_price_map = defaultdict(dict)
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id.in_(needed_upids)).all():
        up_price_map[p.customer_id][(p.channel_id, p.operator_id)] = float(p.price)
    orders_info = {}
    for order in orders_by_tid.values():
        tid = order.task_id
        cust = cust_map.get(order.customer_id)
        ch = ch_map.get(order.channel_id)
        op = op_map.get(order.operator_id)
        up_price = 0.0
        if order.upstream_id and order.channel_id:
            pm = up_price_map.get(order.upstream_id, {})
            up_price = pm.get((order.channel_id, order.operator_id))
            if up_price is None:
                up_price = pm.get((order.channel_id, None), 0.0)
        orders_info[tid] = {
            "group_name": order.group_name or "", "task_name": order.task_name or "",
            "order_no": order.order_no or "", "channel": ch.name if ch else "",
            "operator": op.name if op else "", "cust_code": cust.code if cust else "",
            "cust_name": cust.name if cust else "", "price": float(order.price) if order.price is not None else 0,
            "customer_id": order.customer_id, "add_name": bool(order.add_name),
            "upstream_price": up_price, "end_date": order.end_date,
        }
    selected = set(customer_ids) if customer_ids else None
    expired_tids = {tid for tid, info in orders_info.items()
                    if info.get("end_date") and date_obj > info["end_date"] + timedelta(days=1)}
    bills = {}
    for tid, cnt in rows:
        cnt = int(cnt or 0)
        if not tid or tid not in orders_info or tid in expired_tids:
            continue
        info = orders_info[tid]
        if selected is not None and info.get("customer_id") not in selected:
            continue
        if (info.get("price") or 0) <= 0:
            continue  # 单价为 0 的订单不生成账单/明细
        cc = info["cust_code"]
        if cc not in bills:
            label = (f"{cc} {info['cust_name']}".strip() if info["cust_name"] else cc)
            cust = cust_map.get(info["customer_id"])
            bills[cc] = {"customer_id": info["customer_id"], "cust_label": label,
                         "qty": 0, "sales": 0.0, "cost": 0.0,
                         "balance": float(cust.balance) if cust and cust.balance is not None else 0,
                         "groups": {}}
        eff_price = info["price"] + (0.01 if info["add_name"] else 0)
        bills[cc]["qty"] += cnt
        bills[cc]["sales"] += eff_price * cnt
        bills[cc]["cost"] += info["upstream_price"] * cnt
        gg = bills[cc]["groups"].setdefault(tid,
                                            {"task_name": info["task_name"] or info["group_name"],
                                             "order_no": info["order_no"], "channel": info["channel"],
                                             "operator": info["operator"], "qty": 0, "amount": 0.0,
                                             "unit": info["price"], "add_name": info["add_name"]})
        gg["qty"] += cnt
        gg["amount"] += eff_price * cnt
    # 产品销售合并
    for ps, prod in db.query(ProductSale, Product).join(
            Product, ProductSale.product_id == Product.id).filter(
            ProductSale.sale_date == date_obj,
            ProductSale.deleted_at.is_(None),
            Product.deleted_at.is_(None)).all():
        cid = ps.customer_id
        if not cid:
            continue
        if selected is not None and cid not in selected:
            continue
        cc = next((_cc for _cc, b in bills.items() if b.get("customer_id") == cid), None)
        if cc is None:
            cust = db.query(Customer).filter(Customer.id == cid).first()
            label = (f"{cust.code} {cust.name}".strip() if cust and cust.name else (cust.code or str(cid)))
            cc = cust.code if cust and cust.code else str(cid)
            bills[cc] = {"customer_id": cid, "cust_label": label, "qty": 0, "sales": 0.0,
                         "cost": 0.0, "balance": float(cust.balance) if cust and cust.balance is not None else 0,
                         "groups": {}}
        b = bills[cc]
        pqty = int(ps.qty or 0)
        psales = float(ps.sales_amount or 0)
        b["qty"] += pqty
        b["sales"] += psales
        b["cost"] += round(pqty * float(ps.purchase_price or 0), 2)
        key = f"product_{ps.id}"
        b["groups"][key] = {"task_name": prod.name, "order_no": "/", "channel": "", "operator": "",
                            "qty": pqty, "amount": psales, "unit": float(ps.sales_price or 0),
                            "add_name": None, "is_product": True}
    # 充值信息
    recharge_map = defaultdict(list)
    for rec in db.query(CustomerRecharge).filter(CustomerRecharge.recharge_date == date_obj).all():
        tlabel = "补数" if (rec.recharge_type or "转账") == "补数" else "充值"
        amt = float(rec.amount_rmb or 0)
        reason = (rec.reason or "").strip()
        s = f"{tlabel} {amt:g}元"
        if reason:
            s += f"（{reason}）"
        recharge_map[rec.customer_id].append(s)
    for cc, b in bills.items():
        b["recharge_info"] = recharge_map.get(b.get("customer_id"), [])
    # 余额：优先用账单表里当日已保存的余额（与当月账单一致），避免用当前余额重复扣减销售金额
    saved_balance = {x.customer_id: x.balance for x in db.query(Bill).filter(Bill.biz_date == date_obj, Bill.deleted_at.is_(None)).all()}
    for cc, b in bills.items():
        cid = b.get("customer_id")
        if cid is not None and cid in saved_balance and saved_balance[cid] is not None:
            b["balance"] = round(float(saved_balance[cid]), 2)
        else:
            b["balance"] = round(float(b.get("balance") or 0) - float(b["sales"]), 2)
    return bills


def _regen_downstream_bills(db, date_obj, customer_ids, user_label=""):
    """重新生成指定客户在指定日期的下游账单（分数+产品销售合并），并同步客户余额"""
    ids = list(customer_ids) or []
    bills = _build_daily_bills(db, date_obj, ids)
    handled = set()
    for cc, b in bills.items():
        cid = b.get("customer_id")
        if not cid:
            continue
        handled.add(cid)
        cust = db.query(Customer).filter(Customer.id == cid).first()
        existing = db.query(Bill).filter(Bill.customer_id == cid, Bill.biz_date == date_obj, Bill.deleted_at.is_(None)).first()
        balance = float(cust.balance) if cust and cust.balance is not None else 0
        # 重复生成时：把上次已扣的销售金额加回来，恢复到扣减前的基础余额，避免重复扣减
        if existing and existing.sales is not None:
            balance = round(balance + float(existing.sales), 2)
        bill_balance = round(balance - float(b["sales"]), 2)
        b["balance"] = bill_balance
        profit = round(float(b["sales"]) - float(b["cost"]), 2)
        detail_json = json.dumps({"groups": b.get("groups", {}),
                                  "recharge_info": b.get("recharge_info", [])},
                                 ensure_ascii=False)
        if existing:
            existing.purchase_qty = b["qty"]
            existing.sales = b["sales"]
            existing.balance = bill_balance
            existing.profit = profit
            existing.detail_json = detail_json
            if user_label:
                existing.updated_by = user_label
        else:
            db.add(Bill(customer_id=cid, biz_date=date_obj, purchase_qty=b["qty"],
                        sales=b["sales"], balance=bill_balance, profit=profit,
                        detail_json=detail_json,
                        created_by=user_label, updated_by=user_label))
        if cust is not None:
            cust.balance = bill_balance
    # 当日已无任何数据（分数+产品销售）的客户：清除其账单并还原余额
    for cid in ids:
        if cid in handled:
            continue
        existing = db.query(Bill).filter(Bill.customer_id == cid, Bill.biz_date == date_obj, Bill.deleted_at.is_(None)).first()
        if existing:
            cust = db.query(Customer).filter(Customer.id == cid).first()
            if cust is not None and existing.sales is not None:
                cust.balance = round(float(cust.balance or 0) + float(existing.sales), 2)
            existing.deleted_at = datetime.now()
    db.commit()


@router.post("/daily-data/distribute")
def distribute_daily(body: GenerateBody, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """分数据：按小组名分组，用出数模版生成文件，按一级代理建文件夹"""
    if not body.date:
        raise HTTPException(status_code=422, detail="请选择数据日期")
    try:
        date_obj = datetime.strptime(body.date, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=422, detail="日期格式错误")

    rows = db.query(DailyData).filter(DailyData.biz_date == date_obj).all()
    if not rows:
        raise HTTPException(status_code=404, detail=f"数据日期 {body.date} 没有可分发的数据")

    # 预加载参考数据，消除逐工单 N+1 查询
    task_ids = {d.task_id for d in rows if d.task_id}
    orders_by_tid = {o.task_id: o for o in db.query(Order).filter(Order.task_id.in_(task_ids)).all()}
    customers = {c.id: c for c in db.query(Customer).all()}
    channels = {c.id: c for c in db.query(Channel).all()}
    operators = {o.id: o for o in db.query(Operator).all()}
    templates = {t.id: t for t in db.query(Template).all()}
    prices = {}
    for p in db.query(CustomerPrice).all():
        prices[(p.customer_id, p.channel_id, p.operator_id)] = float(p.price)

    # 1. 先把分数据对象（日活数据）全部与订单数据关联（按任务id）
    orders_info = {}
    for tid in task_ids:
        order = orders_by_tid.get(tid)
        if not order:
            continue
        cust = customers.get(order.customer_id)
        tpl = templates.get(order.tpl_id) if order.tpl_id else None
        ch = channels.get(order.channel_id)
        op = operators.get(order.operator_id)
        # 上游单价（成本）：客户管理里「上游」客户按渠道（dpi 渠道再按运营商）设定的单价
        up_price = 0.0
        if order.upstream_id and order.channel_id:
            v = prices.get((order.upstream_id, order.channel_id, order.operator_id))
            if v is None:
                v = prices.get((order.upstream_id, order.channel_id, None))
            up_price = v or 0.0
        orders_info[tid] = {
            "tpl_code": tpl.code if tpl else "",
            "group_name": order.group_name or "",
            "task_name": order.task_name or "",
            "order_no": order.order_no or "",
            "stop_date": order.stop_date,
            "end_date": order.end_date,
            "city": order.city or "",
            "province": order.province or "",
            "channel": ch.name if ch else "",
            "operator": op.name if op else "",
            "cust_code": cust.code if cust else "",
            "cust_name": cust.name if cust else "",
            "sec_agent": order.secondary_agent or "",
            "tpl_bytes": bytes(tpl.file_data) if tpl and tpl.file_data else None,
            "price": float(order.price) if order.price is not None else 0,
            "customer_id": order.customer_id,
            "add_name": bool(order.add_name),
            "check_collision": bool(order.check_collision),
            "upstream_price": up_price,
        }

    if not orders_info:
        raise HTTPException(status_code=404, detail="没有可分发的数据（任务id未关联到订单）")

    # 指定一级代理时，只分发这些代理的数据
    if body.customer_ids:
        selected = set(body.customer_ids)
        orders_info = {tid: info for tid, info in orders_info.items() if info.get("customer_id") in selected}
        if not orders_info:
            raise HTTPException(status_code=404, detail="所选一级代理没有可分发的数据")

    # 仅分发指定「截止日期」的订单（补发场景：如 10-14 比 10-15 多出的截止日期=10-13 订单）
    if body.end_date:
        try:
            end_obj = datetime.strptime(body.end_date, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(status_code=422, detail="截止日期格式错误")
        orders_info = {tid: info for tid, info in orders_info.items() if info.get("end_date") == end_obj}
        if not orders_info:
            raise HTTPException(status_code=404, detail=f"没有「截止日期={body.end_date}」的可分发订单")

    # 截止日期过了 T+1（末次出数日期 end_date+1）之后仍在出数的数据：不分发、不记账
    expired_tids = {tid for tid, info in orders_info.items()
                    if info.get("end_date") and date_obj > info["end_date"] + timedelta(days=1)}

    # 2. 校验：小组名/一级代理 不能为空
    for tid, info in orders_info.items():
        if not info["group_name"]:
            raise HTTPException(status_code=400, detail=f"订单「{tid}」的小组名为空，无法分数据")
        if not info["cust_code"]:
            raise HTTPException(status_code=400, detail=f"订单「{tid}」的一级代理为空，无法分数据")

    # 3. 按小组聚合，校验出数模版/一级代理一致性
    group_meta = {}
    for tid, info in orders_info.items():
        g = info["group_name"]
        m = group_meta.setdefault(g, {"tpl": set(), "cust": set(), "set_tid": None})
        m["tpl"].add(info["tpl_code"])
        m["cust"].add(info["cust_code"])
        if info["tpl_code"] and m["set_tid"] is None:
            m["set_tid"] = tid

    soft_issues = []
    for g, m in group_meta.items():
        set_tpls = {t for t in m["tpl"] if t}
        if len(set_tpls) > 1:
            raise HTTPException(status_code=400, detail=f"小组「{g}」对应多个出数模版：{'、'.join(sorted(set_tpls))}")
        if not set_tpls:
            raise HTTPException(status_code=400, detail=f"小组「{g}」的出数模版为空，无法分数据")
        if len(m["cust"]) > 1:
            raise HTTPException(status_code=400, detail=f"小组「{g}」对应多个一级代理：{'、'.join(sorted(m['cust']))}")
        if "" in m["tpl"]:
            unset_count = sum(1 for info in orders_info.values() if info["group_name"] == g and not info["tpl_code"])
            soft_issues.append({"group_name": g, "tpl_code": sorted(set_tpls)[0], "unset_count": unset_count})

    # 4. 有未设定出数模版的订单时，提示是否继续
    if soft_issues and not body.confirm:
        detail = "、".join(f"「{s['group_name']}」{s['unset_count']}个订单未设定出数模版" for s in soft_issues)
        return {"code": 0, "data": {"needs_confirm": True, "issues": soft_issues},
                "msg": f"以下小组存在未设定出数模版的订单：{detail}"}

    # 5. 按小组名分组（同一小组的数据合并到一个文件，使用小组已设定的出数模版）
    jobs = {}
    for d in rows:
        if not d.task_id or d.task_id not in orders_info or d.task_id in expired_tids:
            continue
        info = orders_info[d.task_id]
        g = info["group_name"]
        if g not in jobs:
            src = orders_info[group_meta[g]["set_tid"]]
            jobs[g] = {
                "records": [],
                "cust_code": src["cust_code"],
                "sec_agent": src["sec_agent"],
                "group_name": g,
                "tpl_bytes": src["tpl_bytes"],
                "tpl_code": src["tpl_code"],
            }
        jobs[g]["records"].append({
            "url": d.phone or "",
            "name": d.name or "",
            "province": d.province or "",
            "city": d.city or "",
            "operator": d.operator or "",
            "channel": d.channel or "",
            "platform": d.platform or "",
            "cat1": d.cat1 or "",
            "task_name": d.task_name or "",
            "task_id": d.task_id or "",
            "check_collision": info["check_collision"],
        })

    jobs = list(jobs.values())
    jobs.sort(key=lambda j: j["cust_code"])

    # 6. 计算各代理账单（进货量、销售金额 = 手机号 × 订单单价）
    bills = {}
    if not body.files_only:
        for d in rows:
            if not d.task_id or d.task_id not in orders_info or d.task_id in expired_tids:
                continue
            info = orders_info[d.task_id]
            if (info.get("price") or 0) <= 0:
                continue  # 单价为 0 的订单不生成账单/明细
            cc = info["cust_code"]
            if cc not in bills:
                label = (f"{cc} {info['cust_name']}".strip() if info["cust_name"] else cc)
                cust = customers.get(info["customer_id"])
                bills[cc] = {"customer_id": info["customer_id"], "cust_label": label,
                             "qty": 0, "sales": 0.0, "cost": 0.0,
                             "balance": float(cust.balance) if cust and cust.balance is not None else 0,
                             "groups": {}}
            bills[cc]["qty"] += 1
            eff_price = info["price"] + (0.01 if info["add_name"] else 0)
            bills[cc]["sales"] += eff_price
            bills[cc]["cost"] += info["upstream_price"]
            gg = bills[cc]["groups"].setdefault(d.task_id,
                                                {"task_name": info["task_name"] or info["group_name"],
                                                 "order_no": info["order_no"],
                                                 "channel": info["channel"], "operator": info["operator"],
                                                 "qty": 0, "amount": 0.0, "unit": info["price"], "add_name": info["add_name"]})
            gg["qty"] += 1
            gg["amount"] += eff_price

        # 6.4 产品销售合并（同客户同日并入账单：进货量/销售金额/利润合并，明细行任务名=产品名）
        for ps, prod in db.query(ProductSale, Product).join(
                Product, ProductSale.product_id == Product.id).filter(
                ProductSale.sale_date == date_obj,
                ProductSale.deleted_at.is_(None),
                Product.deleted_at.is_(None)).all():
            cid = ps.customer_id
            if not cid:
                continue
            cc = next((_cc for _cc, b in bills.items() if b.get("customer_id") == cid), None)
            if cc is None:
                cust = customers.get(cid)
                label = (f"{cust.code} {cust.name}".strip() if cust and cust.name else (cust.code or str(cid)))
                cc = cust.code if cust and cust.code else str(cid)
                bills[cc] = {"customer_id": cid, "cust_label": label,
                             "qty": 0, "sales": 0.0, "cost": 0.0,
                             "balance": float(cust.balance) if cust and cust.balance is not None else 0,
                             "groups": {}}
            b = bills[cc]
            pqty = int(ps.qty or 0)
            psales = float(ps.sales_amount or 0)
            b["qty"] += pqty
            b["sales"] += psales
            b["cost"] += round(pqty * float(ps.purchase_price or 0), 2)
            key = f"product_{ps.id}"
            b["groups"][key] = {"task_name": prod.name, "order_no": "/", "channel": "",
                                "operator": "", "qty": pqty, "amount": psales,
                                "unit": float(ps.sales_price or 0), "add_name": None,
                                "is_product": True}

        # 6.5 各代理当日充值/补数信息（金额+说明）
        recharge_map = defaultdict(list)
        for rec in db.query(CustomerRecharge).filter(CustomerRecharge.recharge_date == date_obj).all():
            tlabel = "补数" if (rec.recharge_type or "转账") == "补数" else "充值"
            amt = float(rec.amount_rmb or 0)
            reason = (rec.reason or "").strip()
            s = f"{tlabel} {amt:g}元"
            if reason:
                s += f"（{reason}）"
            recharge_map[rec.customer_id].append(s)
        for cc, b in bills.items():
            b["recharge_info"] = recharge_map.get(b.get("customer_id"), [])

    user_label = user.nickname or user.username
    task_id = uuid.uuid4().hex
    _DISTRIBUTE_TASKS[task_id] = {"total": len(jobs), "done": 0, "status": "running",
                                  "error": "", "buffer": None, "cancel": False,
                                  "date_str": date_obj.strftime("%m%d")}

    def run():
        db2 = SessionLocal()
        try:
            # 7. 保存账单（按客户+日期 upsert；余额 = 客户余额 - 销售金额），并同步客户余额
            if not body.files_only:
                for cc, b in bills.items():
                    cid = b.get("customer_id")
                    if not cid:
                        continue
                    cust = db2.query(Customer).filter(Customer.id == cid).first()
                    existing = db2.query(Bill).filter(Bill.customer_id == cid, Bill.biz_date == date_obj).first()
                    balance = float(cust.balance) if cust and cust.balance is not None else 0
                    # 重复分发时：把上次已扣的销售金额加回来，恢复到扣减前的基础余额，避免重复扣减
                    if existing and existing.sales is not None:
                        balance = round(balance + float(existing.sales), 2)
                    bill_balance = round(balance - float(b["sales"]), 2)
                    b["balance"] = bill_balance
                    profit = round(float(b["sales"]) - float(b["cost"]), 2)
                    detail_json = json.dumps({"groups": b.get("groups", {}),
                                              "recharge_info": b.get("recharge_info", [])},
                                             ensure_ascii=False)
                    if existing:
                        existing.purchase_qty = b["qty"]
                        existing.sales = b["sales"]
                        existing.balance = bill_balance
                        existing.profit = profit
                        existing.detail_json = detail_json
                        existing.updated_by = user_label
                    else:
                        db2.add(Bill(customer_id=cid, biz_date=date_obj, purchase_qty=b["qty"],
                                     sales=b["sales"], balance=bill_balance, profit=profit,
                                     detail_json=detail_json,
                                     created_by=user_label, updated_by=user_label))
                    if cust is not None:
                        cust.balance = bill_balance

                # 7.5 生成甲方账单数据（按甲方+日期 upsert；余额 = 甲方余额 - 销售金额），并同步甲方余额
                for item in _build_upstream_bills(db2, date_obj):
                    up_id = item["up_id"]
                    up = db2.query(Customer).filter(Customer.id == up_id).first()
                    existing = db2.query(Bill).filter(Bill.customer_id == up_id, Bill.biz_date == date_obj).first()
                    balance = float(up.balance) if up and up.balance is not None else 0
                    if existing and existing.sales is not None:
                        balance = round(balance + float(existing.sales), 2)
                    bill_balance = round(balance - float(item["sales"]), 2)
                    up_detail = json.dumps({"groups": item["groups"],
                                            "recharge_info": item["recharge_info"]},
                                           ensure_ascii=False)
                    if existing:
                        existing.purchase_qty = item["qty"]
                        existing.sales = item["sales"]
                        existing.balance = bill_balance
                        existing.profit = round(float(item["sales"]), 2)
                        existing.detail_json = up_detail
                        existing.updated_by = user_label
                    else:
                        db2.add(Bill(customer_id=up_id, biz_date=date_obj, purchase_qty=item["qty"],
                                     sales=item["sales"], balance=bill_balance,
                                     profit=round(float(item["sales"]), 2),
                                     detail_json=up_detail,
                                     created_by=user_label, updated_by=user_label))
                    if up is not None:
                        up.balance = bill_balance
                db2.commit()

            _attach_history(db2, jobs, date_obj)
            buf = _build_distribute_zip(db2, jobs, date_obj.strftime("%m%d"),
                                        lambda done: _DISTRIBUTE_TASKS[task_id].update(done=done),
                                        should_cancel=lambda: _DISTRIBUTE_TASKS[task_id].get("cancel"))
            data = buf.getvalue()
            db2.add(SourceFile(filename=f"分发数据-{date_obj.strftime('%Y-%m-%d')}.zip",
                               file_data=data, file_type="zip", category="出货",
                               biz_date=date_obj.strftime("%Y-%m-%d"), size=len(data)))
            db2.commit()
            _DISTRIBUTE_TASKS[task_id]["buffer"] = buf
            _DISTRIBUTE_TASKS[task_id]["status"] = "done"
        except DistributeCancelled:
            _DISTRIBUTE_TASKS[task_id]["status"] = "cancelled"
        except Exception as e:  # noqa
            _DISTRIBUTE_TASKS[task_id]["status"] = "error"
            _DISTRIBUTE_TASKS[task_id]["error"] = str(e)
        finally:
            db2.close()

    threading.Thread(target=run, daemon=True).start()
    return {"code": 0, "data": {"task_id": task_id, "total": len(jobs)}, "msg": "分数据任务已启动"}


@router.get("/daily-data/distribute/progress/{task_id}")
def distribute_progress(task_id: str, _=Depends(get_current_user)):
    t = _DISTRIBUTE_TASKS.get(task_id)
    if not t:
        raise HTTPException(status_code=404, detail="任务不存在")
    return {"code": 0, "data": {"total": t["total"], "done": t["done"], "status": t["status"], "error": t["error"]}, "msg": "ok"}


@router.post("/daily-data/distribute/cancel/{task_id}")
def distribute_cancel(task_id: str, _=Depends(get_current_user)):
    t = _DISTRIBUTE_TASKS.get(task_id)
    if not t:
        raise HTTPException(status_code=404, detail="任务不存在")
    if t["status"] == "running":
        t["cancel"] = True
    return {"code": 0, "data": None, "msg": "已请求取消"}


@router.get("/daily-data/distribute/download/{task_id}")
def distribute_download(task_id: str, _=Depends(get_current_user)):
    from urllib.parse import quote
    from fastapi.responses import Response
    t = _DISTRIBUTE_TASKS.get(task_id)
    if not t or t["status"] != "done" or not t["buffer"]:
        raise HTTPException(status_code=404, detail="任务未完成或不存在")
    buf = t["buffer"]
    buf.seek(0)
    date_str = t.get("date_str", "")
    return Response(content=buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('分数据-' + date_str + '.zip')}"})


# ============ 源文件管理 ============
@router.get("/source-files")
def list_source_files(db: Session = Depends(get_db), _=Depends(get_current_user),
                      q: str = "", biz_date: str = "", page: int = 1, per_page: int = 10):
    qy = db.query(SourceFile)
    if q:
        qy = qy.filter(SourceFile.filename.like(f"%{q}%"))
    if biz_date:
        qy = qy.filter(SourceFile.biz_date == biz_date)
    total, rows = paginate(qy.order_by(SourceFile.id.desc()), page, per_page)
    data = [{"id": s.id, "filename": s.filename, "file_type": s.file_type,
             "category": s.category or "", "biz_date": s.biz_date, "party": s.party, "size": s.size,
             "created_at": fmt_dt(s.created_at)} for s in rows]
    return ok_page(data, total)


@router.get("/source-files/{sid}/download")
def download_source_file(sid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    from urllib.parse import quote
    from fastapi.responses import Response
    s = db.query(SourceFile).filter(SourceFile.id == sid).first()
    if not s:
        raise HTTPException(status_code=404, detail="文件不存在")
    media = "application/zip" if s.file_type == "zip" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(content=bytes(s.file_data), media_type=media,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(s.filename)}"})


@router.post("/source-files/upload")
async def upload_source_file(file: UploadFile = File(...), category: str = Form("手动"),
                             db: Session = Depends(get_db), _=Depends(get_current_user)):
    fn = file.filename or ""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="文件内容为空")
    ext = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
    cat = (category or "手动").strip() or "手动"
    db.add(SourceFile(filename=fn, file_data=content, file_type=ext or "file",
                      category=cat, size=len(content)))
    db.commit()
    return {"code": 0, "data": None, "msg": "上传成功"}


@router.delete("/source-files/{sid}")
def delete_source_file(sid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    s = db.query(SourceFile).filter(SourceFile.id == sid).first()
    if not s:
        raise HTTPException(status_code=404, detail="文件不存在")
    db.delete(s)
    db.commit()
    return {"code": 0, "data": None, "msg": "已删除"}


# ============ 公积金 ============
@router.get("/fund")
def list_fund(db: Session = Depends(get_db), _=Depends(get_current_user),
              q: str = "", phone: str = "", name: str = "", id_card: str = "", gender: str = "",
              province: str = "", city: str = "", company: str = "", company_type: str = "",
              deposit_status: str = "", operator: str = "", source_file: str = "",
              page: int = 1, per_page: int = 10):
    qy = db.query(Fund)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Fund.phone.like(like), Fund.name.like(like),
                           Fund.company.like(like), Fund.province.like(like), Fund.id_card.like(like)))
    if phone:
        qy = qy.filter(Fund.phone.like(f"%{phone}%"))
    if name:
        qy = qy.filter(Fund.name.like(f"%{name}%"))
    if id_card:
        qy = qy.filter(Fund.id_card.like(f"%{id_card}%"))
    if gender:
        qy = qy.filter(Fund.gender == gender)
    if province:
        qy = qy.filter(Fund.province.like(f"%{province}%"))
    if city:
        qy = qy.filter(Fund.city.like(f"%{city}%"))
    if company:
        qy = qy.filter(Fund.company.like(f"%{company}%"))
    if company_type:
        qy = qy.filter(Fund.company_type.like(f"%{company_type}%"))
    if deposit_status:
        qy = qy.filter(Fund.deposit_status == deposit_status)
    if operator:
        qy = qy.filter(Fund.operator == operator)
    if source_file:
        qy = qy.filter(Fund.source_file.like(f"%{source_file}%"))
    total, rows = paginate(qy.order_by(Fund.id.desc()), page, per_page)
    return ok_page([_fund_dict(f) for f in rows], total)


@router.delete("/fund/{fid}")
def delete_fund(fid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    db.query(Fund).filter(Fund.id == fid).delete()
    db.commit()
    return {"code": 0, "data": None, "msg": "ok"}


@router.post("/fund/import")
async def import_fund(file: UploadFile = File(...), db: Session = Depends(get_db), _=Depends(get_current_user)):
    fn = file.filename or ""
    content = await file.read()
    if fn.lower().endswith(".csv"):
        rows = _parse_fund_csv(content)
    elif fn.lower().endswith((".xlsx", ".xls")):
        rows = _parse_fund_xlsx(content)
    else:
        raise HTTPException(status_code=422, detail="仅支持 csv / xlsx 文件")
    n = 0
    for r in rows:
        phone = r.get("phone")
        if not phone:
            continue
        db.add(Fund(
            phone=phone, name=r.get("name"), id_card=r.get("id_card"), gender=r.get("gender"),
            province=r.get("province"), city=r.get("city"), company=r.get("company"),
            company_type=r.get("company_type"),
            base=_num(r.get("base")), ratio=r.get("ratio"), monthly=_num(r.get("monthly")),
            balance=_num(r.get("balance")), deposit_status=r.get("deposit_status"),
            open_date=r.get("open_date"), pay_to=r.get("pay_to"), operator=r.get("operator"),
            source_file=r.get("source_file") or fn))
        n += 1
    db.commit()
    return {"code": 0, "data": {"imported": n}, "msg": f"导入 {n} 条"}


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _map_row(header, values):
    idx = {str(h).strip(): i for i, h in enumerate(header)}
    row = {}
    for name, field in FUND_HEADER_MAP.items():
        if name in idx and idx[name] < len(values):
            v = values[idx[name]]
            row[field] = None if v is None else str(v).strip()
    return row


def _parse_fund_csv(content):
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    if not header:
        return []
    return [_map_row(header, values) for values in reader]


def _parse_fund_xlsx(content):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    header = None
    for values in ws.iter_rows(values_only=True):
        if header is None:
            header = values
            continue
        rows.append(_map_row(header, values))
    wb.close()
    return rows


def _fund_dict(f: Fund):
    return {
        "id": f.id, "phone": f.phone, "name": f.name, "id_card": f.id_card, "gender": f.gender,
        "province": f.province, "city": f.city, "company": f.company, "company_type": f.company_type,
        "base": float(f.base) if f.base is not None else None, "ratio": f.ratio,
        "monthly": float(f.monthly) if f.monthly is not None else None,
        "balance": float(f.balance) if f.balance is not None else None,
        "deposit_status": f.deposit_status, "open_date": f.open_date, "pay_to": f.pay_to,
        "operator": f.operator, "source_file": f.source_file,
    }


@router.get("/fund/export")
def export_fund(db: Session = Depends(get_db), _=Depends(get_current_user),
                q: str = "", phone: str = "", name: str = "", gender: str = "",
                province: str = "", city: str = "", company: str = "",
                deposit_status: str = ""):
    qy = db.query(Fund)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Fund.phone.like(like), Fund.name.like(like),
                           Fund.company.like(like), Fund.province.like(like)))
    if phone:
        qy = qy.filter(Fund.phone.like(f"%{phone}%"))
    if name:
        qy = qy.filter(Fund.name.like(f"%{name}%"))
    if gender:
        qy = qy.filter(Fund.gender == gender)
    if province:
        qy = qy.filter(Fund.province.like(f"%{province}%"))
    if city:
        qy = qy.filter(Fund.city.like(f"%{city}%"))
    if company:
        qy = qy.filter(Fund.company.like(f"%{company}%"))
    if deposit_status:
        qy = qy.filter(Fund.deposit_status == deposit_status)
    rows = qy.order_by(Fund.id.desc())
    headers = ["手机号", "姓名", "身份证号", "性别", "省份", "地市", "单位名称", "单位性质",
               "缴存基数", "缴存比例", "月缴存额", "账户余额", "缴存状态", "开户日期",
               "缴至年月", "运营商", "来源文件"]

    def gen():
        for f in rows.yield_per(2000):
            yield [f.phone, f.name, f.id_card, f.gender, f.province, f.city, f.company, f.company_type,
                   float(f.base) if f.base is not None else None, f.ratio,
                   float(f.monthly) if f.monthly is not None else None,
                   float(f.balance) if f.balance is not None else None,
                   f.deposit_status, f.open_date, f.pay_to, f.operator, f.source_file]

    return xlsx_response(headers, gen(), "公积金.xlsx", "公积金")


# ============ 账单 ============
@router.get("/bills")
def list_bills(db: Session = Depends(get_db), _=Depends(get_current_user),
               q: str = "", customer: str = "", customer_id: int = 0,
               start_date: str = "", end_date: str = "",
               page: int = 1, per_page: int = 10):
    qy = db.query(Bill).outerjoin(Customer, Bill.customer_id == Customer.id).filter(Bill.deleted_at.is_(None))
    if q:
        qy = qy.filter(or_(Customer.name.like(f"%{q}%"), Customer.code.like(f"%{q}%")))
    if customer:
        qy = qy.filter(Customer.code.like(f"%{customer}%"))
    if customer_id:
        qy = qy.filter(Bill.customer_id == customer_id)
    if start_date:
        qy = qy.filter(Bill.biz_date >= start_date)
    if end_date:
        qy = qy.filter(Bill.biz_date <= end_date)
    total, rows = paginate(qy.order_by(Bill.biz_date.desc()), page, per_page)
    data = []
    for b in rows:
        c = db.query(Customer).filter(Customer.id == b.customer_id).first()
        data.append({
            "id": b.id, "customer": f"{c.code} {c.name}" if c else "", "customer_id": b.customer_id,
            "biz_date": fmt_date(b.biz_date), "purchase_qty": b.purchase_qty,
            "sales": float(b.sales) if b.sales is not None else 0,
            "balance": float(b.balance) if b.balance is not None else None,
            "profit": float(b.profit) if b.profit is not None else 0,
            "created_at": fmt_dt(b.created_at),
            "created_by": b.created_by or "",
            "updated_at": fmt_dt(b.updated_at),
            "updated_by": b.updated_by or "",
        })
    return ok_page(data, total)


def _restore_bill_balance(db, bill):
    """删除账单时恢复客户余额：加回该账单扣减的销售金额"""
    if not bill.customer_id:
        return
    cust = db.query(Customer).filter(Customer.id == bill.customer_id).first()
    if cust and bill.sales is not None:
        cust.balance = round(float(cust.balance or 0) + float(bill.sales), 2)


@router.post("/bills/batch-delete")
def batch_delete_bills(body: BatchStopBody, db: Session = Depends(get_db), _=Depends(require_admin)):
    rows = db.query(Bill).filter(Bill.id.in_(body.ids)).all()
    now = datetime.now()
    for r in rows:
        _restore_bill_balance(db, r)
        r.deleted_at = now
    db.commit()
    return {"code": 0, "data": {"deleted": len(rows)}, "msg": f"已删除 {len(rows)} 条账单"}


@router.delete("/bills")
def clear_bills(db: Session = Depends(get_db), _=Depends(get_current_user)):
    rows = db.query(Bill).filter(Bill.deleted_at.is_(None)).all()
    now = datetime.now()
    for r in rows:
        _restore_bill_balance(db, r)
        r.deleted_at = now
    db.commit()
    return {"code": 0, "data": {"cleared": len(rows)}, "msg": f"已清空 {len(rows)} 条账单"}


@router.post("/bills/import")
async def import_bills(file: UploadFile = File(...), db: Session = Depends(get_db), _=Depends(get_current_user)):
    """导入账单：读取「余额总览」sheet，按编号匹配客户，更新客户余额与账单
    （销售金额=本周出货金额，余额=整体所剩余额；文件名需含日期，如 账单-2026-09-29.xlsx）"""
    fn = file.filename or ""
    content = await file.read()
    if not fn.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=422, detail="仅支持 xlsx / xls 文件")
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", fn)
    if not m:
        raise HTTPException(status_code=422, detail="文件名需包含日期，如 账单-2026-09-29.xlsx")
    try:
        biz_date = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        raise HTTPException(status_code=422, detail="文件名日期无效")

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = None
    for name in wb.sheetnames:
        if "总览" in name or "汇总" in name:
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    header_row = None
    code_col = sales_col = balance_col = None
    for r in range(1, ws.max_row + 1):
        rowvals = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                rowvals[c] = str(v).strip()
        if any("编号" in v for v in rowvals.values()) and any("名称" in v for v in rowvals.values()):
            header_row = r
            for c, v in rowvals.items():
                if "编号" in v:
                    code_col = c
                elif "出货" in v:
                    sales_col = c
                elif "余额" in v and "上月" not in v:
                    balance_col = c
            break

    if header_row is None or code_col is None:
        raise HTTPException(status_code=422, detail="未找到账单汇总表头（需含「编号」「名称」列）")

    updated = 0
    errors = []
    for r in range(header_row + 1, ws.max_row + 1):
        code = str(ws.cell(row=r, column=code_col).value or "").strip() if code_col else ""
        if not code:
            continue
        cust = db.query(Customer).filter(Customer.code == code).first()
        if not cust:
            name = code.replace("甲方", "").strip()
            cust = db.query(Customer).filter(Customer.name == name).first() if name else None
        if not cust:
            errors.append(f"{code}: 客户不存在")
            continue
        sales = _num(ws.cell(row=r, column=sales_col).value) if sales_col else None
        balance = _num(ws.cell(row=r, column=balance_col).value) if balance_col else None
        if sales is None and balance is None:
            errors.append(f"{code}: 出货金额/余额为空（请在 Excel 中打开并保存后再导入，以缓存公式结果）")
            continue
        if balance is not None:
            cust.balance = balance
        bill = db.query(Bill).filter(Bill.customer_id == cust.id, Bill.biz_date == biz_date, Bill.deleted_at.is_(None)).first()
        if bill:
            if sales is not None:
                bill.sales = sales
            if balance is not None:
                bill.balance = balance
        else:
            db.add(Bill(customer_id=cust.id, biz_date=biz_date, purchase_qty=0,
                        sales=(sales if sales is not None else 0),
                        balance=(balance if balance is not None else 0), profit=0))
        updated += 1
    db.commit()
    return {"code": 0, "data": {"updated": updated, "errors": errors},
            "msg": f"已更新 {updated} 个客户账单，失败 {len(errors)} 个"}


@router.get("/bills/export")
def export_bills(db: Session = Depends(get_db), _=Depends(get_current_user),
                 start_date: str = "", end_date: str = "", customer_id: int = 0):
    """按日期期间导出账单明细（含充值记录和余额）；customer_id=0 导出全部代理"""
    from urllib.parse import quote
    from fastapi.responses import Response
    import openpyxl

    cust_map = {c.id: c for c in db.query(Customer).all()}

    qy = db.query(Bill).order_by(Bill.biz_date, Bill.customer_id)
    if customer_id:
        qy = qy.filter(Bill.customer_id == customer_id)
    if start_date:
        qy = qy.filter(Bill.biz_date >= start_date)
    if end_date:
        qy = qy.filter(Bill.biz_date <= end_date)
    bills = qy.all()

    rqy = db.query(CustomerRecharge).order_by(CustomerRecharge.recharge_date, CustomerRecharge.customer_id)
    if customer_id:
        rqy = rqy.filter(CustomerRecharge.customer_id == customer_id)
    if start_date:
        rqy = rqy.filter(CustomerRecharge.recharge_date >= start_date)
    if end_date:
        rqy = rqy.filter(CustomerRecharge.recharge_date <= end_date)
    recharges = rqy.all()

    wb = openpyxl.Workbook()
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    zebra_fill = PatternFill("solid", fgColor="F2F6FC")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    money_fmt = '#,##0.00'

    # ---- 账单明细 ----
    ws1 = wb.active
    ws1.title = "账单明细"
    headers = ["客户编号", "客户名称", "业务日期", "进货量", "销售金额(元)", "余额(元)"]
    widths = [12, 16, 12, 10, 14, 14]
    period = f"{start_date or '...'} ~ {end_date or '...'}"
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    tc = ws1.cell(row=1, column=1, value=f"账单明细（{period}）")
    tc.font = title_font
    tc.alignment = center
    ws1.row_dimensions[1].height = 26
    for j, h in enumerate(headers, 1):
        c = ws1.cell(row=2, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = center
        ws1.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws1.row_dimensions[2].height = 20
    r = 3
    for b in bills:
        c = cust_map.get(b.customer_id)
        row = [
            c.code if c else "", c.name if c else "",
            fmt_date(b.biz_date) if b.biz_date else "",
            b.purchase_qty if b.purchase_qty is not None else 0,
            float(b.sales) if b.sales is not None else 0,
            float(b.balance) if b.balance is not None else 0,
        ]
        for j, v in enumerate(row, 1):
            cell = ws1.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = right if j in (4, 5, 6) else center
            if j in (5, 6):
                cell.number_format = money_fmt
            if r % 2 == 0:
                cell.fill = zebra_fill
        r += 1

    # ---- 充值记录 ----
    ws2 = wb.create_sheet("充值记录")
    headers2 = ["客户编号", "客户名称", "充值日期", "充值金额(U)", "折算人民币(元)", "备注"]
    widths2 = [12, 16, 12, 14, 16, 24]
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers2))
    tc2 = ws2.cell(row=1, column=1, value=f"充值记录（{period}）")
    tc2.font = title_font
    tc2.alignment = center
    ws2.row_dimensions[1].height = 26
    for j, h in enumerate(headers2, 1):
        c = ws2.cell(row=2, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = center
        ws2.column_dimensions[get_column_letter(j)].width = widths2[j - 1]
    ws2.row_dimensions[2].height = 20
    r = 3
    for rec in recharges:
        c = cust_map.get(rec.customer_id)
        row = [
            c.code if c else "", c.name if c else "",
            fmt_date(rec.recharge_date) if rec.recharge_date else "",
            float(rec.amount_u) if rec.amount_u is not None else 0,
            float(rec.amount_rmb) if rec.amount_rmb is not None else 0,
            rec.note or "",
        ]
        for j, v in enumerate(row, 1):
            cell = ws2.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = right if j in (4, 5) else center
            if j in (4, 5):
                cell.number_format = money_fmt
            if r % 2 == 0:
                cell.fill = zebra_fill
        r += 1

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    period = f"{start_date or '起'}-{end_date or '止'}"
    filename = f"账单明细-{period}.xlsx"
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.get("/bills/summary-export")
def export_bills_summary(db: Session = Depends(get_db), _=Depends(get_current_user),
                         date: str = "", start_date: str = "", end_date: str = "",
                         customer_ids: str = ""):
    """导出账单汇总：
    - 特定日期(date)：上游(牛/新按渠道套单价)+下游账单，顶部附当日补数金额/理由；
    - 日期区间(start_date~end_date)：每天上游+下游汇总，末列备注=当天该客户补数理由。
    """
    from urllib.parse import quote
    from fastapi.responses import Response
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    if date:
        date_obj = _parse_biz_date(date)
        if not date_obj:
            raise HTTPException(status_code=422, detail="请选择日期")
        dates = [date_obj]
        mode = "day"
    elif start_date and end_date:
        sd = _parse_biz_date(start_date)
        ed = _parse_biz_date(end_date)
        if not sd or not ed:
            raise HTTPException(status_code=422, detail="请选择日期区间")
        if ed < sd:
            sd, ed = ed, sd
        dates = []
        d = sd
        while d <= ed:
            dates.append(d)
            d += timedelta(days=1)
        mode = "range"
    else:
        raise HTTPException(status_code=422, detail="请选择日期或日期区间")

    up_customers = {c.id: c for c in db.query(Customer).filter(Customer.ctype == "upstream", Customer.status == 1).all()}
    up_name_by_id = {c.id: c.name for c in up_customers.values()}
    ch_by_id = {ch.id: ch.name for ch in db.query(Channel).all()}
    op_by_id = {op.id: op.name for op in db.query(Operator).all()}
    cust_map = {c.id: c for c in db.query(Customer).all()}
    selected_ids = [int(x) for x in customer_ids.split(",") if x.strip().isdigit()] if customer_ids else []
    if selected_ids:
        up_customers = {c.id: c for c in db.query(Customer).filter(
            Customer.ctype == "upstream", Customer.status == 1,
            Customer.id.in_(selected_ids)).all()}
        up_name_by_id = {c.id: c.name for c in up_customers.values()}
        down_customers = db.query(Customer).filter(
            Customer.ctype == "downstream", Customer.status == 1,
            Customer.id.in_(selected_ids)).order_by(Customer.code).all()
    else:
        down_customers = db.query(Customer).filter(Customer.ctype == "downstream", Customer.status == 1) \
            .order_by(Customer.code).all()
    price_map = defaultdict(dict)
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id.in_(up_customers.keys())).all():
        price_map[p.customer_id][(p.channel_id, p.operator_id)] = float(p.price)

    # 充值信息（区间内）：(customer_id, date) -> [类型+金额+理由]
    recharge_info = defaultdict(list)
    for rec in db.query(CustomerRecharge).filter(
            CustomerRecharge.recharge_date >= dates[0],
            CustomerRecharge.recharge_date <= dates[-1]).all():
        tlabel = "补数" if (rec.recharge_type or "转账") == "补数" else "充值"
        amt = float(rec.amount_rmb or 0)
        reason = (rec.reason or "").strip()
        s = f"{tlabel} {amt:g}元"
        if reason:
            s += f"（{reason}）"
        recharge_info[(rec.customer_id, rec.recharge_date)].append(s)

    def channel_label(ch_id, op_id):
        ch_name = ch_by_id.get(ch_id, "")
        op_name = op_by_id.get(op_id, "")
        return f"{ch_name}-{op_name}" if ("dpi" in (ch_name or "").lower() and op_name) else (ch_name or "未知")

    def up_agg_for(d):
        task_ids = [r[0] for r in db.query(DailyData.task_id)
                    .filter(DailyData.biz_date == d, DailyData.task_id.isnot(None)).distinct().all()]
        order_task_map = {}
        for o in db.query(Order).filter(Order.task_id.in_(task_ids)).all():
            order_task_map[o.task_id] = (o.upstream_id, o.channel_id, o.operator_id)
        agg = defaultdict(lambda: defaultdict(int))
        for (tid,) in db.query(DailyData.task_id).filter(DailyData.biz_date == d).all():
            if tid not in order_task_map:
                continue
            up_id, ch_id, op_id = order_task_map[tid]
            agg[up_id][(ch_id, op_id)] += 1
        return agg

    def down_bills_for(d):
        return {b.customer_id: b for b in db.query(Bill).filter(Bill.biz_date == d, Bill.deleted_at.is_(None)).all()}

    wb = openpyxl.Workbook()
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    total_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    money_fmt = '#,##0.00'

    # ================= Sheet1 上游账单 =================
    ws1 = wb.active
    ws1.title = "上游账单"
    if mode == "day":
        d0 = dates[0]
        up_headers = ["甲方", "渠道", "数量", "单价(元)", "金额(元)", "充值信息"]
        ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(up_headers))
        tc = ws1.cell(row=1, column=1, value=f"上游账单汇总（{date}）")
        tc.font = title_font
        tc.alignment = center
        ws1.row_dimensions[1].height = 26
        hr = 2
        for j, h in enumerate(up_headers, 1):
            c = ws1.cell(row=hr, column=j, value=h)
            c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
        r = hr + 1
        agg = up_agg_for(d0)
        total_qty = total_amt = 0
        for up_id in sorted(agg.keys()):
            up_name = up_name_by_id.get(up_id, "")
            combos = agg[up_id]
            sub_qty = sum(combos.values())
            sub_amt = 0.0
            up_note = "；".join(recharge_info.get((up_id, d0), []))
            for (ch_id, op_id) in sorted(combos.keys(), key=lambda x: (x[0] or 0, x[1] or 0)):
                cnt = combos[(ch_id, op_id)]
                label = channel_label(ch_id, op_id)
                price = price_map.get(up_id, {}).get((ch_id, op_id))
                if price is None:
                    price = price_map.get(up_id, {}).get((ch_id, None), 0.0)
                amt = round(cnt * price, 2)
                sub_amt = round(sub_amt + amt, 2)
                vals = [up_name, label, cnt, price, amt, up_note]
                for j, v in enumerate(vals, 1):
                    cell = ws1.cell(row=r, column=j, value=v)
                    cell.border = border
                    cell.alignment = right if j in (3, 4, 5) else (left if j == 6 else center)
                    if j == 5:
                        cell.number_format = money_fmt
                r += 1
            for j, v in enumerate(["", "", sub_qty, "", round(sub_amt, 2), ""], 1):
                cell = ws1.cell(row=r, column=j, value=v)
                cell.font = total_font; cell.fill = total_fill; cell.border = border
                cell.alignment = right if j in (3, 4, 5) else center
                if j == 5:
                    cell.number_format = money_fmt
            r += 1
            total_qty += sub_qty
            total_amt = round(total_amt + sub_amt, 2)
        for j, v in enumerate(["合计", "", total_qty, "", round(total_amt, 2), ""], 1):
            cell = ws1.cell(row=r, column=j, value=v)
            cell.font = total_font; cell.fill = total_fill; cell.border = border
            cell.alignment = right if j in (3, 4, 5) else center
            if j == 5:
                cell.number_format = money_fmt
        for j, w in enumerate([10, 12, 10, 10, 12, 30], 1):
            ws1.column_dimensions[get_column_letter(j)].width = w
    else:
        up_headers = ["日期", "甲方", "渠道", "数量", "单价(元)", "金额(元)", "充值信息"]
        ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(up_headers))
        tc = ws1.cell(row=1, column=1, value=f"上游账单汇总（{start_date} ~ {end_date}）")
        tc.font = title_font
        tc.alignment = center
        ws1.row_dimensions[1].height = 26
        for j, h in enumerate(up_headers, 1):
            c = ws1.cell(row=2, column=j, value=h)
            c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
        r = 3
        for d in dates:
            agg = up_agg_for(d)
            for up_id in sorted(agg.keys()):
                up_name = up_name_by_id.get(up_id, "")
                combos = agg[up_id]
                for (ch_id, op_id) in sorted(combos.keys(), key=lambda x: (x[0] or 0, x[1] or 0)):
                    cnt = combos[(ch_id, op_id)]
                    label = channel_label(ch_id, op_id)
                    price = price_map.get(up_id, {}).get((ch_id, op_id))
                    if price is None:
                        price = price_map.get(up_id, {}).get((ch_id, None), 0.0)
                    amt = round(cnt * price, 2)
                    note = "；".join(recharge_info.get((up_id, d), []))
                    vals = [str(d), up_name, label, cnt, price, amt, note or ""]
                    for j, v in enumerate(vals, 1):
                        cell = ws1.cell(row=r, column=j, value=v)
                        cell.border = border
                        cell.alignment = right if j in (4, 5, 6) else (left if j == 7 else center)
                        if j == 6:
                            cell.number_format = money_fmt
                    r += 1
        for j, w in enumerate([12, 10, 12, 10, 10, 12, 30], 1):
            ws1.column_dimensions[get_column_letter(j)].width = w

    # ================= Sheet2 下游账单 =================
    ws2 = wb.create_sheet("下游账单")
    if mode == "day":
        d0 = dates[0]
        down_bills = down_bills_for(d0)
        down_headers = ["代理编号", "代理名称", "进货量", "销售金额(元)", "利润(元)", "余额(元)", "充值信息"]
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(down_headers))
        tc2 = ws2.cell(row=1, column=1, value=f"下游账单（{date}）")
        tc2.font = title_font
        tc2.alignment = center
        ws2.row_dimensions[1].height = 26
        for j, h in enumerate(down_headers, 1):
            c = ws2.cell(row=2, column=j, value=h)
            c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
        r = 3
        d_qty = d_sales = d_profit = d_bal = 0
        for c in down_customers:
            b = down_bills.get(c.id)
            if b:
                qty = b.purchase_qty or 0
                sales = float(b.sales) if b.sales is not None else 0
                profit = float(b.profit) if b.profit is not None else 0
                balance = float(b.balance) if b.balance is not None else 0
            else:
                qty = 0
                sales = 0.0
                profit = 0.0
                balance = float(c.balance) if c.balance is not None else 0
            note = "；".join(recharge_info.get((c.id, d0), []))
            vals = [c.code if c else "", c.name if c else "", qty, sales, profit, balance, note]
            d_qty += vals[2]; d_sales = round(d_sales + vals[3], 2)
            d_profit = round(d_profit + vals[4], 2); d_bal = round(d_bal + vals[5], 2)
            for j, v in enumerate(vals, 1):
                cell = ws2.cell(row=r, column=j, value=v)
                cell.border = border
                cell.alignment = right if j in (3, 4, 5, 6) else (left if j == 7 else center)
                if j in (4, 5, 6):
                    cell.number_format = money_fmt
            r += 1
        for j, v in enumerate(["合计", "", d_qty, round(d_sales, 2), round(d_profit, 2), round(d_bal, 2), ""], 1):
            cell = ws2.cell(row=r, column=j, value=v)
            cell.font = total_font; cell.fill = total_fill; cell.border = border
            cell.alignment = right if j in (3, 4, 5, 6) else center
            if j in (4, 5, 6):
                cell.number_format = money_fmt
        for j, w in enumerate([12, 16, 10, 14, 14, 14, 30], 1):
            ws2.column_dimensions[get_column_letter(j)].width = w
    else:
        down_headers = ["日期", "代理编号", "代理名称", "进货量", "销售金额(元)", "利润(元)", "余额(元)", "充值信息"]
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(down_headers))
        tc2 = ws2.cell(row=1, column=1, value=f"下游账单（{start_date} ~ {end_date}）")
        tc2.font = title_font
        tc2.alignment = center
        ws2.row_dimensions[1].height = 26
        for j, h in enumerate(down_headers, 1):
            c = ws2.cell(row=2, column=j, value=h)
            c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
        r = 3
        for d in dates:
            down_bills = down_bills_for(d)
            for c in down_customers:
                b = down_bills.get(c.id)
                if b:
                    qty = b.purchase_qty or 0
                    sales = float(b.sales) if b.sales is not None else 0
                    profit = float(b.profit) if b.profit is not None else 0
                    balance = float(b.balance) if b.balance is not None else 0
                else:
                    qty = 0
                    sales = 0.0
                    profit = 0.0
                    balance = float(c.balance) if c.balance is not None else 0
                note = "；".join(recharge_info.get((c.id, d), []))
                vals = [str(d), c.code if c else "", c.name if c else "", qty, sales, profit, balance, note or ""]
                for j, v in enumerate(vals, 1):
                    cell = ws2.cell(row=r, column=j, value=v)
                    cell.border = border
                    cell.alignment = right if j in (4, 5, 6, 7) else (left if j == 8 else center)
                    if j in (5, 6, 7):
                        cell.number_format = money_fmt
                r += 1
        for j, w in enumerate([12, 12, 16, 10, 14, 14, 14, 30], 1):
            ws2.column_dimensions[get_column_letter(j)].width = w

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    period = date.replace('-', '') if date else f"{start_date.replace('-', '')}-{end_date.replace('-', '')}"
    filename = f"账单汇总-{period}.xlsx"
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


def _build_summary_wb_day(db, date_obj, customer_ids=None):
    """生成某日期账单汇总工作簿（上游账单 + 下游账单，与账单汇总导出一致）"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    selected_ids = list(customer_ids) if customer_ids else []
    up_customers = {c.id: c for c in db.query(Customer).filter(Customer.ctype == "upstream", Customer.status == 1).all()}
    if selected_ids:
        up_customers = {c.id: c for c in db.query(Customer).filter(
            Customer.ctype == "upstream", Customer.status == 1,
            Customer.id.in_(selected_ids)).all()}
    up_name_by_id = {c.id: c.name for c in up_customers.values()}
    ch_by_id = {ch.id: ch.name for ch in db.query(Channel).all()}
    op_by_id = {op.id: op.name for op in db.query(Operator).all()}
    if selected_ids:
        down_customers = db.query(Customer).filter(
            Customer.ctype == "downstream", Customer.status == 1,
            Customer.id.in_(selected_ids)).order_by(Customer.code).all()
    else:
        down_customers = db.query(Customer).filter(
            Customer.ctype == "downstream", Customer.status == 1).order_by(Customer.code).all()
    price_map = defaultdict(dict)
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id.in_(up_customers.keys())).all():
        price_map[p.customer_id][(p.channel_id, p.operator_id)] = float(p.price)
    recharge_info = defaultdict(list)
    for rec in db.query(CustomerRecharge).filter(CustomerRecharge.recharge_date == date_obj).all():
        tlabel = "补数" if (rec.recharge_type or "转账") == "补数" else "充值"
        amt = float(rec.amount_rmb or 0)
        reason = (rec.reason or "").strip()
        s = f"{tlabel} {amt:g}元"
        if reason:
            s += f"（{reason}）"
        recharge_info[rec.customer_id].append(s)

    def channel_label(ch_id, op_id):
        ch_name = ch_by_id.get(ch_id, "")
        op_name = op_by_id.get(op_id, "")
        return f"{ch_name}-{op_name}" if ("dpi" in (ch_name or "").lower() and op_name) else (ch_name or "未知")

    task_ids = [r[0] for r in db.query(DailyData.task_id).filter(
        DailyData.biz_date == date_obj, DailyData.task_id.isnot(None)).distinct().all()]
    order_task_map = {}
    for o in db.query(Order).filter(Order.task_id.in_(task_ids)).all():
        order_task_map[o.task_id] = (o.upstream_id, o.channel_id, o.operator_id)
    agg = defaultdict(lambda: defaultdict(int))
    for (tid,) in db.query(DailyData.task_id).filter(DailyData.biz_date == date_obj).all():
        if tid not in order_task_map:
            continue
        up_id, ch_id, op_id = order_task_map[tid]
        if selected_ids and up_id not in selected_ids:
            continue
        agg[up_id][(ch_id, op_id)] += 1
    down_bills = {b.customer_id: b for b in db.query(Bill).filter(Bill.biz_date == date_obj, Bill.deleted_at.is_(None)).all()}

    wb = openpyxl.Workbook()
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    total_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    money_fmt = '#,##0.00'

    date_str = date_obj.strftime("%Y-%m-%d")
    # 上游账单
    ws1 = wb.active
    ws1.title = "上游账单"
    up_headers = ["甲方", "渠道", "数量", "单价(元)", "金额(元)", "充值信息"]
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(up_headers))
    tc = ws1.cell(row=1, column=1, value=f"上游账单汇总（{date_str}）")
    tc.font = title_font; tc.alignment = center
    ws1.row_dimensions[1].height = 26
    for j, h in enumerate(up_headers, 1):
        c = ws1.cell(row=2, column=j, value=h)
        c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
    r = 3
    total_qty = total_amt = 0
    for up_id in sorted(agg.keys()):
        up_name = up_name_by_id.get(up_id, "")
        combos = agg[up_id]
        sub_qty = sum(combos.values())
        sub_amt = 0.0
        up_note = "；".join(recharge_info.get(up_id, []))
        for (ch_id, op_id) in sorted(combos.keys(), key=lambda x: (x[0] or 0, x[1] or 0)):
            cnt = combos[(ch_id, op_id)]
            label = channel_label(ch_id, op_id)
            price = price_map.get(up_id, {}).get((ch_id, op_id))
            if price is None:
                price = price_map.get(up_id, {}).get((ch_id, None), 0.0)
            amt = round(cnt * price, 2)
            sub_amt = round(sub_amt + amt, 2)
            vals = [up_name, label, cnt, price, amt, up_note]
            for j, v in enumerate(vals, 1):
                cell = ws1.cell(row=r, column=j, value=v)
                cell.border = border
                cell.alignment = right if j in (3, 4, 5) else (left if j == 6 else center)
                if j == 5:
                    cell.number_format = money_fmt
            r += 1
        for j, v in enumerate(["", "", sub_qty, "", round(sub_amt, 2), ""], 1):
            cell = ws1.cell(row=r, column=j, value=v)
            cell.font = total_font; cell.fill = total_fill; cell.border = border
            cell.alignment = right if j in (3, 4, 5) else center
            if j == 5:
                cell.number_format = money_fmt
        r += 1
        total_qty += sub_qty
        total_amt = round(total_amt + sub_amt, 2)
    for j, v in enumerate(["合计", "", total_qty, "", round(total_amt, 2), ""], 1):
        cell = ws1.cell(row=r, column=j, value=v)
        cell.font = total_font; cell.fill = total_fill; cell.border = border
        cell.alignment = right if j in (3, 4, 5) else center
        if j == 5:
            cell.number_format = money_fmt
    for j, w in enumerate([10, 12, 10, 10, 12, 30], 1):
        ws1.column_dimensions[get_column_letter(j)].width = w

    # 下游账单
    ws2 = wb.create_sheet("下游账单")
    down_headers = ["代理编号", "代理名称", "进货量", "销售金额(元)", "利润(元)", "余额(元)", "充值信息"]
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(down_headers))
    tc2 = ws2.cell(row=1, column=1, value=f"下游账单（{date_str}）")
    tc2.font = title_font; tc2.alignment = center
    ws2.row_dimensions[1].height = 26
    for j, h in enumerate(down_headers, 1):
        c = ws2.cell(row=2, column=j, value=h)
        c.font = header_font; c.fill = header_fill; c.border = border; c.alignment = center
    r = 3
    d_qty = d_sales = d_profit = d_bal = 0
    for c in down_customers:
        b = down_bills.get(c.id)
        if b:
            qty = b.purchase_qty or 0
            sales = float(b.sales) if b.sales is not None else 0
            profit = float(b.profit) if b.profit is not None else 0
            balance = float(b.balance) if b.balance is not None else 0
        else:
            qty = sales = profit = 0.0
            balance = float(c.balance) if c.balance is not None else 0
        note = "；".join(recharge_info.get(c.id, []))
        vals = [c.code if c else "", c.name if c else "", qty, sales, profit, balance, note]
        d_qty += vals[2]; d_sales = round(d_sales + vals[3], 2)
        d_profit = round(d_profit + vals[4], 2); d_bal = round(d_bal + vals[5], 2)
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = right if j in (3, 4, 5, 6) else (left if j == 7 else center)
            if j in (4, 5, 6):
                cell.number_format = money_fmt
        r += 1
    for j, v in enumerate(["合计", "", d_qty, round(d_sales, 2), round(d_profit, 2), round(d_bal, 2), ""], 1):
        cell = ws2.cell(row=r, column=j, value=v)
        cell.font = total_font; cell.fill = total_fill; cell.border = border
        cell.alignment = right if j in (3, 4, 5, 6) else center
        if j in (4, 5, 6):
            cell.number_format = money_fmt
    for j, w in enumerate([12, 16, 10, 14, 14, 14, 30], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    return wb


def _build_upstream_bills(db, date_obj, customer_ids=None):
    """按甲方计算当日账单数据（按渠道+运营商聚合），返回 [{up_id, up_code, up_name, qty, sales, balance, groups, recharge_info}]"""
    from collections import defaultdict

    selected_ids = list(customer_ids) if customer_ids else []
    up_customers = {c.id: c for c in db.query(Customer).filter(
        Customer.ctype == "upstream", Customer.status == 1, Customer.deleted_at.is_(None)).all()}
    if selected_ids:
        up_customers = {c.id: c for c in db.query(Customer).filter(
            Customer.ctype == "upstream", Customer.status == 1,
            Customer.deleted_at.is_(None),
            Customer.id.in_(selected_ids)).all()}
    up_by_name = {c.name: c for c in up_customers.values() if c.name}
    ch_by_name = {c.name: c for c in db.query(Channel).all()}
    op_by_name = {o.name: o for o in db.query(Operator).all()}
    price_map = defaultdict(dict)
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id.in_(up_customers.keys())).all():
        price_map[p.customer_id][(p.channel_id, p.operator_id)] = float(p.price)

    # 按 甲方名 + 渠道 + 运营商 聚合当日数量（甲方名来自日活数据反查的甲方快照字段）
    up_agg = defaultdict(lambda: defaultdict(int))
    for up_name, ch_name, op_name, cnt in db.query(
            DailyData.upstream, DailyData.channel, DailyData.operator, func.count(DailyData.id)).filter(
            DailyData.biz_date == date_obj).group_by(
            DailyData.upstream, DailyData.channel, DailyData.operator).all():
        if up_name and up_name in up_by_name:
            up_agg[up_name][(ch_name or "", op_name or "")] += int(cnt)

    recharge_map = defaultdict(list)
    for rec in db.query(CustomerRecharge).filter(CustomerRecharge.recharge_date == date_obj).all():
        tlabel = "补数" if (rec.recharge_type or "转账") == "补数" else "充值"
        amt = float(rec.amount_rmb or 0)
        reason = (rec.reason or "").strip()
        s = f"{tlabel} {amt:g}元"
        if reason:
            s += f"（{reason}）"
        recharge_map[rec.customer_id].append(s)

    result = []
    for up_name, ch_op_qty in up_agg.items():
        up = up_by_name[up_name]
        up_id = up.id
        total_qty = 0
        total_amt = 0.0
        groups = {}
        for (ch_name, op_name), cnt in ch_op_qty.items():
            ch = ch_by_name.get(ch_name)
            ch_id = ch.id if ch else None
            op_id = op_by_name.get(op_name).id if op_name else None
            price = price_map[up_id].get((ch_id, op_id))
            if price is None:
                price = price_map[up_id].get((ch_id, None), 0.0)
            amt = round(cnt * price, 2)
            total_amt = round(total_amt + amt, 2)
            total_qty += cnt
            key = f"{ch_name}-{op_name}"
            groups[key] = {"order_no": "", "task_name": key, "qty": cnt,
                           "unit": price, "amount": amt, "add_name": None}
        result.append({
            "up_id": up_id, "up_code": up.code or str(up_id), "up_name": up.name,
            "qty": total_qty, "sales": total_amt,
            "balance": float(up.balance) if up.balance is not None else 0,
            "recharge_info": recharge_map.get(up_id, []),
            "groups": groups,
        })
    return result


def _build_upstream_bill_files(db, summary_start, summary_end, detail_days, date_label, customer_ids=None):
    """按甲方生成账单（MB-032 模版：汇总 + 每日明细），返回 [(filename, bytes)]"""
    tpl = db.query(Template).filter(Template.code == UPSTREAM_BILL_TPL_CODE).first()
    tpl_bytes = bytes(tpl.file_data) if tpl and tpl.file_data else None
    q = db.query(Customer).filter(Customer.ctype == "upstream", Customer.status == 1, Customer.deleted_at.is_(None))
    if customer_ids:
        q = q.filter(Customer.id.in_(customer_ids))
    up_customers = q.order_by(Customer.code).all()
    files = []
    for up in up_customers:
        if not tpl_bytes:
            continue
        has_data = db.query(Bill).filter(
            Bill.customer_id == up.id, Bill.biz_date >= summary_start,
            Bill.biz_date <= summary_end, Bill.deleted_at.is_(None)).first() is not None
        if not has_data:
            has_data = db.query(CustomerRecharge).filter(
                CustomerRecharge.customer_id == up.id,
                CustomerRecharge.recharge_date >= summary_start,
                CustomerRecharge.recharge_date <= summary_end).first() is not None
        if not has_data:
            continue
        bio, end_balance = _fill_mb031_bill(db, up, summary_start, summary_end, detail_days, tpl_bytes, _upstream_groups)
        files.append((f"【账单】{up.code}-{date_label}{_balance_suffix(end_balance)}.xlsx", bio.read()))
    return files


def _parse_mb032_channel(name):
    """把模版渠道名解析为 (渠道名, 运营商)：如 dpi-白-移动 -> (dpi-白, 移动)"""
    for suffix, op in (("-移动", "移动"), ("-联通", "联通"), ("-电信", "电信")):
        if name.endswith(suffix):
            return name[:-len(suffix)], op
    return name, None


def _fill_mb032_bill(db, up, date_obj, tpl_bytes):
    """按 MB-032 模版生成甲方月度进货账单（渠道 × 每日数量矩阵）"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(tpl_bytes))
    ws = wb.active

    month_start = date_obj.replace(day=1)
    if date_obj.month == 12:
        month_end = date(date_obj.year + 1, 1, 1) - timedelta(days=1)
    else:
        month_end = date(date_obj.year, date_obj.month + 1, 1) - timedelta(days=1)
    days_in_month = month_end.day

    ch_by_name = {c.name: c for c in db.query(Channel).all()}
    op_by_name = {o.name: o for o in db.query(Operator).all()}
    prices = {}
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id == up.id).all():
        prices[(p.channel_id, p.operator_id)] = float(p.price)

    # 月度按 (渠道, 运营商) × 日 聚合数量（按日活数据的甲方快照字段匹配）
    rows = db.query(DailyData.biz_date, DailyData.channel, DailyData.operator).filter(
        DailyData.biz_date >= month_start, DailyData.biz_date <= month_end,
        DailyData.upstream == up.name).all()
    qty_map = defaultdict(lambda: defaultdict(int))
    for biz_date, ch_name, op_name in rows:
        qty_map[(ch_name or "", op_name or "")][biz_date.day] += 1

    # 充值(U)/补数(元) 按日
    recharge_u = defaultdict(float)       # 充值 U
    recharge_rmb = defaultdict(float)     # 补数 RMB
    recharge_u_rmb = defaultdict(float)   # 充值 RMB（实际到账金额）
    for rec in db.query(CustomerRecharge).filter(
            CustomerRecharge.customer_id == up.id,
            CustomerRecharge.recharge_date >= month_start,
            CustomerRecharge.recharge_date <= month_end).all():
        d = rec.recharge_date.day
        if (rec.recharge_type or "转账") == "补数":
            recharge_rmb[d] += float(rec.amount_rmb or 0)
        else:
            recharge_u[d] += float(rec.amount_u or 0)
            recharge_u_rmb[d] += float(rec.amount_rmb or 0)

    # 汇率：按系统配置 recharge_rate（U→人民币）
    rate_cfg = db.query(SysConfig).filter(SysConfig.key == "recharge_rate").first()
    try:
        rate = float(rate_cfg.value) if rate_cfg and rate_cfg.value else 6.7
    except (TypeError, ValueError):
        rate = 6.7

    # 标题 + 客户名称
    ws["B1"] = f"LM 进货账单\u3000{date_obj.year} 年 {date_obj.month} 月"
    ws["B2"] = up.name or ""

    # 上月剩余余额：上月最后一张账单余额，取不到默认 0
    prev = db.query(Bill).filter(Bill.customer_id == up.id, Bill.biz_date < month_start, Bill.deleted_at.is_(None)).order_by(Bill.biz_date.desc()).first()
    prev_balance = float(prev.balance) if prev and prev.balance is not None else 0.0
    ws["D2"] = round(prev_balance, 2)
    ws["E2"] = round(prev_balance / rate, 2) if rate else 0.0

    def _day_col(day):
        return 5 + day  # F=6 -> day1

    def _day_qty(ch_name, op_name):
        if op_name:
            return qty_map.get((ch_name, op_name), {})
        # 无运营商拆分的渠道（106/小程序/直播间/公积金白名单/短信回复）：汇总该渠道下所有运营商
        merged = defaultdict(int)
        for (cn, on), dq in qty_map.items():
            if cn == ch_name:
                for day, q in dq.items():
                    merged[day] += q
        return merged

    # 渠道行（第 8~21 行）
    for r in range(8, 22):
        label = ws.cell(row=r, column=2).value
        if label is None:
            continue
        ch_name, op_name = _parse_mb032_channel(str(label).strip())
        ch = ch_by_name.get(ch_name)
        ch_id = ch.id if ch else None
        op_id = op_by_name.get(op_name).id if op_name else None
        price = prices.get((ch_id, op_id))
        if price is None:
            price = prices.get((ch_id, None), 0.0)
        ws.cell(row=r, column=4, value=price)
        day_qty = _day_qty(ch_name, op_name)
        for day in range(1, days_in_month + 1):
            q = day_qty.get(day, 0)
            if q:
                ws.cell(row=r, column=_day_col(day), value=q)

    # 小计行（第 22 行）：每日数量小计 = 各渠道每日之和
    from openpyxl.utils import get_column_letter
    for day in range(1, days_in_month + 1):
        col_letter = get_column_letter(_day_col(day))
        ws.cell(row=22, column=_day_col(day), value=f"=SUM({col_letter}8:{col_letter}21)")

    # 本月充值金额（第 4 行，每日 U）
    for day in range(1, days_in_month + 1):
        u = recharge_u.get(day, 0)
        if u:
            ws.cell(row=4, column=_day_col(day), value=round(u, 4))

    # 补数（元）（第 23 行，每日元）
    for day in range(1, days_in_month + 1):
        rmb = recharge_rmb.get(day, 0)
        if rmb:
            ws.cell(row=23, column=_day_col(day), value=round(rmb, 2))

    # 本月充值金额（人民币）：按实际到账金额合计，覆盖模板写死的 6.7 汇率
    ws["D4"] = round(sum(recharge_u_rmb.values()), 2)
    # U 列换算按系统汇率（覆盖模板 D2/6.7、D3/6.7、D5/6.7）
    ws["E3"] = f"=D3/{rate}"
    ws["E5"] = f"=D5/{rate}"

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio



@router.get("/bills/generate-zip")
def generate_bills_zip(db: Session = Depends(get_db), _=Depends(get_current_user),
                       date: str = "", customer_ids: str = ""):
    """生成账单：汇总文件 + 每个客户(下游)与甲方(上游)的当日账单文件，打包 ZIP 下载"""
    from urllib.parse import quote
    from fastapi.responses import Response
    import zipfile
    if not date:
        raise HTTPException(status_code=422, detail="请选择日期")
    date_obj = _parse_biz_date(date)
    if not date_obj:
        raise HTTPException(status_code=422, detail="日期格式错误")
    selected_ids = [int(x) for x in customer_ids.split(",") if x.strip().isdigit()] if customer_ids else []
    bills = _build_daily_bills(db, date_obj, selected_ids or None)
    mb031_tpl = db.query(Template).filter(Template.code == "MB-031").first()
    mb031_bytes = bytes(mb031_tpl.file_data) if mb031_tpl and mb031_tpl.file_data else None
    cust_map = {c.id: c for c in db.query(Customer).all()}

    buf = io.BytesIO()
    date_str = date_obj.strftime("%m%d")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # 汇总
        wb = _build_summary_wb_day(db, date_obj, selected_ids or None)
        bio = io.BytesIO()
        wb.save(bio)
        zf.writestr(f"账单汇总-{date_obj.strftime('%Y%m%d')}.xlsx", bio.getvalue())
        # 各客户账单（MB-031 模版，与账单管理一致）
        used_names = set()
        for cc, b in bills.items():
            cust = cust_map.get(b.get("customer_id"))
            if mb031_bytes and cust is not None:
                bio, _ = _fill_mb031_bill(db, cust, date_obj.replace(day=1), date_obj, [date_obj], mb031_bytes)
            else:
                bio = _fill_bill_file(b.get("cust_label") or cc, date_str, b, None)
            base = f"【账单】{cc}-{date_obj.strftime('%Y%m%d')}{_balance_suffix(b.get('balance'))}"
            filename = f"{base}.xlsx"
            n = 1
            while filename in used_names:
                n += 1
                filename = f"{base}-{n}.xlsx"
            used_names.add(filename)
            zf.writestr(filename, bio.read())
        # 各甲方账单
        for filename, data in _build_upstream_bill_files(db, date_obj.replace(day=1), date_obj, [date_obj], date_obj.strftime('%Y%m%d'), selected_ids or None):
            zf.writestr(filename, data)
    buf.seek(0)
    return Response(content=buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('账单-' + date_obj.strftime('%Y%m%d') + '.zip')}"})


def _bill_groups(db, cust, day, rt_bills=None):
    """取某客户某日的账单明细 groups（实时重建优先，日活缺失时回退快照 detail_json）"""
    if rt_bills is None:
        rt_bills = _build_daily_bills(db, day, [cust.id])
    g = rt_bills.get(cust.code) or {}
    groups = g.get("groups") or {}
    if not groups:
        b = db.query(Bill).filter(Bill.customer_id == cust.id, Bill.biz_date == day, Bill.deleted_at.is_(None)).first()
        if b and b.detail_json:
            try:
                groups = json.loads(b.detail_json).get("groups") or {}
            except (ValueError, TypeError):
                groups = {}
    return groups


def _upstream_groups(db, up, day, rt_bills=None):
    """甲方某日按工单号聚合的明细 groups（任务名=订单编号、工单号=工单号）"""
    rows = db.query(DailyData).filter(DailyData.biz_date == day, DailyData.upstream == up.name).all()
    tids = {d.task_id for d in rows if d.task_id}
    orders_by_tid = {o.task_id: o for o in db.query(Order).filter(
        Order.task_id.in_(tids)).all()} if tids else {}
    price_map = {}
    for p in db.query(CustomerPrice).filter(CustomerPrice.customer_id == up.id).all():
        price_map[(p.channel_id, p.operator_id)] = float(p.price)
    groups = {}
    for d in rows:
        tid = d.task_id
        if not tid or tid not in orders_by_tid:
            continue
        o = orders_by_tid[tid]
        if tid not in groups:
            price = price_map.get((o.channel_id, o.operator_id))
            if price is None:
                price = price_map.get((o.channel_id, None), 0.0)
            groups[tid] = {"task_name": o.order_no or "", "order_no": tid,
                           "qty": 0, "unit": price, "add_name": bool(o.add_name), "amount": 0.0}
        g = groups[tid]
        g["qty"] += 1
        g["amount"] = round(g["qty"] * g["unit"], 2)
    return groups


def _fill_mb031_bill(db, cust, summary_start, summary_end, detail_days, tpl_bytes, groups_fn=None, rt_bills_by_day=None):
    """MB-031/032 客户账单：sheet1 账单汇总（区间内每日消耗/补账/充值/余额）+ 每日明细 sheet"""
    from openpyxl import load_workbook
    from openpyxl.styles import Border, Side, Alignment
    wb = load_workbook(io.BytesIO(tpl_bytes))
    if len(wb.sheetnames) < 2:
        raise ValueError("MB-031 模版缺少明细 sheet")
    ws1 = wb.worksheets[0]   # 账单汇总
    tpl_ws = wb.worksheets[1]  # yyyymmdd 明细模版

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 期初余额：汇总起始日前的最后一张账单余额，取不到默认 0
    prev = db.query(Bill).filter(Bill.customer_id == cust.id, Bill.biz_date < summary_start, Bill.deleted_at.is_(None)).order_by(Bill.biz_date.desc()).first()
    opening = round(float(prev.balance), 2) if prev and prev.balance is not None else 0.0

    # 汇总期间账单 / 充值补数
    bills = {b.biz_date: b for b in db.query(Bill).filter(
        Bill.customer_id == cust.id, Bill.biz_date >= summary_start,
        Bill.biz_date <= summary_end, Bill.deleted_at.is_(None)).all()}
    bu_by_day = defaultdict(float)          # 补数 RMB
    u_by_day = defaultdict(float)           # 充值 U
    chong_rmb_by_day = defaultdict(float)   # 充值 RMB（余额累计用）
    recharge_info_by_day = defaultdict(list)
    for rec in db.query(CustomerRecharge).filter(
            CustomerRecharge.customer_id == cust.id,
            CustomerRecharge.recharge_date >= summary_start,
            CustomerRecharge.recharge_date <= summary_end).all():
        d = rec.recharge_date
        if (rec.recharge_type or "转账") == "补数":
            amt = float(rec.amount_rmb or 0)
            bu_by_day[d] += amt
            s = f"补数 {amt:g}元"
            if (rec.reason or "").strip():
                s += f"（{rec.reason.strip()}）"
            recharge_info_by_day[d].append(s)
        else:
            u = float(rec.amount_u or 0)
            u_by_day[d] += u
            chong_rmb_by_day[d] += float(rec.amount_rmb or 0)
            s = f"充值 {u:g}U"
            if (rec.note or rec.reason or "").strip():
                s += f"（{(rec.note or rec.reason).strip()}）"
            recharge_info_by_day[d].append(s)

    # 写汇总 sheet（表头第 1 行，数据从第 2 行开始）
    r = 2
    bal = opening
    d = summary_start
    while d <= summary_end:
        b = bills.get(d)
        sales = round(float(b.sales), 2) if b and b.sales is not None else 0.0
        bu = round(bu_by_day.get(d, 0.0), 2)
        u = round(u_by_day.get(d, 0.0), 4)
        if b is not None and b.balance is not None:
            # 以 Bill.balance 为锚点（已含当日销售/充值/手动调整），避免逐日递推漂移
            bal = round(float(b.balance), 2)
        else:
            # 无账单的日期：仅按当日充值/补数递推余额（销售为 0）
            bal = round(bal + bu + chong_rmb_by_day.get(d, 0.0), 2)
        rowvals = [d.strftime("%Y-%m-%d"), sales, bu, u, bal]
        for j, v in enumerate(rowvals, 1):
            cell = ws1.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = center
        r += 1
        d += timedelta(days=1)
    end_balance = bal
    # 删除汇总 sheet 多余的空行（模板预留了 30 行）
    if r <= ws1.max_row:
        ws1.delete_rows(r, ws1.max_row - r + 1)

    # 每日明细 sheet：复制 yyyymmdd 模版，仅生成有明细或充值补数的日期（最近日期在前）
    for day in reversed(detail_days):
        info = recharge_info_by_day.get(day, [])
        rt_bills = rt_bills_by_day.get(day) if rt_bills_by_day else None
        groups = (groups_fn or _bill_groups)(db, cust, day, rt_bills)
        if not groups and not info:
            continue
        new_ws = wb.copy_worksheet(tpl_ws)
        new_ws.title = day.strftime("%Y%m%d")
        c = new_ws["B2"]
        c.value = day.strftime("%Y-%m-%d")
        c.alignment = left_wrap
        c = new_ws["B5"]
        c.value = "\n".join(info)
        c.alignment = left_wrap
        if info:
            # 多行时行高自适应
            new_ws.row_dimensions[5].height = max(21.75, len(info) * 15)
        sorted_groups = sorted(groups.items(), key=lambda kv: _pinyin_key(kv[1].get("task_name") or kv[0]))
        # 明细列结构：MB-031 含「是否洗名」列（金额在第 6 列 F），MB-032 无（金额在第 5 列 E）
        has_add_name = any("洗名" in str(new_ws.cell(row=7, column=j).value or "") for j in range(1, 7))
        amount_col = 6 if has_add_name else 5
        amount_letter = "F" if has_add_name else "E"
        rr = 8
        for g_key, g in sorted_groups:
            if rr > 17:
                # 超过模版预留的 10 行，动态在合计行前插入一行
                new_ws.insert_rows(rr, 1)
            vals = [g.get("task_name") or g_key, g.get("order_no") or "",
                    int(g.get("qty") or 0),
                    round(float(g["unit"]), 4) if g.get("unit") is not None else ""]
            if has_add_name:
                vals.append("/" if g.get("is_product") else ("是" if g.get("add_name") else "否"))
            vals.append(round(float(g.get("amount") or 0), 2))
            for j, v in enumerate(vals, 1):
                cell = new_ws.cell(row=rr, column=j, value=v)
                cell.border = border
                cell.alignment = center
            rr += 1
        # 删除未使用的明细空行（8~17 未填部分），合计行自动上移
        if rr <= 17:
            new_ws.delete_rows(rr, 17 - rr + 1)
        # 合计行现在位于 rr，修复合计公式与 B3/B4 引用
        if rr > 8:
            new_ws.cell(row=rr, column=3, value=f"=SUM(C8:C{rr - 1})")
            new_ws.cell(row=rr, column=amount_col, value=f"=SUM({amount_letter}8:{amount_letter}{rr - 1})")
        else:
            new_ws.cell(row=rr, column=3, value=0)
            new_ws.cell(row=rr, column=amount_col, value=0)
        new_ws["B3"] = f"=C{rr}"
        new_ws["B4"] = f"={amount_letter}{rr}"
    # 删除明细模版 sheet
    if "yyyymmdd" in wb.sheetnames:
        del wb["yyyymmdd"]

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, end_balance


@router.get("/bills/generate-bill")
def generate_bill_files(db: Session = Depends(get_db), _=Depends(get_current_user), start_date: str = "", end_date: str = ""):
    """生成代理账单（MB-031 模版）：可指定日期区间或单日；汇总 sheet + 每日明细 sheet；并附带甲方账单（MB-032）打包 ZIP"""
    from urllib.parse import quote
    from fastapi.responses import Response
    import zipfile
    sd = _parse_biz_date(start_date) if start_date else date.today()
    if not sd:
        raise HTTPException(status_code=422, detail="开始日期格式错误")
    if end_date:
        ed = _parse_biz_date(end_date)
        if not ed:
            raise HTTPException(status_code=422, detail="截止日期格式错误")
        if ed < sd:
            raise HTTPException(status_code=422, detail="截止日期不能早于开始日期")
        summary_start, summary_end = sd, ed
        detail_days = [sd + timedelta(days=i) for i in range((ed - sd).days + 1)]
        date_label = f"{sd.strftime('%Y%m%d')}-{ed.strftime('%Y%m%d')}"
    else:
        summary_start = sd.replace(day=1)
        summary_end = sd
        detail_days = [sd]
        date_label = sd.strftime("%Y%m%d")

    tpl = db.query(Template).filter(Template.code == "MB-031").first()
    if not tpl or not tpl.file_data:
        raise HTTPException(status_code=404, detail="MB-031 账单模版不存在或未上传文件")
    tpl_bytes = bytes(tpl.file_data)
    down_customers = db.query(Customer).filter(Customer.ctype == "downstream", Customer.deleted_at.is_(None)).order_by(Customer.code).all()

    # 预计算区间内每天的实时账单（按日一次，供所有代理共享，避免每个客户每天重复重建）
    rt_bills_by_day = {day: _build_daily_bills(db, day, None) for day in detail_days}

    buf = io.BytesIO()
    used_names = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for cust in down_customers:
            bio, end_balance = _fill_mb031_bill(db, cust, summary_start, summary_end, detail_days, tpl_bytes, rt_bills_by_day=rt_bills_by_day)
            base = f"【账单】{cust.code}-{date_label}{_balance_suffix(end_balance)}"
            filename = f"{base}.xlsx"
            n = 1
            while filename in used_names:
                n += 1
                filename = f"{base}-{n}.xlsx"
            used_names.add(filename)
            zf.writestr(filename, bio.read())
        # 甲方账单文件（MB-032 模版，按月，按结束日生成）
        for filename, data in _build_upstream_bill_files(db, summary_start, summary_end, detail_days, date_label):
            fn = filename
            n = 1
            while fn in used_names:
                n += 1
                fn = f"{filename[:-5]}-{n}.xlsx"
            used_names.add(fn)
            zf.writestr(fn, data)
    buf.seek(0)
    return Response(content=buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('账单-' + date_label + '.zip')}"})


@router.get("/bills/{bid}")
def bill_detail(bid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    b = db.query(Bill).filter(Bill.id == bid).first()
    if not b:
        raise HTTPException(status_code=404, detail="账单不存在")
    c = db.query(Customer).filter(Customer.id == b.customer_id).first()
    rows = (db.query(DailyData, Order)
            .join(Order, DailyData.task_id == Order.task_id)
            .filter(DailyData.biz_date == b.biz_date, Order.customer_id == b.customer_id)
            .all())
    groups = {}
    for d, o in rows:
        ch = db.query(Channel).filter(Channel.id == o.channel_id).first()
        op = db.query(Operator).filter(Operator.id == o.operator_id).first()
        up_price = 0.0
        if o.upstream_id and o.channel_id:
            up_q = db.query(CustomerPrice).filter(CustomerPrice.customer_id == o.upstream_id,
                                                  CustomerPrice.channel_id == o.channel_id)
            up_cp = up_q.filter(CustomerPrice.operator_id == o.operator_id).first()
            if not up_cp:
                up_cp = up_q.filter(CustomerPrice.operator_id.is_(None)).first()
            if up_cp:
                up_price = float(up_cp.price)
        task_label = o.task_name or o.group_name or "未分组"
        gg = groups.setdefault(o.task_id or "", {
            "task_name": task_label,
            "order_no": o.order_no or "",
            "channel": ch.name if ch else "",
            "operator": op.name if op else "",
            "qty": 0, "amount": 0.0,
            "purchase_price": up_price,
            "sales_price": float(o.price) if o.price is not None else 0,
        })
        gg["qty"] += 1
        eff_price = (float(o.price) if o.price is not None else 0) + (0.01 if o.add_name else 0)
        gg["amount"] += eff_price
    # 产品销售明细
    for ps, prod in db.query(ProductSale, Product).join(
            Product, ProductSale.product_id == Product.id).filter(
            ProductSale.sale_date == b.biz_date, ProductSale.customer_id == b.customer_id,
            ProductSale.deleted_at.is_(None)).all():
        groups[f"product_{ps.id}"] = {
            "task_name": prod.name, "order_no": "/", "channel": "", "operator": "",
            "qty": int(ps.qty or 0), "amount": float(ps.sales_amount or 0),
            "purchase_price": float(ps.purchase_price or 0),
            "sales_price": float(ps.sales_price or 0), "is_product": True,
        }
    detail = {
        "id": b.id,
        "customer": (f"{c.code} {c.name}" if c else ""),
        "biz_date": fmt_date(b.biz_date),
        "purchase_qty": b.purchase_qty,
        "sales": float(b.sales) if b.sales is not None else 0,
        "balance": float(b.balance) if b.balance is not None else 0,
        "profit": float(b.profit) if b.profit is not None else 0,
        "created_at": fmt_dt(b.created_at),
        "groups": [{"task_name": gg["task_name"], "task_id": g, "order_no": gg["order_no"],
                    "channel": gg["channel"], "operator": gg["operator"], "qty": gg["qty"],
                    "amount": round(float(gg["amount"]), 2),
                    "purchase_price": round(float(gg.get("purchase_price") or 0), 4),
                    "sales_price": round(float(gg.get("sales_price") or 0), 4)}
                   for g, gg in sorted(groups.items(), key=lambda kv: _pinyin_key(kv[1]["task_name"]))],
    }
    return {"code": 0, "data": detail, "msg": "ok"}


@router.put("/bills/{bid}/balance")
def update_bill_balance(bid: int, body: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    b = db.query(Bill).filter(Bill.id == bid).first()
    if not b:
        raise HTTPException(status_code=404, detail="账单不存在")
    try:
        b.balance = round(float(body.get("balance")), 2)
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="余额格式不正确")
    b.updated_by = user.nickname or user.username
    db.commit()
    return {"code": 0, "data": None, "msg": "余额已更新"}


# ============ 预警 ============
@router.get("/alerts")
def list_alerts(db: Session = Depends(get_db), _=Depends(get_current_user),
                q: str = "", level: str = "", type: str = "", status: str = "未处理",
                customer_name: str = "", order_no: str = "", task_name: str = "", content: str = "",
                page: int = 1, per_page: int = 10):
    qy = db.query(Alert).outerjoin(Customer, Alert.customer_id == Customer.id)
    if q:
        like = f"%{q}%"
        qy = qy.filter(or_(Alert.content.like(like), Alert.task_name.like(like),
                           Customer.name.like(like)))
    if customer_name:
        like = f"%{customer_name}%"
        qy = qy.filter(or_(Customer.name.like(like), Customer.code.like(like)))
    if order_no:
        qy = qy.filter(Alert.order_no.like(f"%{order_no}%"))
    if task_name:
        qy = qy.filter(Alert.task_name.like(f"%{task_name}%"))
    if content:
        qy = qy.filter(Alert.content.like(f"%{content}%"))
    if level:
        qy = qy.filter(Alert.level == level)
    if type:
        qy = qy.filter(Alert.type == type)
    if status:
        qy = qy.filter(Alert.status == status)
    total, rows = paginate(qy.order_by(Alert.id.desc()), page, per_page)
    data = []
    for a in rows:
        c = db.query(Customer).filter(Customer.id == a.customer_id).first()
        data.append({
            "id": a.id, "type": a.type, "customer_id": a.customer_id,
            "customer_name": f"{c.code} {c.name}" if c else "", "task_name": a.task_name,
            "order_no": a.order_no or "", "order_id": a.order_id,
            "content": a.content, "time": fmt_dt(a.trigger_time),
            "status": a.status,
        })
    return ok_page(data, total)


@router.put("/alerts/{aid}")
def handle_alert(aid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    a = db.query(Alert).filter(Alert.id == aid).first()
    if not a:
        raise HTTPException(status_code=404, detail="预警不存在")
    a.status = "已处理"
    db.commit()
    return {"code": 0, "data": None, "msg": "已处理"}


def _alert_days(db):
    cfg = db.query(SysConfig).filter(SysConfig.key == "alert_days").first()
    try:
        return int(cfg.value) if cfg and cfg.value else 0
    except (TypeError, ValueError):
        return 0


@router.get("/alert-config")
def get_alert_config(db: Session = Depends(get_db), _=Depends(get_current_user)):
    cfg = db.query(SysConfig).filter(SysConfig.key == "alert_trigger_time").first()
    return {"code": 0, "data": {"trigger_time": cfg.value if cfg else "00:00",
                                "alert_days": _alert_days(db)}, "msg": "ok"}


@router.put("/alert-config")
def set_alert_config(body: dict, db: Session = Depends(get_db), _=Depends(get_current_user)):
    value = (body.get("trigger_time") or "00:00").strip()
    cfg = db.query(SysConfig).filter(SysConfig.key == "alert_trigger_time").first()
    if cfg:
        cfg.value = value
    else:
        db.add(SysConfig(key="alert_trigger_time", value=value))
    if "alert_days" in body:
        try:
            days = int(body.get("alert_days"))
        except (TypeError, ValueError):
            days = 0
        dc = db.query(SysConfig).filter(SysConfig.key == "alert_days").first()
        if dc:
            dc.value = str(days)
        else:
            db.add(SysConfig(key="alert_days", value=str(days)))
    db.commit()
    return {"code": 0, "data": None, "msg": "已保存"}


@router.post("/alerts/reset")
def reset_alerts(body: dict, db: Session = Depends(get_db), _=Depends(require_admin)):
    """重置预警：按指定日期清除预警数据后重新扫描生成新预警"""
    date_val = (body.get("date") or "").strip()
    scope = (body.get("scope") or "stop").strip()
    date_obj = _parse_biz_date(date_val) if date_val else None
    q = db.query(Alert)
    if scope != "all":
        q = q.filter(Alert.type == "停单提醒")
    if date_obj:
        cutoff = datetime.combine(date_obj + timedelta(days=1), datetime.min.time())
        q = q.filter(Alert.trigger_time < cutoff)
    deleted = q.delete(synchronize_session=False)
    db.commit()
    run_alert_scan(db)
    remaining = db.query(func.count(Alert.id)).filter(Alert.status == "未处理").scalar() or 0
    return {"code": 0, "data": {"deleted": deleted, "remaining": int(remaining)},
            "msg": f"已清除 {deleted} 条预警并重新扫描，当前未处理 {remaining} 条"}


@router.get("/daily-data/filter-options")
def daily_filter_options(db: Session = Depends(get_db), _=Depends(get_current_user)):
    """日活数据筛选级联选项：品类（一级→二级→平台）、一级代理→二级代理/渠道"""
    # 一级品类 → 二级品类
    cat1 = []
    cat2_map = {}
    for c in db.query(Category).filter(Category.level == 1).order_by(Category.id).all():
        cat1.append(c.name)
        children = db.query(Category).filter(Category.level == 2, Category.parent_id == c.id).order_by(Category.id).all()
        cat2_map[c.name] = [x.name for x in children]
    # 二级品类 → 平台
    platform_map = {}
    for p in db.query(Platform).order_by(Platform.id).all():
        c2 = db.query(Category).filter(Category.id == p.cat_id).first()
        if c2:
            platform_map.setdefault(c2.name, []).append(p.name)
    # 一级代理 → 二级代理 / 渠道（customer 存储为「编号 名称」，用编号作为级联键）
    agent_map = {}
    channel_map = {}
    for cu, sa in db.query(DailyData.customer, DailyData.secondary_agent).filter(
            DailyData.customer.isnot(None), DailyData.customer != "",
            DailyData.secondary_agent.isnot(None), DailyData.secondary_agent != "").distinct().all():
        agent_map.setdefault(cu.split()[0], []).append(sa)
    for cu, ch in db.query(DailyData.customer, DailyData.channel).filter(
            DailyData.customer.isnot(None), DailyData.customer != "",
            DailyData.channel.isnot(None), DailyData.channel != "").distinct().all():
        channel_map.setdefault(cu.split()[0], []).append(ch)
    for k in agent_map:
        agent_map[k] = sorted(set(agent_map[k]))
    for k in channel_map:
        channel_map[k] = sorted(set(channel_map[k]))
    group_name = sorted({g for g, in db.query(DailyData.group_name).filter(
        DailyData.group_name.isnot(None), DailyData.group_name != "").distinct().all()})
    return {"code": 0, "data": {
        "cat1": cat1, "cat2_map": cat2_map, "platform_map": platform_map,
        "agent_map": agent_map, "channel_map": channel_map, "group_name": group_name,
    }, "msg": "ok"}


@router.get("/distinct")
def distinct_values(db: Session = Depends(get_db), _=Depends(get_current_user),
                    model: str = "", field: str = ""):
    allowed = {
        "daily_data": {"province", "city", "platform", "upstream", "wash_status"},
        "orders": {"province", "city", "start_date", "end_date"},
        "fund": {"province", "city", "company_type", "operator", "source_file"},
    }
    if model not in allowed or field not in allowed[model]:
        return {"code": 0, "data": [], "msg": "ok"}
    table = {"daily_data": DailyData, "orders": Order, "fund": Fund, "customer": Customer}[model]
    col = getattr(table, field)
    # 日期字段不参与「空字符串」过滤（Date 列无法与 "" 比较）
    if field in ("start_date", "end_date"):
        rows = db.query(col).filter(col.isnot(None)).distinct().order_by(col).all()
    else:
        rows = db.query(col).filter(col.isnot(None), col != "").distinct().order_by(col).all()
    return {"code": 0, "data": [str(r[0]) if r[0] is not None else "" for r in rows], "msg": "ok"}


def _apply_daily_filters(qy, f):
    f = f or {}
    if f.get("q"):
        like = f"%{f['q']}%"
        qy = qy.filter(or_(DailyData.phone.like(like), DailyData.task_name.like(like),
                           DailyData.upstream.like(like), DailyData.task_id.like(like)))
    if f.get("phone"):
        qy = qy.filter(DailyData.phone.like(f"%{f['phone']}%"))
    if f.get("upstream"):
        qy = qy.filter(DailyData.upstream.like(f"%{f['upstream']}%"))
    if f.get("task_id"):
        qy = qy.filter(DailyData.task_id.like(f"%{f['task_id']}%"))
    if f.get("task_name"):
        qy = qy.filter(DailyData.task_name.like(f"%{f['task_name']}%"))
    if f.get("province"):
        provs = [p for p in f["province"].split(",") if p.strip()]
        if provs:
            qy = qy.filter(or_(*[DailyData.province.like(f"%{p}%") for p in provs]))
    if f.get("city"):
        cities = [c for c in f["city"].split(",") if c.strip()]
        if cities:
            qy = qy.filter(or_(*[DailyData.city.like(f"%{c}%") for c in cities]))
    if f.get("operator"):
        qy = qy.filter(DailyData.operator.like(f"%{f['operator']}%"))
    if f.get("cat1"):
        qy = qy.filter(DailyData.cat1.like(f"%{f['cat1']}%"))
    if f.get("cat2"):
        qy = qy.filter(DailyData.cat2.like(f"%{f['cat2']}%"))
    if f.get("platform"):
        qy = qy.filter(DailyData.platform.like(f"%{f['platform']}%"))
    if f.get("platform_empty") in ("1", "有", "是", "true", "True", "yes", "YES"):
        qy = qy.filter(or_(DailyData.platform.is_(None), DailyData.platform == ""))
    if f.get("customer"):
        qy = qy.filter(DailyData.customer.like(f"%{f['customer']}%"))
    if f.get("secondary_agent"):
        qy = qy.filter(DailyData.secondary_agent.like(f"%{f['secondary_agent']}%"))
    if f.get("channel"):
        qy = qy.filter(DailyData.channel.like(f"%{f['channel']}%"))
    if f.get("group_name"):
        qy = qy.filter(DailyData.group_name.like(f"%{f['group_name']}%"))
    if f.get("source_file"):
        qy = qy.filter(DailyData.source_file.like(f"%{f['source_file']}%"))
    if f.get("name"):
        qy = qy.filter(DailyData.name.like(f"%{f['name']}%"))
    if f.get("biz_date"):
        qy = qy.filter(DailyData.biz_date == f["biz_date"])
    if f.get("start_date"):
        sd = _parse_biz_date(f["start_date"])
        if sd:
            qy = qy.filter(DailyData.biz_date >= sd)
    if f.get("end_date"):
        ed = _parse_biz_date(f["end_date"])
        if ed:
            qy = qy.filter(DailyData.biz_date <= ed)
    return qy


@router.post("/daily-data/extract")
def extract_daily(body: ExtractBody, db: Session = Depends(get_db), _=Depends(get_current_user)):
    """单独抽取：按条件/时间段抽取日活数据，生成数据文件与账单文件，并单独记一笔账单"""
    from urllib.parse import quote
    from fastapi.responses import Response
    import zipfile
    import openpyxl

    qy = _apply_daily_filters(db.query(DailyData), body.filters)
    if body.start_date:
        sd = _parse_biz_date(body.start_date)
        if sd:
            qy = qy.filter(DailyData.biz_date >= sd)
    if body.end_date:
        ed = _parse_biz_date(body.end_date)
        if ed:
            qy = qy.filter(DailyData.biz_date <= ed)
    rows = qy.order_by(DailyData.id.desc()).all()
    if body.limit and body.limit > 0:
        rows = rows[:body.limit]
    if not rows:
        raise HTTPException(status_code=404, detail="没有符合条件的数据")

    qty = len(rows)
    price = float(body.price or 0)
    sales = round(price * qty, 2)

    cust = db.query(Customer).filter(Customer.id == body.customer_id).first() if body.customer_id else None
    target_label = (f"{cust.code} {cust.name}" if cust else (body.target_name or "其他人"))

    balance = 0
    if cust is not None:
        balance = round((float(cust.balance) if cust.balance is not None else 0) - sales, 2)
        cust.balance = balance
    db.add(Bill(customer_id=body.customer_id, biz_date=date.today(),
                purchase_qty=qty, sales=sales, balance=balance,
                profit=sales, bill_type="单独抽取",
                target_name=(body.target_name or None) if not cust else None))
    db.commit()

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        dwb = openpyxl.Workbook()
        dws = dwb.active
        dws.title = "抽取数据"
        dheaders = ["数据日期", "甲方", "任务id", "任务名", "手机号", "姓名", "省", "市", "运营商", "渠道", "平台"]
        dws.append(dheaders)
        for d in rows:
            dws.append([fmt_date(d.biz_date) if d.biz_date else "", d.upstream or "", d.task_id or "",
                        d.task_name or "", d.phone or "", d.name or "", d.province or "", d.city or "",
                        d.operator or "", d.channel or "", d.platform or ""])
        dbio = io.BytesIO()
        dwb.save(dbio)
        zf.writestr("抽取数据.xlsx", dbio.getvalue())

        bwb = openpyxl.Workbook()
        bws = bwb.active
        bws.title = "账单"
        bws.append(["收款对象", target_label])
        bws.append(["日期", date.today().strftime("%Y-%m-%d")])
        bws.append(["进货量", qty])
        bws.append(["单价(元/条)", price])
        bws.append(["销售金额(元)", sales])
        bws.append(["利润(元)", sales])
        if body.note:
            bws.append(["备注", body.note])
        bbio = io.BytesIO()
        bwb.save(bbio)
        zf.writestr("账单.xlsx", bbio.getvalue())

    bio.seek(0)
    filename = f"单独抽取-{date.today().strftime('%Y%m%d')}.zip"
    return Response(content=bio.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


def _order_alert_days(db):
    """订单预警天数（截止日期在该天数内到期即提醒，0 天=仅截止已过期）"""
    cfg = db.query(SysConfig).filter(SysConfig.key == "order_alert_days").first()
    try:
        return int(cfg.value) if cfg and cfg.value else 0
    except (TypeError, ValueError):
        return 0


def _regenerate_stop_alerts(db, with_notify=False):
    """清除未处理停单提醒，再按「订单预警天数」重新生成（在执 + 截止日期 <= 今天+预警天数）"""
    today = date.today()
    days = _order_alert_days(db)
    target = today + timedelta(days=days)
    cleared = db.query(Alert).filter(Alert.type == "停单提醒",
                                     Alert.status == "未处理").delete(synchronize_session=False)
    created = 0
    for o in db.query(Order).filter(Order.status == "在执", Order.end_date.isnot(None),
                                    Order.end_date <= target).all():
        db.add(Alert(level="提醒", type="停单提醒", customer_id=o.customer_id,
                     task_name=o.task_name, order_no=o.order_no, order_id=o.id,
                     content=f"截止时间 {o.end_date} 已到期（{o.task_name or o.order_no}），请确认是否处理",
                     trigger_time=datetime.now()))
        if with_notify:
            task_label = o.task_name or o.order_no or ""
            if o.task_id:
                task_label += f"（{o.task_id}）"
            end_str = o.end_date.strftime("%m/%d") if o.end_date else ""
            db.add(Notification(order_no=o.order_no, task_id=o.task_id, order_id=o.id,
                                content=f"{task_label}于 {end_str} 到期，请及时处理停单。",
                                notify_time=datetime.now()))
        created += 1
    return cleared, created


def run_alert_scan(db: Session):
    """定时：按「订单预警天数」重建停单提醒（截止日期到期）；余额低于预警额度 -> 账单预警"""
    _regenerate_stop_alerts(db, with_notify=True)
    # 账单预警
    for c in db.query(Customer).filter(Customer.ctype == "downstream", Customer.status == 1).all():
        if c.warn_amount and (c.balance or 0) < c.warn_amount:
            if not db.query(Alert).filter(Alert.type == "账单预警", Alert.customer_id == c.id,
                                          Alert.status == "未处理").first():
                db.add(Alert(level="预警", type="账单预警", customer_id=c.id,
                             content=f"余额 {c.balance} 已低于预警额度 {c.warn_amount}",
                             trigger_time=datetime.now()))
    # 自动标记已解决的账单预警（余额已补足）
    for a in db.query(Alert).filter(Alert.type == "账单预警", Alert.status == "未处理").all():
        c = db.query(Customer).filter(Customer.id == a.customer_id).first()
        if not c or not c.warn_amount or (c.balance or 0) >= c.warn_amount:
            a.status = "已处理"
    db.commit()


def _sync_order_stop_alert(db: Session, o):
    """同步订单停单提醒：数据变化后，截止日期>今天则取消预警；在执且截止日期<=今天则创建预警"""
    today = date.today()
    should = o.status == "在执" and o.end_date is not None and o.end_date <= today
    existing = db.query(Alert).filter(Alert.type == "停单提醒", Alert.status == "未处理",
                                      Alert.order_id == o.id).first()
    if should and not existing:
        db.add(Alert(level="提醒", type="停单提醒", customer_id=o.customer_id,
                     task_name=o.task_name, order_no=o.order_no, order_id=o.id,
                     content=f"截止时间 {o.end_date} 已到期（{o.task_name or o.order_no}），请确认是否处理",
                     trigger_time=datetime.now()))
        return True
    if not should and existing:
        existing.status = "已处理"
    return False
